"""FR-021: a report written with render_xlsx reads back with openpyxl (Step 8).

The report uses the bundled IPAex Mincho, the default registry and the
library's components end to end; the expectations are what the mapping
promises (cells, numbers, rich text, merges, borders, page setup, header and
footer), read from the file rather than from the model.
"""

import re
import zipfile

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from PIL import Image as PILImage
import pytest

from gridlayoutpdf import (
    A3_LANDSCAPE_REFERENCE, A4_PORTRAIT_REFERENCE, Cell, Document, Equation, EquationAnchors, EquationGroup,
    Figure, Fraction, Image, Label, MathBlock, Operator, PageDecoration, PageProfile, PageSection, Paragraph,
    Row, Rule, Script, Table, TextBox, render_xlsx,
)

PADDING = 2.4
A4 = PageProfile("A4 body", 595, 842, A4_PORTRAIT_REFERENCE.grid, 2, 45)
A3 = PageProfile("A3 body", 1191, 842, A3_LANDSCAPE_REFERENCE.grid, 2, 45)


def padding(size):
    return PADDING * max(1, size / 12)


def report(directory):
    image = directory / "dot.png"
    PILImage.new("RGB", (8, 8), "blue").save(image)
    return Document(
        Paragraph("1. 断面の諸元"),
        EquationGroup(
            Equation("I", Fraction(Row("4.50", Operator("×"), Script("2.70", superscript="3")), "12"),
                     result="7.381", unit=Script("m", superscript="4")),
            anchors=EquationAnchors(3, 15, 21, 22)),
        Paragraph("−2.000"),
        Table([4, 4], [[Cell("荷重"), Cell("12.150")]], numeric_columns=(1,)),
        Rule("medium"),
        PageSection(A3, Paragraph("横の頁")),
        Figure(10, 5, TextBox(4, 4, 80, 30, ("箱",)), Label(100, 20, "ラベル")),
        Image(image, col_span=2, rows=2),
        profile=A4,
        header=PageDecoration("計算書 {page}/{pages}", align="right"),
        footer=PageDecoration("- {page} -", align="center"),
    )


@pytest.fixture(scope="module")
def written(tmp_path_factory):
    directory = tmp_path_factory.mktemp("xlsx")
    path = directory / "report.xlsx"
    result = report(directory).render_xlsx(path, "断面", cell_padding_pt=padding)
    return path, result


def load(path, **options):
    return openpyxl.load_workbook(path, **options)


def by_value(sheet, value):
    return next(cell for row in sheet.iter_rows() for cell in row if cell.value == value)


def test_sheets_follow_the_paper_changes(written):
    path, result = written
    workbook = load(path)
    assert workbook.sheetnames == ["断面", "断面_2", "断面_3"]
    assert [(info.first_page, info.last_page) for info in result.sheets] == [(1, 1), (2, 2), (3, 3)]
    assert not result.fallbacks and not result.untagged


def test_cells_numbers_and_merges(written):
    path, _ = written
    sheet = load(path)["断面"]
    heading = by_value(sheet, "1. 断面の諸元")
    assert heading.font.name == "IPAex明朝" and heading.font.size == 10
    minus = by_value(sheet, -2.0)
    assert minus.number_format == '0.000;"−"0.000' and minus.alignment.horizontal == "left"
    amount = by_value(sheet, 12.15)
    assert amount.number_format == "0.000" and amount.alignment.horizontal == "right"
    lhs = by_value(sheet, "I")
    merged = next(str(merge) for merge in sheet.merged_cells.ranges if lhs.coordinate in merge)
    first, last = re.findall(r"\d+", merged)
    assert int(last) == int(first) + 1  # merged over the numerator and denominator rows
    assert lhs.alignment.vertical == "center" and lhs.alignment.horizontal == "right"


def test_scripts_are_rich_text_runs(written):
    path, _ = written
    sheet = load(path, rich_text=True)["断面"]
    numerator = next(cell for row in sheet.iter_rows() for cell in row if isinstance(cell.value, CellRichText))
    blocks = [block for block in numerator.value if isinstance(block, TextBlock)]
    assert "".join(str(block.text) for block in blocks).endswith("3")
    assert blocks[-1].font.vertAlign == "superscript"


def test_borders_come_from_the_table_and_the_rule(written):
    path, _ = written
    sheet = load(path)["断面"]
    amount = by_value(sheet, 12.15)
    styles = {cell.border.top.style for row in sheet.iter_rows() for cell in row} - {None}
    assert {"thin", "medium"} <= styles
    assert sheet.cell(amount.row, 1).border.left.style == "thin"


def test_page_setup_header_and_footer(written):
    path, _ = written
    workbook = load(path)
    first, second = workbook["断面"], workbook["断面_2"]
    assert int(first.page_setup.paperSize) == int(first.PAPERSIZE_A4) and first.page_setup.orientation == "portrait"
    assert int(second.page_setup.paperSize) == int(second.PAPERSIZE_A3) and second.page_setup.orientation == "landscape"
    assert first.page_setup.scale == 100
    assert first.page_margins.left == pytest.approx(50.88 / 72) and first.page_margins.top == pytest.approx(54.598 / 72)
    assert "&P/3" in first.oddHeader.right.text and "計算書 " in first.oddHeader.right.text
    assert first.oddFooter.center.text.startswith("- &P")
    assert int(second.page_setup.firstPageNumber) == 2 and second.page_setup.useFirstPageNumber
    assert first.row_dimensions[47].height == pytest.approx(15.75) and first.column_dimensions["AD"].width
    assert first.sheet_view.showGridLines in (None, True)  # shown on screen (the default)


def test_figures_and_images_are_drawings(written):
    path, _ = written
    with zipfile.ZipFile(path) as package:
        names = package.namelist()
        drawing = next(package.read(name).decode() for name in names
                       if name.startswith("xl/drawings/drawing") and name.endswith(".xml") and
                       "箱" in package.read(name).decode())
    assert "<xdr:grpSp>" in drawing and "ラベル" in drawing
    assert any(name.startswith("xl/media/") for name in names)


def test_the_same_report_gives_the_same_bytes(written, tmp_path):
    path, _ = written
    again = tmp_path / "again.xlsx"
    report(tmp_path).render_xlsx(again, "断面", cell_padding_pt=padding)
    assert again.read_bytes() == path.read_bytes()
