"""DrawingML shapes: kinds, cell anchors and EMU geometry read from the XML."""

import io
import struct
import xml.etree.ElementTree as ET
import zipfile
import zlib

import openpyxl
import pytest
from PIL import Image as PILImage

from gridlayoutpdf.renderers.xlsx._model import (
    Arrowhead,
    Diamond,
    Ellipse,
    Font,
    Group,
    Line,
    Picture,
    Polyline,
    Rect,
    Run,
    Sheet,
    TextBox,
    TextFrame,
    Workbook,
)
from gridlayoutpdf.renderers.xlsx._package import write_workbook

NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
EMU = 12_700
MINCHO = Font("IPAex明朝", 10.5)
# 16.08 pt columns are 21 px = 15.75 pt in Excel; 15.48 pt rows are 21 px = 15.75 pt.
WIDTH, HEIGHT = 16.08, 15.48
ACTUAL = 15.75


def _png(width=2, height=1):
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _jpeg():
    buffer = io.BytesIO()
    PILImage.new("RGB", (2, 1), (0, 128, 255)).save(buffer, "JPEG")
    return buffer.getvalue()


JPEG = _jpeg()


def _sheet(drawings, name="S", columns=10, rows=10):
    return Sheet(name, (WIDTH,) * columns, (HEIGHT,) * rows, drawings=drawings)


def _archive(*sheets):
    return zipfile.ZipFile(io.BytesIO(write_workbook(Workbook(sheets))))


def _drawing(drawings, index=1):
    with _archive(_sheet(drawings)) as archive:
        return ET.fromstring(archive.read(f"xl/drawings/drawing{index}.xml"))


def _anchor_cells(anchor):
    def point(tag):
        node = anchor.find(f"xdr:{tag}", NS)
        return tuple(int(node.find(f"xdr:{name}", NS).text) for name in ("col", "colOff", "row", "rowOff"))
    return point("from"), point("to")


def _xfrm(node):
    xfrm = node.find(".//a:xfrm", NS)
    off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
    return (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")),
            xfrm.get("flipH"), xfrm.get("flipV"))


def test_sheet_without_drawings_has_no_drawing_part():
    with _archive(Sheet("S", (WIDTH,), (HEIGHT,))) as archive:
        assert not [name for name in archive.namelist() if "drawing" in name]


def test_drawing_parts_relationships_and_content_types():
    with _archive(_sheet((Rect(0, 0, WIDTH, HEIGHT),)), _sheet((), name="T"),
                  _sheet((Line(0, 0, 10, 0),), name="U")) as archive:
        names = archive.namelist()
        assert "xl/drawings/drawing1.xml" in names and "xl/drawings/drawing2.xml" in names
        assert "xl/drawings/drawing3.xml" not in names
        assert 'Target="../drawings/drawing2.xml"' in archive.read("xl/worksheets/_rels/sheet3.xml.rels").decode()
        assert "xl/worksheets/_rels/sheet2.xml.rels" not in names
        assert '<drawing r:id="rId1"/>' in archive.read("xl/worksheets/sheet1.xml").decode()
        types = archive.read("[Content_Types].xml").decode()
        assert 'PartName="/xl/drawings/drawing1.xml"' in types
        assert "application/vnd.openxmlformats-officedocument.drawing+xml" in types


def test_cell_anchor_keeps_the_fraction_of_the_cell_in_excel_pixels():
    # x = 1.5 columns, y = 2.25 rows in the model's pt; Excel's cells are 15.75 pt.
    root = _drawing((Rect(1.5 * WIDTH, 2.25 * HEIGHT, 2 * WIDTH, HEIGHT),))
    anchor = root.find("xdr:twoCellAnchor", NS)
    start, end = _anchor_cells(anchor)
    half, quarter = round(0.5 * ACTUAL * EMU), round(0.25 * ACTUAL * EMU)
    assert start == (1, half, 2, quarter)
    assert end == (3, half, 3, quarter)
    x, y, cx, cy, _, _ = _xfrm(anchor.find("xdr:sp", NS))
    cell = round(ACTUAL * EMU)
    assert (x, y) == (cell + half, 2 * cell + quarter)
    assert (cx, cy) == (3 * cell + half - x, 3 * cell + quarter - y)


def test_anchor_on_the_right_and_bottom_edge_of_the_sheet():
    root = _drawing((Rect(0, 0, 10 * WIDTH, 10 * HEIGHT),))
    start, end = _anchor_cells(root.find("xdr:twoCellAnchor", NS))
    assert start == (0, 0, 0, 0)
    assert end == (9, round(ACTUAL * EMU), 9, round(ACTUAL * EMU))


@pytest.mark.parametrize(("line", "flips"), [
    (Line(0, 0, 2 * WIDTH, HEIGHT), (None, None)),
    (Line(2 * WIDTH, 0, 0, HEIGHT), ("1", None)),
    (Line(0, HEIGHT, 2 * WIDTH, 0), (None, "1")),
    (Line(2 * WIDTH, HEIGHT, 0, 0), ("1", "1")),
])
def test_lines_are_connectors_with_flips_for_their_direction(line, flips):
    connector = _drawing((line,)).find("xdr:twoCellAnchor/xdr:cxnSp", NS)
    assert connector.find(".//a:prstGeom", NS).get("prst") == "line"
    x, y, cx, cy, flip_h, flip_v = _xfrm(connector)
    assert (x, y, cx, cy) == (0, 0, round(2 * ACTUAL * EMU), round(ACTUAL * EMU))
    assert (flip_h, flip_v) == flips


def test_line_width_color_dash_and_arrowheads():
    line = Line(0, 0, 50, 0, stroke_width_pt=1.5, color=(1, 0, 0), dash=(3, 1.5),
                start_arrow=Arrowhead("oval", "sm", "lg"), end_arrow=Arrowhead("triangle"))
    ln = _drawing((line,)).find(".//a:ln", NS)
    assert ln.get("w") == str(round(1.5 * EMU)) and ln.get("cap") == "flat"
    assert ln.find("a:solidFill/a:srgbClr", NS).get("val") == "FF0000"
    dashes = [(item.get("d"), item.get("sp")) for item in ln.findall("a:custDash/a:ds", NS)]
    assert dashes == [("200000", "100000")]
    assert ln.find("a:round", NS) is not None
    head, tail = ln.find("a:headEnd", NS), ln.find("a:tailEnd", NS)
    assert (head.get("type"), head.get("w"), head.get("len")) == ("oval", "sm", "lg")
    assert (tail.get("type"), tail.get("w"), tail.get("len")) == ("triangle", "med", "med")


def test_odd_dash_pattern_repeats_like_pdf():
    ln = _drawing((Line(0, 0, 50, 0, stroke_width_pt=1, dash=(2, 1, 3)),)).find(".//a:ln", NS)
    assert [(item.get("d"), item.get("sp")) for item in ln.findall("a:custDash/a:ds", NS)] == [
        ("200000", "100000"), ("300000", "200000"), ("100000", "300000")]


def test_polyline_is_a_freeform_path_in_local_emu():
    shape = _drawing((Polyline(((WIDTH, HEIGHT), (3 * WIDTH, HEIGHT), (2 * WIDTH, 3 * HEIGHT)), closed=True,
                               fill=(0, 0, 1)),)).find("xdr:twoCellAnchor/xdr:sp", NS)
    geometry = shape.find(".//a:custGeom", NS)
    path = geometry.find("a:pathLst/a:path", NS)
    assert (int(path.get("w")), int(path.get("h"))) == (round(2 * ACTUAL * EMU), round(2 * ACTUAL * EMU))
    points = [(int(pt.get("x")), int(pt.get("y"))) for pt in path.iter(f"{{{NS['a']}}}pt")]
    assert points == [(0, 0), (round(2 * ACTUAL * EMU), 0), (round(ACTUAL * EMU), round(2 * ACTUAL * EMU))]
    assert [child.tag.split("}")[1] for child in path] == ["moveTo", "lnTo", "lnTo", "close"]
    assert shape.find("xdr:spPr/a:solidFill/a:srgbClr", NS).get("val") == "0000FF"


def test_open_polyline_has_no_close_and_no_fill():
    shape = _drawing((Polyline(((0, 0), (WIDTH, HEIGHT), (2 * WIDTH, 0))),)).find(".//xdr:sp", NS)
    path = shape.find(".//a:path", NS)
    assert path.find("a:close", NS) is None and path.get("fill") == "none"
    assert shape.find("xdr:spPr/a:noFill", NS) is not None


def test_horizontal_polyline_has_a_nonzero_path_box():
    path = _drawing((Polyline(((0, HEIGHT), (2 * WIDTH, HEIGHT))),)).find(".//a:path", NS)
    assert int(path.get("h")) >= 1


@pytest.mark.parametrize(("shape", "preset"), [
    (Rect(0, 0, 30, 20), "rect"), (Rect(0, 0, 30, 20, corner_radius_pt=4), "roundRect"),
    (Ellipse(0, 0, 30, 20), "ellipse"), (Diamond(0, 0, 30, 20), "diamond"),
])
def test_preset_shapes(shape, preset):
    assert _drawing((shape,)).find(".//a:prstGeom", NS).get("prst") == preset


def test_corner_radius_becomes_the_adjust_value_of_the_shorter_side():
    # 4 pt on a 2 x 1 cell box: the shorter side is 15.75 pt in Excel.
    guide = _drawing((Rect(0, 0, 2 * WIDTH, HEIGHT, corner_radius_pt=4),)).find(".//a:prstGeom/a:avLst/a:gd", NS)
    assert guide.get("name") == "adj" and guide.get("fmla") == f"val {round(4 / ACTUAL * 100000)}"


def test_stroke_and_fill_of_area_shapes():
    shape = _drawing((Rect(0, 0, 30, 20, stroke=False, fill=(0.5, 0.5, 0.5)),)).find(".//xdr:sp", NS)
    assert shape.find("xdr:spPr/a:solidFill/a:srgbClr", NS).get("val") == "808080"
    assert shape.find("xdr:spPr/a:ln/a:noFill", NS) is not None
    stroked = _drawing((Ellipse(0, 0, 30, 20, stroke_width_pt=0.75, color=(0, 1, 0)),)).find(".//xdr:sp", NS)
    assert stroked.find("xdr:spPr/a:noFill", NS) is not None
    assert stroked.find("xdr:spPr/a:ln", NS).get("w") == str(round(0.75 * EMU))


def test_text_box_has_no_line_or_fill_and_writes_runs():
    frame = TextFrame(((Run("E"), Run("c", vert_align="subscript"), Run("2", Font("IPAexゴシック", 7, color=(1, 0, 0)),
                                                                        vert_align="superscript")),
                       (), (Run("次の行"),)),
                      font=MINCHO, h_align="center", v_align="bottom")
    shape = _drawing((TextBox(0, 0, 60, 40, frame),)).find(".//xdr:sp", NS)
    assert shape.find("xdr:nvSpPr/xdr:cNvSpPr", NS).get("txBox") == "1"
    assert shape.find("xdr:spPr/a:noFill", NS) is not None
    assert shape.find("xdr:spPr/a:ln/a:noFill", NS) is not None
    body = shape.find("xdr:txBody", NS)
    assert body.find("a:bodyPr", NS).get("anchor") == "b"
    assert body.find("a:bodyPr", NS).get("wrap") == "none"
    assert [body.find("a:bodyPr", NS).get(name) for name in ("lIns", "tIns", "rIns", "bIns")] == ["0"] * 4
    paragraphs = body.findall("a:p", NS)
    assert len(paragraphs) == 3
    assert paragraphs[0].find("a:pPr", NS).get("algn") == "ctr"
    runs = paragraphs[0].findall("a:r", NS)
    assert [run.find("a:t", NS).text for run in runs] == ["E", "c", "2"]
    properties = [run.find("a:rPr", NS) for run in runs]
    assert [prop.get("sz") for prop in properties] == ["1050", "1050", "700"]
    assert [prop.get("baseline") for prop in properties] == [None, "-25000", "30000"]
    assert {prop.find("a:latin", NS).get("typeface") for prop in properties[:2]} == {"IPAex明朝"}
    assert properties[2].find("a:ea", NS).get("typeface") == "IPAexゴシック"
    assert properties[2].find("a:solidFill/a:srgbClr", NS).get("val") == "FF0000"
    assert paragraphs[1].find("a:r", NS) is None and paragraphs[1].find("a:endParaRPr", NS) is not None


@pytest.mark.parametrize(("rotation", "vert", "rot"), [(0, "horz", None), (90, "vert270", None),
                                                        (270, "vert", None), (180, "horz", "10800000")])
def test_text_rotation_counterclockwise(rotation, vert, rot):
    frame = TextFrame(((Run("回転"),),), font=MINCHO, rotation=rotation)
    body = _drawing((TextBox(0, 0, 60, 40, frame),)).find(".//a:bodyPr", NS)
    assert (body.get("vert"), body.get("rot")) == (vert, rot)


def test_text_inside_preset_shapes():
    frame = TextFrame(((Run("判定"),),), font=MINCHO, h_align="center", v_align="center",
                      insets_pt=(1, 2, 3, 4), wrap=True)
    for shape in (Rect(0, 0, 60, 40, text=frame), Diamond(0, 0, 60, 40, text=frame),
                  Ellipse(0, 0, 60, 40, text=frame)):
        body = _drawing((shape,)).find(".//xdr:txBody/a:bodyPr", NS)
        assert body.get("anchor") == "ctr" and body.get("wrap") == "square"
        assert [body.get(name) for name in ("lIns", "tIns", "rIns", "bIns")] == [
            str(1 * EMU), str(2 * EMU), str(3 * EMU), str(4 * EMU)]


def test_group_contains_children_without_anchors_in_an_identity_child_space():
    group = Group((Rect(WIDTH, HEIGHT, WIDTH, HEIGHT), Group((Line(0, 0, WIDTH, 0), Ellipse(WIDTH, 0, WIDTH, HEIGHT)))))
    root = _drawing((group,))
    anchors = root.findall("xdr:twoCellAnchor", NS)
    assert len(anchors) == 1
    start, end = _anchor_cells(anchors[0])
    assert start == (0, 0, 0, 0) and end == (2, 0, 2, 0)
    outer = anchors[0].find("xdr:grpSp", NS)
    xfrm = outer.find("xdr:grpSpPr/a:xfrm", NS)
    values = {name: (int(xfrm.find(f"a:{name}", NS).get("x" if "ff" in name.lower() else "cx")),
                     int(xfrm.find(f"a:{name}", NS).get("y" if "ff" in name.lower() else "cy")))
              for name in ("off", "ext", "chOff", "chExt")}
    assert values["off"] == values["chOff"] == (0, 0)
    assert values["ext"] == values["chExt"] == (round(2 * ACTUAL * EMU), round(2 * ACTUAL * EMU))
    assert [child.tag.split("}")[1] for child in outer if child.tag.split("}")[1] != "nvGrpSpPr"] == [
        "grpSpPr", "sp", "grpSp"]
    inner = outer.find("xdr:grpSp", NS)
    assert [child.tag.split("}")[1] for child in inner][2:] == ["cxnSp", "sp"]


def test_shape_ids_are_unique_and_sequential_per_drawing():
    root = _drawing((Group((Rect(0, 0, 10, 10), Line(0, 0, 10, 10))), Ellipse(0, 0, 10, 10)))
    ids = [int(node.get("id")) for node in root.iter(f"{{{NS['xdr']}}}cNvPr")]
    assert ids == [2, 3, 4, 5]


def test_pictures_embed_deduplicated_media():
    png = _png()
    with _archive(_sheet((Picture(0, 0, 30, 15, png, "png"), Picture(40, 0, 30, 15, JPEG, "jpeg"),
                          Picture(0, 40, 30, 15, png, "png"))),
                  _sheet((Picture(0, 0, 30, 15, png, "png"),), name="T")) as archive:
        names = archive.namelist()
        assert [name for name in names if name.startswith("xl/media/")] == ["xl/media/image1.png",
                                                                            "xl/media/image2.jpeg"]
        assert archive.read("xl/media/image1.png") == png
        rels = archive.read("xl/drawings/_rels/drawing1.xml.rels").decode()
        assert rels.count("<Relationship ") == 2 and 'Target="../media/image2.jpeg"' in rels
        assert 'Target="../media/image1.png"' in archive.read("xl/drawings/_rels/drawing2.xml.rels").decode()
        types = archive.read("[Content_Types].xml").decode()
        assert 'Extension="png" ContentType="image/png"' in types
        assert 'Extension="jpeg" ContentType="image/jpeg"' in types
        root = ET.fromstring(archive.read("xl/drawings/drawing1.xml"))
        embeds = [blip.get(f"{{{NS['r']}}}embed") for blip in root.iter(f"{{{NS['a']}}}blip")]
        assert embeds == ["rId1", "rId2", "rId1"]


def test_workbook_with_drawings_still_loads_in_openpyxl():
    sheet = _sheet((Rect(0, 0, 30, 20), Picture(0, 30, 30, 15, _png(), "png"), Line(0, 0, 40, 40)))
    book = openpyxl.load_workbook(io.BytesIO(write_workbook(Workbook((sheet,)))))
    assert len(book["S"]._images) == 1


def test_same_drawings_give_identical_bytes():
    def book():
        return Workbook((_sheet((Group((Rect(0, 0, 30, 20, text=TextFrame(((Run("x"),),), font=MINCHO)),
                                        Picture(0, 30, 30, 15, _png(), "png"))),)),))
    assert write_workbook(book()) == write_workbook(book())


# --- validation -----------------------------------------------------------


@pytest.mark.parametrize("make", [
    lambda: Line(0, 0, 0, 0),
    lambda: Line(0, 0, 10, 0, stroke_width_pt=0),
    lambda: Line(0, 0, 10, 0, stroke_width_pt=1585),
    lambda: Line(0, 0, 10, 0, dash=None),
    lambda: Line(0, 0, 10, 0, dash=(1, -1)),
    lambda: Line(0, 0, 10, 0, start_arrow="triangle"),
    lambda: Line(float("nan"), 0, 10, 0),
    lambda: Arrowhead("barb"),
    lambda: Arrowhead("triangle", width="xl"),
    lambda: Polyline(((0, 0),)),
    lambda: Polyline(((0, 0), (1, 1)), fill=(1, 0, 0)),
    lambda: Polyline(((0, 0), (0, 0))),
    lambda: Polyline(((0, 0), (1,))),
    lambda: Rect(0, 0, 0, 10),
    lambda: Rect(0, 0, 10, 10, corner_radius_pt=5.1),
    lambda: Rect(0, 0, 10, 10, corner_radius_pt=-1),
    lambda: Rect(0, 0, 10, 10, stroke="yes"),
    lambda: Rect(0, 0, 10, 10, text="x"),
    lambda: Ellipse(0, 0, 10, -1),
    lambda: Diamond(0, 0, 10, 10, fill=(1, 1)),
    lambda: TextBox(0, 0, 10, 10, None),
    lambda: TextFrame((), font=MINCHO),
    lambda: TextFrame(((Run("x"),),), font=None),
    lambda: TextFrame(((Run("x"),),), font=MINCHO, rotation=45),
    lambda: TextFrame(((Run("x"),),), font=MINCHO, h_align="justify"),
    lambda: TextFrame(((Run("x"),),), font=MINCHO, insets_pt=(0, 0, 0)),
    lambda: TextFrame((("x",),), font=MINCHO),
    lambda: Picture(0, 0, 10, 10, b"", "png"),
    lambda: Picture(0, 0, 10, 10, JPEG, "png"),
    lambda: Picture(0, 0, 10, 10, _png(), "gif"),
    lambda: Group(()),
    lambda: Group(("x",)),
])
def test_drawings_reject_invalid_input(make):
    with pytest.raises(ValueError):
        make()


@pytest.mark.parametrize("drawing", [
    Rect(-1, 0, 10, 10), Rect(0, 0, 10 * WIDTH + 0.01, 10), Line(0, 0, 0, 10 * HEIGHT + 1),
    Group((Rect(0, 0, 10, 10), Ellipse(10 * WIDTH, 0, 1, 1))),
])
def test_drawings_must_lie_on_the_sheet(drawing):
    with pytest.raises(ValueError, match="outside the sheet"):
        _sheet((drawing,))


def test_drawings_must_be_drawing_items():
    with pytest.raises(ValueError, match="drawing"):
        _sheet(("rect",))


# --- review round 1 ----------------------------------------------------------


def test_solid_line_uses_an_empty_dash_like_the_scene():
    ln = _drawing((Line(0, 0, 50, 0, dash=()),)).find(".//a:ln", NS)
    assert ln.find("a:custDash", NS) is None


@pytest.mark.parametrize("make", [
    lambda: Line(0, 0, 10, 0, stroke_width_pt=1e-6),
    lambda: Line(0, 0, 10, 0, stroke_width_pt=0.25, dash=(6000, 1)),
    lambda: Line(0, 0, 10, 0, stroke_width_pt=1000, dash=(0.001, 1)),
    lambda: TextFrame(((Run("x"),),), font=MINCHO, insets_pt=(1e6, 0, 0, 0)),
    lambda: Rect(0, 0, 10, 10, stroke=False),
    lambda: Ellipse(0, 0, 10, 10, stroke=False),
    lambda: TextFrame(((Run("a\nb"),),), font=MINCHO),
    lambda: TextFrame(((Run("a\tb"),),), font=MINCHO),
    lambda: Picture(0, 0, 10, 10, b"\x89PNG\r\n\x1a\n", "png"),
    lambda: Picture(0, 0, 10, 10, _png()[:30], "png"),
    lambda: Picture(0, 0, 10, 10, b"\xff\xd8\xff" + _png(), "jpeg"),
])
def test_drawings_excel_would_change_or_drop_are_errors(make):
    with pytest.raises(ValueError):
        make()


def test_invisible_box_with_text_is_allowed():
    assert Rect(0, 0, 10, 10, stroke=False, text=TextFrame(((Run("x"),),), font=MINCHO))


@pytest.mark.parametrize("row", [3, 31, 47])
def test_coordinates_on_a_cell_boundary_anchor_at_the_start_of_the_next_cell(row):
    from gridlayoutpdf.renderers.xlsx._drawing import SheetGeometry
    geometry = SheetGeometry(Sheet("S", (WIDTH,) * 50, (HEIGHT,) * 50))
    for value in (row * HEIGHT, sum((HEIGHT,) * row)):
        assert geometry.y.anchor(value) == (row, 0)
    assert geometry.y.anchor(50 * HEIGHT) == (49, round(ACTUAL * EMU))


# --- review round 2 ----------------------------------------------------------


def test_jpeg_with_an_mpo_extension_is_a_jpeg():
    buffer = io.BytesIO()
    frames = [PILImage.new("RGB", (2, 1), color) for color in ((255, 0, 0), (0, 0, 255))]
    frames[0].save(buffer, "MPO", save_all=True, append_images=frames[1:])
    data = buffer.getvalue()
    with PILImage.open(io.BytesIO(data)) as image:
        assert image.format == "MPO"
    assert Picture(0, 0, 10, 10, data, "jpeg").data == data
    with pytest.raises(ValueError, match="not png"):
        Picture(0, 0, 10, 10, data, "png")


def test_drawing_text_is_not_spreadsheetml_escaped():
    frame = TextFrame(((Run("_x0041_"),),), font=MINCHO)
    run = _drawing((TextBox(0, 0, 60, 20, frame),)).find(".//a:t", NS)
    assert run.text == "_x0041_"


# --- review round 3 ----------------------------------------------------------


def test_truncated_png_is_an_error_even_when_pillow_loads_truncated_images(monkeypatch):
    from PIL import ImageFile

    monkeypatch.setattr(ImageFile, "LOAD_TRUNCATED_IMAGES", True)
    data = _png(8, 8)
    with pytest.raises(ValueError):
        Picture(0, 0, 10, 10, data[:-20], "png")
    assert Picture(0, 0, 10, 10, data, "png").data == data


# --- review round 4 ----------------------------------------------------------


@pytest.mark.parametrize(("width", "dash"), [(1.0, (5e-6, 1)), (0.01, (5e-8, 5e-8))])
def test_dash_lengths_of_half_a_unit_are_written_as_one(width, dash):
    # A ratio of exactly 0.5 (1/1000 %) rounds half up to 1, never to 0.
    ln = _drawing((Line(0, 0, 10, 0, stroke_width_pt=width, dash=dash),)).find(".//a:ln", NS)
    assert ln.find("a:custDash/a:ds", NS).get("d") == "1"


@pytest.mark.parametrize(("width", "dash"), [(1.0, (4.9999e-6, 1)), (2.5 / 12_700, (21_000 * 2.5 / 12_700, 1))])
def test_dash_lengths_outside_drawingml_values_are_errors(width, dash):
    # The second: a 2.5 EMU width is written as 2 EMU, so 21000 times it is above 2^31 - 1 thousandths of a %.
    with pytest.raises(ValueError, match="dash"):
        Line(0, 0, 10, 0, stroke_width_pt=width, dash=dash)


def test_dash_values_are_relative_to_the_written_line_width():
    ln = _drawing((Line(0, 0, 50, 0, stroke_width_pt=0.75, dash=(1.5, 0.75)),)).find(".//a:ln", NS)
    assert int(ln.get("w")) == 9525
    assert [(item.get("d"), item.get("sp")) for item in ln.findall("a:custDash/a:ds", NS)] == [("200000", "100000")]


def test_cut_jpeg_is_an_error_even_when_pillow_loads_truncated_images(monkeypatch):
    from PIL import ImageFile

    buffer = io.BytesIO()
    PILImage.effect_noise((64, 64), 64).convert("RGB").save(buffer, "JPEG")
    data = buffer.getvalue()
    monkeypatch.setattr(ImageFile, "LOAD_TRUNCATED_IMAGES", True)
    for cut in (data[:len(data) // 2], data[:-2]):
        with pytest.raises(ValueError):
            Picture(0, 0, 10, 10, cut, "jpeg")
    assert Picture(0, 0, 10, 10, data, "jpeg").data == data


# --- review round 5 ----------------------------------------------------------


def _noise_jpeg(**options):
    buffer = io.BytesIO()
    PILImage.effect_noise((64, 64), 64).convert("RGB").save(buffer, "JPEG", **options)
    return buffer.getvalue()


@pytest.mark.parametrize("trailer", [b"\n", b"\xff\xff", b"SEFH" + b"\x00" * 20, b"ftypmp42" * 8])
def test_jpeg_with_data_after_its_end_marker_is_accepted_unchanged(trailer):
    data = _noise_jpeg() + trailer
    assert Picture(0, 0, 10, 10, data, "jpeg").data == data


@pytest.mark.parametrize("options", [{}, {"progressive": True}, {"restart_marker_blocks": 1}])
def test_cut_jpeg_scans_are_errors_even_when_pillow_loads_truncated_images(monkeypatch, options):
    from PIL import ImageFile

    data = _noise_jpeg(**options)
    assert Picture(0, 0, 10, 10, data, "jpeg").data == data
    monkeypatch.setattr(ImageFile, "LOAD_TRUNCATED_IMAGES", True)
    with pytest.raises(ValueError, match="cut"):
        Picture(0, 0, 10, 10, data[:len(data) * 3 // 4], "jpeg")


def test_jpeg_cut_after_an_embedded_thumbnail_end_marker_is_an_error(monkeypatch):
    from PIL import ImageFile

    thumbnail = _noise_jpeg()
    main = _noise_jpeg()
    # An APP1 segment carrying a whole JPEG (with its own EOI), as an EXIF thumbnail does.
    segment = b"\xff\xe1" + (len(thumbnail) + 2).to_bytes(2, "big") + thumbnail
    data = main[:2] + segment + main[2:]
    assert Picture(0, 0, 10, 10, data, "jpeg").data == data
    monkeypatch.setattr(ImageFile, "LOAD_TRUNCATED_IMAGES", True)
    with pytest.raises(ValueError, match="cut"):
        Picture(0, 0, 10, 10, data[:len(data) - len(main) // 2], "jpeg")


def test_mpo_cut_after_its_first_frame_is_an_error(monkeypatch):
    from PIL import ImageFile

    buffer = io.BytesIO()
    frames = [PILImage.effect_noise((32, 32), 64).convert("RGB") for _ in range(2)]
    frames[0].save(buffer, "MPO", save_all=True, append_images=frames[1:])
    data = buffer.getvalue()
    from gridlayoutpdf.renderers.xlsx._model import _complete_jpeg_frames

    assert _complete_jpeg_frames(data) == 2
    monkeypatch.setattr(ImageFile, "LOAD_TRUNCATED_IMAGES", True)
    second = data.index(b"\xff\xd8", 2)
    with pytest.raises(ValueError, match="cut"):
        Picture(0, 0, 10, 10, data[:second + 100], "jpeg")


# --- review round 6 ----------------------------------------------------------


def test_mpo_with_bytes_between_its_frames_is_accepted():
    buffer = io.BytesIO()
    frames = [PILImage.effect_noise((32, 32), 64).convert("RGB") for _ in range(2)]
    frames[0].save(buffer, "MPO", save_all=True, append_images=frames[1:])
    data = buffer.getvalue()
    second = data.index(b"\xff\xd8\xff", 2)
    from gridlayoutpdf.renderers.xlsx._model import _complete_jpeg_frames

    padded = data[:second] + b"PAD!" + data[second:]
    assert _complete_jpeg_frames(padded) == 2
    assert Picture(0, 0, 10, 10, padded, "jpeg").data == padded


def test_jpeg_with_stray_bytes_between_segments_is_accepted_and_still_checked(monkeypatch):
    from PIL import ImageFile

    data = _noise_jpeg()
    app0_end = 4 + int.from_bytes(data[4:6], "big")  # SOI (2) + APP0 marker (2) + its length
    stray = data[:app0_end] + b"\x00\x00" + data[app0_end:]
    assert Picture(0, 0, 10, 10, stray, "jpeg").data == stray
    monkeypatch.setattr(ImageFile, "LOAD_TRUNCATED_IMAGES", True)
    with pytest.raises(ValueError, match="cut"):
        Picture(0, 0, 10, 10, stray[:len(stray) * 3 // 4], "jpeg")



def test_jpeg_scan_ends_at_a_second_soi_or_a_bogus_length_and_skips_ff00(monkeypatch):
    from PIL import ImageFile

    from gridlayoutpdf.renderers.xlsx._model import _jpeg_frame_end

    first, second = _noise_jpeg(), _noise_jpeg()
    monkeypatch.setattr(ImageFile, "LOAD_TRUNCATED_IMAGES", True)
    with pytest.raises(ValueError):
        Picture(0, 0, 10, 10, first[:len(first) // 3] + second, "jpeg")
    app0_end = 4 + int.from_bytes(first[4:6], "big")
    assert _jpeg_frame_end(first[:app0_end] + b"\xff\xfe\x00\x00" + first[app0_end:], 0) is None
    with_ff00 = first[:app0_end] + b"\xff\x00" + first[app0_end:]
    assert Picture(0, 0, 10, 10, with_ff00, "jpeg").data == with_ff00
