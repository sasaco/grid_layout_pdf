"""ZIP packaging of the standard-library .xlsx writer (ADR-0001)."""

import io
import zipfile

import openpyxl
import pytest

from gridlayoutpdf.renderers.xlsx._model import CellValue, Font, Sheet, Workbook
from gridlayoutpdf.renderers.xlsx._package import save_workbook, write_workbook


def _minimal(name="Sheet1", value="計算書"):
    return Workbook((
        Sheet(name, column_widths_pt=(16.08, 16.08), row_heights_pt=(15.48, 15.48),
              cells=(CellValue(0, 0, value, font=Font("IPAex明朝", 10.5)),)),
    ))


def _load(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data))


def test_minimal_workbook_is_read_back_by_openpyxl():
    book = _load(write_workbook(_minimal()))
    assert book.sheetnames == ["Sheet1"]
    sheet = book["Sheet1"]
    assert sheet["A1"].value == "計算書"
    assert sheet["A1"].font.name == "IPAex明朝"
    assert sheet["A1"].font.sz == 10.5


def test_same_model_gives_identical_bytes():
    assert write_workbook(_minimal()) == write_workbook(_minimal())


def test_zip_entries_have_fixed_dates_order_and_compression():
    data = write_workbook(_minimal())
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
        names = [info.filename for info in infos]
        assert names[0] == "[Content_Types].xml"
        assert names[:5] == [
            "[Content_Types].xml", "_rels/.rels", "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels", "xl/styles.xml",
        ]
        assert "xl/worksheets/sheet1.xml" in names
        assert len(names) == len(set(names))
        for info in infos:
            assert info.date_time == (1980, 1, 1, 0, 0, 0)
            assert info.compress_type == zipfile.ZIP_DEFLATED
            assert info.create_system == 0
            assert info.external_attr == 0o600 << 16
        assert archive.testzip() is None


def test_different_models_give_different_bytes():
    assert write_workbook(_minimal(value="a")) != write_workbook(_minimal(value="b"))


def test_save_workbook_replaces_destination_atomically(tmp_path):
    destination = tmp_path / "nested" / "report.xlsx"
    save_workbook(_minimal(), destination)
    assert destination.read_bytes() == write_workbook(_minimal())
    save_workbook(_minimal(value="更新"), destination)
    assert _load(destination.read_bytes())["Sheet1"]["A1"].value == "更新"
    assert [path.name for path in destination.parent.iterdir()] == ["report.xlsx"]


def test_save_workbook_leaves_destination_untouched_when_the_model_is_invalid(tmp_path):
    destination = tmp_path / "report.xlsx"
    destination.write_bytes(b"old")
    with pytest.raises(ValueError):
        save_workbook("not a workbook", destination)
    assert destination.read_bytes() == b"old"
    assert [path.name for path in tmp_path.iterdir()] == ["report.xlsx"]


def test_write_workbook_rejects_non_workbook():
    with pytest.raises(ValueError, match="Workbook"):
        write_workbook([])


def test_every_part_is_written_in_the_fixed_order():
    from test_xlsx_drawing import _png

    from gridlayoutpdf.renderers.xlsx._model import Picture, Rect

    sheet = Sheet("A", (16.08,) * 3, (15.48,) * 3, cells=(CellValue(0, 0, "x"),),
                  drawings=(Rect(0, 0, 10, 10), Picture(0, 20, 10, 10, _png(), "png")))
    book = Workbook((sheet, Sheet("B", (16.08,), (15.48,)),
                     Sheet("C", (16.08,), (15.48,), drawings=(Rect(0, 0, 5, 5),))))
    with zipfile.ZipFile(io.BytesIO(write_workbook(book))) as archive:
        assert archive.namelist() == [
            "[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels",
            "xl/styles.xml", "xl/sharedStrings.xml",
            "xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml", "xl/worksheets/sheet3.xml",
            "xl/worksheets/_rels/sheet1.xml.rels", "xl/worksheets/_rels/sheet3.xml.rels",
            "xl/drawings/drawing1.xml", "xl/drawings/_rels/drawing1.xml.rels", "xl/drawings/drawing2.xml",
            "xl/media/image1.png",
        ]


def test_failed_replace_keeps_the_destination_and_removes_the_temporary_file(tmp_path, monkeypatch):
    import os

    destination = tmp_path / "report.xlsx"
    destination.write_bytes(b"old")

    def fail(source, target):
        raise PermissionError("the file is open in Excel")

    monkeypatch.setattr(os, "replace", fail)
    with pytest.raises(PermissionError):
        save_workbook(_minimal(), destination)
    assert destination.read_bytes() == b"old"
    assert [path.name for path in tmp_path.iterdir()] == ["report.xlsx"]
