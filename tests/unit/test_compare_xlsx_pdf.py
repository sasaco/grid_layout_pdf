"""Step 10 second stage: the Excel-independent parts of scripts/compare_xlsx_pdf.py.

Text pieces of the library PDF are matched by their text (whitespace removed)
to Excel's PDF; equal texts take the nearest occurrence first, longer texts
are matched before shorter ones, and one printed glyph serves one piece.
Positions are compared at the anchor the workbook cell's alignment gives
(left, right or centre; top, bottom or middle), within one grid cell along the
reading direction and half a cell across it. The verdict is the exit code.
"""

import importlib.util
import json
import math
import sys
import tempfile
import zipfile
from pathlib import Path

import pymupdf
import pytest

from gridlayoutpdf import Cell, Document, PageBreak, Paragraph, Table, render_xlsx
from gridlayoutpdf.renderers.xlsx import XlsxStrictError

SCRIPT = Path(__file__).resolve().parents[2] / "scripts" / "compare_xlsx_pdf.py"
_SPEC = importlib.util.spec_from_file_location("compare_xlsx_pdf", SCRIPT)
compare = importlib.util.module_from_spec(_SPEC)
sys.modules["compare_xlsx_pdf"] = compare
_SPEC.loader.exec_module(compare)

Box, Piece, Glyph, CellArea = compare.Box, compare.Piece, compare.Glyph, compare.CellArea
PX, PY = 16.08, 15.48


def P(text, box, **kwargs):
    """A library piece; 10 pt unless given (the production path always knows the size)."""
    return Piece(text, box, **{"size": 10.0, **kwargs})


def glyphs(text, x, y, width=5.0, height=10.0, vertical=False, size=10.0):
    """One printed row of glyphs, each ``width`` wide, the first starting at ``x``."""
    return tuple(Glyph(char, Box(x + k * width, y, x + (k + 1) * width, y + height), vertical=vertical, size=size)
                 for k, char in enumerate(text))


def cell(x0, y0, x1, y1, horizontal="left", vertical="center", address="A1"):
    return CellArea(Box(x0, y0, x1, y1), horizontal, vertical, address)


def match(pieces, rows, cells=()):
    return compare.match_page(pieces, rows, cells, PX, PY)


# --- anchors and tolerances ------------------------------------------------------------------------

@pytest.mark.parametrize(("horizontal", "vertical", "expected"), [
    ("left", "center", (10, 25)), ("right", "center", (30, 25)), ("center", "center", (20, 25)),
    ("left", "top", (10, 20)), ("left", "bottom", (10, 30)), ("center", "bottom", (20, 30)),
])
def test_the_anchor_follows_the_cell_alignment(horizontal, vertical, expected):
    assert compare.anchor_point(Box(10, 20, 30, 30), horizontal, vertical) == expected


def test_the_tolerance_is_one_cell_along_the_reading_direction_and_half_a_cell_across():
    assert compare.tolerance(P("a", Box(0, 0, 1, 1)), PX, PY) == (PX, PY / 2)
    assert compare.tolerance(P("a", Box(0, 0, 1, 1), vertical=True), PX, PY) == (PX / 2, PY)


def test_find_cell_takes_the_cell_holding_the_printed_text_centre():
    left, right = cell(0, 0, 50, 20, address="A1"), cell(50, 0, 100, 20, "right", address="B1")
    assert compare.find_cell(Box(60, 5, 90, 15), (left, right)) is right
    assert compare.find_cell(Box(10, 5, 20, 15), (left, right)) is left
    assert compare.find_cell(Box(10, 30, 20, 40), (left, right)) is None


def test_the_cell_is_the_one_excel_printed_the_text_in():
    # The map puts a piece in its baseline's row; a table text centred on two grid rows has its box
    # centre in the row above. Excel prints it centred in the baseline row, which gives the cell.
    piece = P("項目", Box(53.88, 188.379, 73.88, 198.379))
    area = cell(50.88, 193.9, 147.36, 209.38, "left", address="A10")
    result = match([piece], [glyphs("項目", 53.42, 196.904, width=10.02, height=9.96)], [area])
    (item,) = result.matches
    assert item.anchor == ("left", "center") and item.cell == "A10"
    assert item.dy == pytest.approx(201.884 - 193.379) and not item.within  # the middles of the two boxes


# --- matching ---------------------------------------------------------------------------------------

def test_identical_positions_match_with_no_offset():
    pieces = [P("12.50", Box(100, 50, 125, 60), figure=True), P("kN", Box(200, 50, 210, 60), figure=True)]
    result = match(pieces, [glyphs("12.50", 100, 50), glyphs("kN", 200, 50)])
    assert result.missing == ()
    assert [(item.piece.text, item.dx, item.dy) for item in result.matches] == [("12.50", 0, 0), ("kN", 0, 0)]
    assert result.outside == ()


def test_equal_texts_take_the_nearest_occurrence_first():
    pieces = [P("1.00", Box(100, 200, 120, 210)), P("1.00", Box(100, 100, 120, 110))]
    rows = [glyphs("1.00", 101, 101), glyphs("1.00", 103, 199)]
    result = match(pieces, rows)
    by_y = {item.piece.box.y0: item for item in result.matches}
    assert (by_y[100].dx, by_y[100].dy) == (1, 1)
    assert (by_y[200].dx, by_y[200].dy) == (3, -1)  # offsets are Excel minus library


def test_the_nearest_pair_is_assigned_before_a_farther_one():
    # The piece at x=100 is nearest to the glyphs at x=101; the piece at x=110 then takes x=130.
    pieces = [P("7", Box(110, 50, 115, 60)), P("7", Box(100, 50, 105, 60))]
    result = match(pieces, [glyphs("7", 101, 50), glyphs("7", 130, 50)])
    assert sorted((item.piece.box.x0, item.box.x0) for item in result.matches) == [(100, 101), (110, 130)]


def test_longer_texts_are_matched_before_their_substrings():
    # "3" (a superscript) lies nearer the "3" of "12.375" than the printed "3", which is farther away.
    pieces = [P("3", Box(116, 50, 120, 60)), P("12.375", Box(100, 50, 130, 60))]
    rows = [glyphs("12.375", 100, 50), glyphs("3", 135, 45)]
    result = match(pieces, rows)
    assert result.missing == ()
    assert {item.piece.text: item.box.x0 for item in result.matches} == {"12.375": 100, "3": 135}


def test_a_printed_glyph_serves_one_piece_only():
    pieces = [P("1", Box(100, 50, 105, 60)), P("1", Box(100, 50, 105, 60))]
    result = match(pieces, [glyphs("1", 100, 50)])
    assert len(result.matches) == 1 and [piece.text for piece in result.missing] == ["1"]


def test_a_piece_found_inside_a_longer_row_is_matched_on_its_glyphs():
    result = match([P("A=", Box(140, 50, 150, 60))], [glyphs("S=12.50A=", 105, 50)])
    (item,) = result.matches
    assert item.box == Box(140, 50, 150, 60)


def test_a_text_missing_from_excel_is_listed():
    result = match([P("断面積", Box(50, 50, 80, 60))], [glyphs("断面", 50, 50)])
    assert result.matches == () and [piece.text for piece in result.missing] == ["断面積"]


def test_offsets_up_to_one_cell_across_and_half_a_cell_down_are_within_the_tolerance():
    pieces = [P("a", Box(100, 50, 105, 60), figure=True), P("b", Box(200, 50, 205, 60), figure=True)]
    result = match(pieces, [glyphs("a", 100 + PX, 50 + PY / 2), glyphs("b", 200 - PX, 50 - PY / 2)])
    assert result.outside == () and len(result.matches) == 2


@pytest.mark.parametrize(("dx", "dy"), [(PX + 0.01, 0), (0, PY / 2 + 0.01), (-PX - 0.01, 0)])
def test_larger_offsets_are_outside_the_tolerance(dx, dy):
    result = match([P("a", Box(100, 50, 105, 60), figure=True)], [glyphs("a", 100 + dx, 50 + dy)])
    (item,) = result.outside
    assert item.dx == pytest.approx(dx) and item.dy == pytest.approx(dy)


def test_vertical_text_has_one_cell_along_y_and_half_a_cell_across_x():
    piece = P("回転", Box(100, 100, 110, 120), vertical=True, figure=True)
    within = match([piece], [glyphs("回転", 100 + PX / 2, 100 + PY, vertical=True)])
    assert within.outside == () and len(within.matches) == 1
    outside = match([piece], [glyphs("回転", 100 + PX / 2 + 0.01, 100, vertical=True)])
    assert len(outside.outside) == 1


def test_a_right_aligned_cell_is_compared_at_the_right_edge():
    # Excel's substituted face prints the number narrower; its right edge stays at the cell's right.
    piece = P("12.150", Box(100, 50, 133.3, 60))
    row = glyphs("12.150", 133.3 - 6 * 4, 50, width=4)
    right = match([piece], [row], [cell(80, 45, 134, 61, "right")])
    assert right.matches[0].dx == pytest.approx(0) and right.matches[0].anchor == ("right", "center")
    left = match([piece], [row], [cell(80, 45, 134, 61, "left")])
    assert left.matches[0].dx == pytest.approx(9.3)


def test_a_centred_cell_is_compared_at_the_centre():
    piece = P("12", Box(100, 50, 112, 60))
    result = match([piece], [glyphs("12", 102, 51, width=4)], [cell(90, 45, 122, 61, "center")])
    assert result.matches[0].dx == pytest.approx(0) and result.matches[0].dy == pytest.approx(1)
    assert result.matches[0].cell == "A1"


def test_a_piece_outside_every_cell_is_compared_at_its_nearest_anchor():
    # A figure's text (from its tags): no cell gives the alignment, so each axis takes the best of three.
    piece = P("ラベル", Box(100, 50, 130, 60), figure=True)
    result = match([piece], [glyphs("ラベル", 112, 50, width=6)])
    (item,) = result.matches
    assert item.anchor == ("nearest", "nearest") and item.cell is None
    assert item.dx == pytest.approx(0) and item.dy == pytest.approx(0)  # right edges 130 and 130


def test_the_page_summary_reports_counts_and_the_largest_offsets():
    pieces = [P("a", Box(100, 50, 105, 60), figure=True), P("b", Box(200, 50, 205, 60), figure=True),
              P("c", Box(0, 0, 5, 10), figure=True)]
    result = match(pieces, [glyphs("a", 103, 50), glyphs("b", 200, 48)])
    summary = compare.page_summary(1, "S", result)
    assert summary["pieces"] == 3 and summary["matched"] == 2
    assert summary["max_dx_pt"] == pytest.approx(3) and summary["max_dy_pt"] == pytest.approx(2)
    assert [item["text"] for item in summary["missing"]] == ["c"]
    assert summary["outside"] == []
    assert summary[compare.FR021] == {"pieces": 0, "matched": 0, "missing": 0, "outside": 0, "max_dx_pt": None,
                                      "max_dy_pt": None, "max_range_dy_pt": None}
    json.dumps(summary, ensure_ascii=False)


# --- FR-021's vertically merged equation parts (lead decision, 2026-10-20) --------------------------------

FR = compare.FR021
MERGED = (50.88, 132.0, 259.92, 162.96)  # a cell merged over two grid rows (15.48 pt each)


def test_a_merged_equation_part_is_judged_against_its_merged_range():
    # Excel centres "I" in rows 5-6; the PDF draws it on the line of the fraction bar, 8.4 pt higher.
    piece = P("I", Box(200.5, 134.2, 204.3, 144.2), category=FR)
    within = match([piece], [glyphs("I", 205.6, 142.66, width=5, height=9.96)], [cell(*MERGED, "right")])
    (item,) = within.matches
    assert item.dy == pytest.approx(8.44, abs=0.01) and item.range_dy == pytest.approx(0.16, abs=0.01)
    assert item.within
    summary = compare.page_summary(1, "S", within)
    assert summary["pieces"] == 0 and summary[FR]["pieces"] == 1 and summary[FR]["outside"] == 0
    assert summary[FR]["max_dy_pt"] == pytest.approx(8.44, abs=0.01)
    assert summary[FR]["max_range_dy_pt"] == pytest.approx(0.16, abs=0.01)


def test_a_merged_equation_part_takes_the_nearest_occurrence_not_the_best_centred_one():
    # Two equations' "=" in merged ranges: the one far below is centred better in its range, but the
    # piece belongs to the near one (the range offset judges, the physical offset assigns).
    piece = P("=", Box(208.3, 134.2, 215.1, 144.2), category=FR)
    near = glyphs("=", 218.45, 141.0, width=5, height=9.96)  # 1.5 pt above its range's middle
    far = glyphs("=", 218.45, 420.0, width=5, height=9.96)  # centred in its range
    cells = [cell(211.68, 132.0, 227.76, 162.96, "center", address="K53"),
             cell(211.68, 409.52, 227.76, 440.48, "center", address="K72")]
    (item,) = match([piece], [far, near], cells).matches
    assert item.cell == "K53" and item.range_dy == pytest.approx(-1.5)


def test_a_merged_equation_part_off_the_range_centre_by_more_than_half_a_cell_is_outside():
    piece = P("I", Box(200.5, 134.2, 204.3, 144.2), category=FR)
    result = match([piece], [glyphs("I", 205.6, 132.2, width=5, height=9.96)], [cell(*MERGED, "right")])
    (item,) = result.outside
    assert item.range_dy == pytest.approx(-10.3, abs=0.01)
    assert compare.page_summary(1, "S", result)["outside"][0]["category"] == FR


def test_a_merged_equation_part_still_has_the_horizontal_tolerance_of_one_cell():
    piece = P("I", Box(200.5, 134.2, 204.3, 144.2), category=FR)
    result = match([piece], [glyphs("I", 199.3 + PX + 0.2, 142.66, width=5, height=9.96)],
                   [cell(*MERGED, "right")])
    assert len(result.outside) == 1 and result.outside[0].range_dy == pytest.approx(0.16, abs=0.01)


def test_a_merged_equation_part_outside_every_cell_is_outside():
    piece = P("I", Box(200.5, 134.2, 204.3, 144.2), category=FR)
    result = match([piece], [glyphs("I", 200.5, 134.2)])
    assert len(result.outside) == 1 and result.outside[0].range_dy is None


def test_a_merged_equation_part_in_a_one_row_cell_takes_the_ordinary_rule():
    # No vertical merge, so nothing excuses a vertical offset (stricter than the range rule).
    piece = P("I", Box(200.5, 134.2, 204.3, 144.2), category=FR)
    one_row = cell(50.88, 132.0, 211.68, 147.48, "right")
    result = match([piece], [glyphs("I", 199.3, 136.75, width=5, height=9.96)], [one_row])
    (item,) = result.matches
    assert item.within and item.range_dy is None
    shifted = match([piece], [glyphs("I", 199.3, 134.2 + PY / 2 + 0.5, width=5, height=9.96)], [one_row])
    assert len(shifted.outside) == 1


def _equation_document():
    from gridlayoutpdf import Equation, EquationAnchors, EquationGroup, Fraction

    return Document(
        EquationGroup(Equation("I", Fraction("4.50", "12"), label="二次", result="0.375", unit="m"),
                      anchors=EquationAnchors(4, 10, 16, 17)),
        EquationGroup(Equation("A", "4.50", label="面積", result="12.150", unit="m"),
                      anchors=EquationAnchors(4, 10, 16, 17)))


def test_pieces_are_classified_by_the_tags_of_the_commands_that_drew_them(tmp_path):
    from gridlayoutpdf import render_pdf

    pages = _equation_document().layout()
    path = tmp_path / "eq.pdf"
    render_pdf(pages, path)
    with pymupdf.open(path) as document:
        pieces = compare.pdf_pieces(document[0], pages[0])
    categories = {piece.text: piece.category for piece in pieces}
    # The line with a fraction: its label, lhs, relations, result and unit; not its numerator/denominator.
    assert {text: categories[text] for text in ("二次", "I", "0.375")} == dict.fromkeys(("二次", "I", "0.375"), FR)
    assert categories["4.50"] == categories["12"] == "standard"
    # The line without a fraction stays ordinary.
    assert {categories[text] for text in ("面積", "A", "12.150")} == {"standard"}
    assert [piece.category for piece in pieces if piece.text == "m"] == [FR, "standard"]
    assert [piece.category for piece in pieces if piece.text == "="] == [FR, FR, "standard", "standard"]


def test_without_the_scene_page_every_piece_is_ordinary(tmp_path):
    from gridlayoutpdf import render_pdf

    pages = _equation_document().layout()
    render_pdf(pages, tmp_path / "eq.pdf")
    with pymupdf.open(tmp_path / "eq.pdf") as document:
        assert {piece.category for piece in compare.pdf_pieces(document[0])} == {"standard"}


# --- reading the PDFs ---------------------------------------------------------------------------------

def _pdf(path, items, size=(300, 200)):
    with pymupdf.open() as document:
        page = document.new_page(width=size[0], height=size[1])
        for point, text, rotate in items:
            page.insert_text(point, text, fontsize=10, rotate=rotate)
        document.save(path)
    return path


def test_library_pieces_are_spans_without_whitespace(tmp_path):
    path = _pdf(tmp_path / "a.pdf", [((20, 40), "4.50 x 2.70", 0), ((20, 100), "Vert", 90)])
    with pymupdf.open(path) as document:
        pieces = compare.pdf_pieces(document[0])
    assert [(piece.text, piece.vertical) for piece in pieces] == [("4.50x2.70", False), ("Vert", True)]
    horizontal = pieces[0]
    assert horizontal.box.x0 == pytest.approx(20, abs=0.5) and horizontal.box.y1 == pytest.approx(40, abs=4)


def test_excel_glyph_rows_join_one_baseline_in_reading_order(tmp_path):
    path = _pdf(tmp_path / "b.pdf", [((120, 40), "kN", 0), ((20, 40), "S =", 0), ((20, 80), "next", 0),
                                     ((200, 150), "up", 90)])
    with pymupdf.open(path) as document:
        rows = compare.pdf_glyph_rows(document[0])
    texts = sorted("".join(glyph.char for glyph in row) for row in rows)
    assert texts == ["S=kN", "next", "up"]


def test_a_script_off_the_baseline_is_a_row_of_its_own(tmp_path):
    # The library's superscripts and subscripts are spans of their own, matched in Excel's script rows.
    path = tmp_path / "c.pdf"
    with pymupdf.open() as document:
        page = document.new_page(width=300, height=200)
        page.insert_text((20, 40), "m", fontsize=10)
        page.insert_text((29, 36.5), "2", fontsize=6.5)  # raised
        page.insert_text((40, 40), "kN", fontsize=10)
        document.save(path)
    with pymupdf.open(path) as document:
        rows = compare.pdf_glyph_rows(document[0])
    assert sorted("".join(glyph.char for glyph in row) for row in rows) == ["2", "mkN"]


# --- the workbook's cells -----------------------------------------------------------------------------

def test_workbook_cells_are_placed_on_their_pages_grid(tmp_path):
    document = Document(Paragraph("一頁目"), PageBreak(),
                        Table([4, 4], [[Cell("荷重"), Cell("12.150")]], numeric_columns=(1,)))
    pages = document.layout()
    result = render_xlsx([("S", pages)], tmp_path / "book.xlsx", strict=True)
    cells = compare.workbook_cells(result, pages)
    assert sorted(cells) == [0, 1]
    first, = [area for area in cells[0]]
    grid = pages[0].profile.grid
    assert first.box.y0 == pytest.approx(grid.y(0)) and first.horizontal == "left"
    number = next(area for area in cells[1] if area.horizontal == "right")
    # The right-aligned number ends at the table's right edge; its cell starts where its ink starts.
    assert number.box.x1 == pytest.approx(grid.x(8)) and grid.x(4) <= number.box.x0 < grid.x(8)
    assert grid.y(0) <= number.box.y0 < number.box.y1 <= grid.y(grid.rows)
    rows = (number.box.y1 - number.box.y0) / grid.pitch_y  # whole grid rows (the map may merge a row band)
    assert rows >= 1 and rows == pytest.approx(round(rows))


# --- review round 1: stricter matching ------------------------------------------------------------------

def test_a_merged_equation_part_far_from_its_library_line_is_outside():
    # Excel's text centred in a merged range 18 rows below the library's line: the range rule alone
    # would pass it (both values are Excel's); the library piece must lie on the range's rows.
    piece = P("=", Box(208.3, 134.2, 215.1, 144.2), category=FR, lookup=(211.7, 141.9))
    far = cell(211.68, 410.64, 227.76, 441.6, "center", address="K72")
    result = match([piece], [glyphs("=", 217.2, 421.1, width=5, height=9.96)], [far])
    (item,) = result.outside
    assert item.rule == "merged_range" and "library_off_range" in item.failures


def test_an_excel_text_outside_every_cell_takes_the_cell_of_the_library_piece():
    # The library piece's baseline is in a left-aligned cell; Excel's narrower text, shifted 25 pt right,
    # left every cell. Taking the nearest anchor would hide the shift.
    piece = P("abcd", Box(100, 50, 140, 60), lookup=(120, 57.9))
    area = cell(90, 45, 141, 61, "left", address="B3")
    result = match([piece], [glyphs("abcd", 140, 50, width=5)], [area])  # Excel's text centre at 150
    (item,) = result.matches
    assert item.anchor == ("left", "center") and item.cell == "B3" and item.dx == pytest.approx(40)
    assert not item.within


def test_a_header_or_footer_piece_is_compared_at_its_decoration_alignment():
    # A right-aligned footer whose left edge happens to match but whose right edge is 30 pt off.
    piece = P("3/5", Box(500, 800, 520, 810), align="right")
    result = match([piece], [glyphs("3/5", 500, 800, width=16.67)])
    (item,) = result.outside
    assert item.anchor == ("right", "center") and item.dx == pytest.approx(30.01, abs=0.01)


def test_text_is_matched_only_to_glyphs_read_in_the_same_direction():
    piece = P("回", Box(100, 100, 110, 110), vertical=True)
    across = (Glyph("回", Box(100, 100, 110, 110), size=10.0),)  # printed upright: the turn was lost
    assert match([piece], [across]).missing == (piece,)
    turned = (Glyph("回", Box(100, 100, 110, 110), size=10.0, vertical=True),)
    assert len(match([piece], [turned]).matches) == 1


def test_a_script_printed_at_the_base_size_is_outside():
    # The library's superscript 2 (6.5 pt); Excel lost the superscript and printed it at 9.96 pt.
    piece = P("2", Box(495, 99.5, 499, 106), size=6.5, figure=True)
    printed = (Glyph("2", Box(495, 100, 500, 110), size=9.96),)
    (item,) = match([piece], [printed]).outside
    assert "size" in item.failures and item.size_ratio == pytest.approx(9.96 / 6.5)
    kept = (Glyph("2", Box(495, 100, 499, 106.6), size=6.6),)
    assert match([piece], [kept]).outside == ()


def test_unused_excel_glyphs_are_reported_as_extra_text():
    pieces = [P("12.50", Box(100, 50, 125, 60))]
    result = match(pieces, [glyphs("12.50", 100, 50), glyphs("&P", 300, 800)])
    assert result.extra == ("&P",)
    assert compare.page_summary(1, "S", result)["extra"] == ["&P"]


def test_vertical_pieces_outside_the_paper_axes_tolerance_are_counted_apart():
    # Within one cell along the (vertical) reading direction, but more than half a cell down on paper.
    piece = P("回転", Box(100, 100, 110, 120), vertical=True, figure=True)
    printed = tuple(Glyph(char, Box(100, 110 + 10 * k, 110, 120 + 10 * k), size=10.0, vertical=True)
                    for k, char in enumerate("回転"))
    result = match([piece], [printed])
    assert result.outside == ()
    assert compare.page_summary(1, "S", result)["outside_if_paper_axes"] == 1


@pytest.mark.parametrize("make", [
    lambda: Box(10, 0, 5, 1), lambda: Box(0, 0, math.inf, 1), lambda: P("a b", Box(0, 0, 1, 1)),
    lambda: P("", Box(0, 0, 1, 1)), lambda: P("a", Box(0, 0, 1, 1), size=0),
    lambda: P("a", Box(0, 0, 1, 1), align="justify"), lambda: Glyph("ab", Box(0, 0, 1, 1)),
    lambda: CellArea(Box(0, 0, 1, 1), "general", "center", "A1"),
    lambda: CellArea(Box(0, 0, 1, 1), "left", "middle", "A1"),
], ids=["box-order", "box-infinite", "piece-space", "piece-empty", "piece-size", "piece-align", "glyph-two",
        "cell-horizontal", "cell-vertical"])
def test_the_records_reject_invalid_values_when_made(make):
    with pytest.raises(ValueError):
        make()


def test_workbook_cells_reject_a_workbook_whose_row_breaks_do_not_give_its_pages(tmp_path):
    import dataclasses

    pages = Document(Paragraph("一頁目"), PageBreak(), Paragraph("二頁目")).layout()
    result = render_xlsx([("S", pages)], tmp_path / "book.xlsx", strict=True)
    wrong = dataclasses.replace(result, sheets=(dataclasses.replace(result.sheets[0], last_page=3),))
    with pytest.raises(ValueError, match="row breaks"):
        compare.workbook_cells(wrong, pages)


# --- the verdict and the report -------------------------------------------------------------------------

def test_the_verdict_passes_only_when_every_check_passes():
    assert compare.verdict({"opens_without_repair": True, "page_count": True}) == compare.EXIT_PASSED
    assert compare.verdict({"opens_without_repair": True, "page_count": False}) == compare.EXIT_FAILED
    assert compare.verdict({}) == compare.EXIT_FAILED  # no check run is no pass


def _target(*blocks):
    document = Document(*(blocks or (Paragraph("断面の諸元"), Table([4, 4], [[Cell("荷重"), Cell("12.150")]],
                                                                   numeric_columns=(1,)),
                                     PageBreak(), Paragraph("二頁目"))))
    return compare.Target("small", "S", lambda: (document.layout(), None))


def _sheet_pages(workbook):
    """The pages of each sheet, from the workbook's row breaks."""
    import openpyxl

    return [len(sheet.row_breaks.brk) + 1 for sheet in openpyxl.load_workbook(workbook).worksheets]


def _copying_excel(*, open_error=None, hashes=(), drop_last_page=False, offset=(0.0, 0.0), page_size=None,
                   garbage=False, numbers_delta=0, shapes_delta=0, drop_key=None, extra_text=None, no_pdf=False,
                   sheet_field=None, clip=None, hidden_text=None):
    """A stand-in for Excel: each sheet's PDF is the library PDF's pages of that sheet (optionally moved)."""

    def run(workbook, sheet_pdfs, timeout):
        workbook = Path(workbook)
        report = {"version": "test", "build": "0", "printer": "Test Printer on Ne00:", "open_error": open_error,
                  "repair_logs": [], "sheets": [], "temp": "C:/Temp"}
        if open_error is None:
            numbers, shapes = compare.workbook_number_cells(workbook), compare.workbook_shapes(workbook)
            first = 0
            with pymupdf.open(workbook.with_name("library.pdf")) as source:
                for index, (pdf, count) in enumerate(zip(sheet_pdfs, _sheet_pages(workbook))):
                    last = first + count - (1 if drop_last_page and index == len(sheet_pdfs) - 1 else 0)
                    if garbage:
                        Path(pdf).write_bytes(b"not a PDF")
                    elif not no_pdf:
                        with pymupdf.open() as part:
                            for number in range(first, last):
                                rect = source[number].rect
                                width, height = page_size or (rect.width, rect.height)
                                page = part.new_page(width=width, height=height)
                                if clip is not None:  # the page as a form XObject shown through a clip
                                    shown = pymupdf.Rect(*clip)
                                    page.show_pdf_page(shown, source, number, clip=shown)
                                else:
                                    page.show_pdf_page(rect + (*offset, *offset), source, number)
                                if extra_text and number == first:
                                    page.insert_text((300, 700), extra_text, fontsize=10)
                                if hidden_text and number == first:  # drawn inside an empty clip
                                    page.insert_text((300, 720), hidden_text, fontsize=10)
                                    streams = page.get_contents()
                                    assert len(streams) > 1, "the hidden text needs a content stream of its own"
                                    data = part.xref_stream(streams[-1])
                                    part.update_stream(streams[-1], b"q 0 0 0 0 re W n\n" + data + b"\nQ\n")
                            part.save(pdf)
                    first += count
                    name = list(numbers)[index]
                    report["sheets"].append({"name": name, "number_cells": numbers[name] + numbers_delta,
                                             "shapes": shapes[name] + shapes_delta, "hashes": list(hashes),
                                             "lookup_error": None})
                    if sheet_field is not None:
                        report["sheets"][-1].update(sheet_field)
        if drop_key is not None:
            report.pop(drop_key, None)
        return report

    return run


def test_a_faithful_excel_print_passes(tmp_path):
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(), zoom=0.5)
    assert code == compare.EXIT_PASSED, report["checks"]
    assert report["passed"] is True and report["exit_code"] == 0
    assert report["checks"] == {"opens_without_repair": True, "readback": True, "strict": True, "no_hashes": True,
                                "page_count": True, "page_size": True, "positions": True, "no_extra_text": True}
    assert report["excel"]["temp"] == "C:/Temp"
    assert report["excel"]["printer"] == "Test Printer on Ne00:" and report["excel"]["version"] == "test"
    assert report["page_count"] == {"library": 2, "excel": 2, "sheets": [{"name": "S", "expected": 2, "excel": 2}]}
    assert [page["missing"] for page in report["pages"]] == [[], []]
    images = sorted(path.name for path in (tmp_path / "side-by-side").iterdir())
    assert images == ["page-01.png", "page-02.png"]
    saved = json.loads((tmp_path / "result.json").read_text(encoding="utf-8"))
    assert saved == report


def test_pages_of_later_sheets_are_compared_with_their_sheets_pdf_pages(tmp_path):
    from gridlayoutpdf import A3_LANDSCAPE_REFERENCE, PageSection

    target = _target(Paragraph("一頁目"), PageSection(A3_LANDSCAPE_REFERENCE, Paragraph("横の頁")),
                     Paragraph("三頁目"))
    code, report = compare.run(target, tmp_path, excel=_copying_excel(), zoom=0.5)
    assert code == compare.EXIT_PASSED, report["checks"]
    assert [(page["sheet"], page["excel_page"]) for page in report["pages"]] == [("S", 1), ("S_2", 1), ("S_3", 1)]


def test_a_print_moved_by_more_than_a_cell_fails_the_positions(tmp_path):
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(offset=(PX + 1, 0)), zoom=0.5)
    assert code == compare.EXIT_FAILED and report["checks"]["positions"] is False
    assert report["pages"][0]["outside"] and report["summary"]["outside"] > 0


def test_a_print_on_another_paper_size_fails(tmp_path):
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(page_size=(612, 792)), zoom=0.5)
    assert code == compare.EXIT_FAILED and report["checks"]["page_size"] is False


def test_a_report_without_text_does_not_pass_the_positions(tmp_path):
    from gridlayoutpdf import Spacer

    code, report = compare.run(_target(Spacer(2)), tmp_path, excel=_copying_excel(), zoom=0.5)
    assert code == compare.EXIT_FAILED and report["checks"]["positions"] is False


def test_a_repair_or_an_unopened_workbook_fails(tmp_path):
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(open_error="file is corrupt"), zoom=0.5)
    assert code == compare.EXIT_FAILED
    assert report["checks"]["opens_without_repair"] is False
    assert report["excel"]["open_error"] == "file is corrupt"
    # The checks that need the open workbook were not run: null, and the checks are not complete.
    assert {name for name, value in report["checks"].items() if value is None} == {
        "readback", "no_hashes", "page_count", "page_size", "positions", "no_extra_text"}
    assert report["checks_complete"] is False


def test_a_repair_log_fails(tmp_path):
    def excel(*args):
        report = _copying_excel()(*args)
        report["repair_logs"] = ["C:/Temp/error123_01.xml"]
        return report

    code, report = compare.run(_target(), tmp_path, excel=excel, zoom=0.5)
    assert code == compare.EXIT_FAILED and report["checks"]["opens_without_repair"] is False


def test_shapes_excel_dropped_fail_the_readback(tmp_path):
    from gridlayoutpdf import Figure, Label

    target = _target(Figure(4, 3, Label(10, 10, "図")), Paragraph("本文"))
    code, report = compare.run(target, tmp_path, excel=_copying_excel(shapes_delta=-1), zoom=0.5)
    assert code == compare.EXIT_FAILED and report["checks"]["readback"] is False
    assert report["shapes"] == {"S": {"workbook": 1, "excel": 0}}


def test_a_missing_page_fails_the_page_count(tmp_path):
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(drop_last_page=True), zoom=0.5)
    assert code == compare.EXIT_FAILED
    assert report["checks"]["page_count"] is False
    assert report["page_count"]["excel"] == 1
    assert report["pages"][1]["excel_page"] is None and report["pages"][1]["matched"] == 0


def test_hashes_in_a_number_cell_fail(tmp_path):
    hashes = [{"address": "E2", "text": "###", "value": 12.15}]
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(hashes=hashes), zoom=0.5)
    assert code == compare.EXIT_FAILED and report["checks"]["no_hashes"] is False
    assert report["hash_cells"] == [{"sheet": "S", "address": "E2", "text": "###", "value": 12.15}]


def test_excel_counting_other_number_cells_than_the_workbook_fails(tmp_path):
    # SpecialCells failing would count 0 numbers and look like "no ###".
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(numbers_delta=-1), zoom=0.5)
    assert code == compare.EXIT_FAILED and report["checks"]["no_hashes"] is False


def test_shifts_written_without_strict_still_fail(tmp_path, monkeypatch):
    import dataclasses

    from gridlayoutpdf.renderers.xlsx import XlsxShift

    original = compare.render_xlsx

    def shifting(*args, **kwargs):
        result = original(*args, **kwargs)
        return dataclasses.replace(result, shifts=(XlsxShift("S", 1, "x", (0, 0), (0, 1), "overlap"),))

    monkeypatch.setattr(compare, "render_xlsx", shifting)
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(), zoom=0.5, strict=False)
    assert code == compare.EXIT_FAILED and report["checks"]["strict"] is False
    assert report["strict"] is False and report["xlsx"]["shifts"] == 1


def test_a_strict_error_is_reported_with_its_own_exit_code(tmp_path):
    from gridlayoutpdf.renderers.xlsx import XlsxShift

    def broken():
        raise XlsxStrictError("1 shift(s)", (XlsxShift("S", 1, "x", (0, 0), (0, 1), "overlap"),))

    code, report = compare.run(compare.Target("broken", "S", broken), tmp_path, excel=_copying_excel())
    assert code == compare.EXIT_LIBRARY_ERROR
    assert report["passed"] is False and report["library_error"]["type"] == "XlsxStrictError"
    assert "1 shift(s)" in report["library_error"]["message"]
    assert report["library_error"]["shift_records"] == [
        {"sheet": "S", "page": 1, "text": "x", "from_cell": [0, 0], "to_cell": [0, 1], "reason": "overlap"}]
    saved = json.loads((tmp_path / "result.json").read_text(encoding="utf-8"))
    assert saved == report and saved["exit_code"] == 4


def test_an_excel_failure_has_its_own_exit_code(tmp_path):
    def excel(*args):
        raise compare.ExcelError("the Excel COM check did not finish in 1 s")

    code, report = compare.run(_target(), tmp_path, excel=excel, zoom=0.5)
    assert code == compare.EXIT_EXCEL_ERROR and "did not finish" in report["excel_error"]


@pytest.mark.parametrize("key", ["version", "build", "printer", "open_error", "repair_logs", "sheets", "temp"])
def test_an_excel_report_missing_a_field_is_an_excel_error(tmp_path, key):
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(drop_key=key), zoom=0.5)
    assert code == compare.EXIT_EXCEL_ERROR and key in report["excel_error"]


def test_an_unexpected_error_after_excel_is_an_internal_error_with_a_result(tmp_path):
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(garbage=True), zoom=0.5)
    assert code == compare.EXIT_INTERNAL_ERROR and report["passed"] is False
    assert report["internal_error"]["type"]
    assert json.loads((tmp_path / "result.json").read_text(encoding="utf-8"))["exit_code"] == compare.EXIT_INTERNAL_ERROR


def test_an_error_of_the_comparison_itself_is_not_reported_as_the_librarys(tmp_path, monkeypatch):
    def unknown(*args):
        raise ValueError("the workbook has a horizontal alignment the comparison does not know: 'fill'")

    monkeypatch.setattr(compare, "workbook_cells", unknown)
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(), zoom=0.5)
    assert code == compare.EXIT_INTERNAL_ERROR and "alignment" in report["internal_error"]["message"]


def test_a_run_removes_the_files_of_an_earlier_run(tmp_path):
    (tmp_path / "excel").mkdir()
    (tmp_path / "excel" / "sheet-03.pdf").write_bytes(b"old")
    (tmp_path / "side-by-side").mkdir()
    (tmp_path / "side-by-side" / "page-09.png").write_bytes(b"old")
    compare.run(_target(), tmp_path, excel=_copying_excel(), zoom=0.5)
    assert not (tmp_path / "excel" / "sheet-03.pdf").exists()
    assert not (tmp_path / "side-by-side" / "page-09.png").exists()


def test_without_excel_the_script_exits_with_the_reason(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(compare, "excel_missing_reason", lambda: "Microsoft Excel is not installed")
    (tmp_path / "library.pdf").write_bytes(b"old")
    code = compare.main(["--example", "representative", "--out", str(tmp_path)])
    assert code == compare.EXIT_NO_EXCEL
    assert "Microsoft Excel is not installed" in capsys.readouterr().err
    saved = json.loads((tmp_path / "result.json").read_text(encoding="utf-8"))
    assert saved["exit_code"] == compare.EXIT_NO_EXCEL and saved["reason"] == "Microsoft Excel is not installed"
    assert saved["checks_complete"] is False
    assert not (tmp_path / "library.pdf").exists()  # no earlier output beside the reason


@pytest.mark.parametrize("argv", [
    ["--example", "cantilever", "--variant", "long"], ["--example", "cantilever", "--font", "MS Mincho"],
    ["--example", "representative", "--pages", "1"], ["--example", "cantilever", "--pages", "0"],
    ["--example", "cantilever", "--pages", "1,1"], ["--example", "cantilever", "--pages", "9,3"],
    ["--example", "cantilever", "--pages", "38"], ["--example", "representative", "--font", ""],
    ["--example", "representative", "--zoom", "9"], ["--example", "representative", "--timeout", "inf"],
    ["--example", "representative", "--timeout", "0"],
], ids=lambda argv: " ".join(argv[1:]))
def test_bad_arguments_exit_with_the_usage_code(argv):
    with pytest.raises(SystemExit) as raised:
        compare.main(argv)
    assert raised.value.code == compare.EXIT_USAGE


def test_the_target_name_is_a_safe_folder_name():
    assert compare.representative_target("long", "MS Mincho/../x").name == "representative-long-MS_Mincho_.._x"


def test_the_default_output_folder_names_the_target():
    assert compare.default_output("representative-standard") == (
        compare.REPOSITORY_ROOT / "build" / "reports" / "xlsx-compare" / "representative-standard")


# --- review round 2 -------------------------------------------------------------------------------------

def test_library_pieces_know_their_size_and_the_point_on_their_baseline(tmp_path):
    # The cell of a piece is looked up on its baseline, nudged to its glyphs' side (as the map does).
    path = _pdf(tmp_path / "e.pdf", [((20, 40), "abc", 0), ((20, 150), "up", 90), ((200, 100), "down", 270)])
    with pymupdf.open(path) as document:
        pieces = {piece.text: piece for piece in compare.pdf_pieces(document[0])}
    flat, up, down = pieces["abc"], pieces["up"], pieces["down"]
    assert flat.size == pytest.approx(10) and flat.lookup == pytest.approx(((flat.box.x0 + flat.box.x1) / 2, 39.99))
    assert up.vertical and up.lookup == pytest.approx((19.99, (up.box.y0 + up.box.y1) / 2))
    assert down.vertical and down.lookup == pytest.approx((200.01, (down.box.y0 + down.box.y1) / 2))
    with pymupdf.open(path) as document:
        assert {glyph.size for row in compare.pdf_glyph_rows(document[0]) for glyph in row} == {10}


def test_an_unknown_size_is_outside():
    piece = Piece("a", Box(100, 50, 105, 60), figure=True)  # no size
    (item,) = match([piece], [glyphs("a", 100, 50)]).outside
    assert item.failures == ("size_unknown",)
    printed = (Glyph("a", Box(100, 50, 105, 60)),)  # Excel's glyph without a size
    (item,) = match([P("a", Box(100, 50, 105, 60), figure=True)], [printed]).outside
    assert item.failures == ("size_unknown",)


def test_an_ordinary_piece_in_no_cell_is_outside_unless_it_is_figure_text():
    plain = P("abc", Box(100, 50, 115, 60))
    (item,) = match([plain], [glyphs("abc", 100, 50)]).outside
    assert item.rule == "no_cell" and item.failures == ("no_cell",)
    figure = P("abc", Box(100, 50, 115, 60), figure=True)
    assert match([figure], [glyphs("abc", 100, 50)]).outside == ()


def test_a_merged_equation_part_does_not_fall_back_to_its_library_cell():
    # FR-021 pieces printed in no cell stay no_cell (lead decision): Excel's "=" 10 pt right of its range.
    piece = P("=", Box(208.3, 134.2, 215.1, 144.2), category=FR, lookup=(211.7, 141.9))
    merged = cell(200.0, 132.0, 220.0, 162.96, "center", address="K53")
    (item,) = match([piece], [glyphs("=", 222.0, 142.66, width=5, height=9.96)], [merged]).outside
    assert item.rule == "no_cell" and item.cell is None


@pytest.mark.parametrize(("y", "inside"), [(132.0 - PY / 2, True), (132.0 - PY / 2 - 0.01, False)])
def test_the_library_piece_may_lie_half_a_cell_off_its_merged_range(y, inside):
    piece = P("=", Box(208.3, y - 8, 215.1, y + 2), category=FR, lookup=(211.7, y))
    merged = cell(211.68, 132.0, 227.76, 162.96, "center", address="K53")
    (item,) = match([piece], [glyphs("=", 217.2, 142.5, width=5, height=9.96)], [merged]).matches
    assert item.library_in_range is inside and ("library_off_range" in item.failures) is not inside


def test_the_cell_of_excels_text_comes_before_the_library_pieces_cell():
    printed_in, lookup_in = cell(100, 45, 150, 61, "left", address="A1"), cell(150, 45, 200, 61, "right", address="B1")
    piece = P("ab", Box(160, 50, 170, 60), lookup=(165, 57.9))
    (item,) = match([piece], [glyphs("ab", 110, 50)], [lookup_in, printed_in]).matches
    assert item.cell == "A1"


def test_a_header_or_footer_takes_its_decoration_alignment_even_over_a_cell():
    piece = P("3/5", Box(500, 800, 520, 810), align="right")
    area = cell(480, 795, 560, 815, "left", address="Z60")
    (item,) = match([piece], [glyphs("3/5", 500, 800, width=10)], [area]).matches
    assert item.anchor == ("right", "center") and item.dx == pytest.approx(10)


def test_a_glyph_excel_clipped_is_outside():
    piece = P("ABCDE", Box(51.62, 50, 90.48, 60))
    row = glyphs("ABCD", 52.44, 50, width=7.8) + (Glyph("E", Box(83.64, 50, 90.48, 60), size=10.0, clipped=True),)
    (item,) = match([piece], [row], [cell(51.62, 45, 86.66, 61)]).outside
    assert item.failures == ("clipped",) and item.clipped_glyphs == ("E",)


def _clipped_pdf(path, texts, clip, size=(300, 200)):
    """A page whose text is drawn inside one clip rectangle (top-left coordinates), as Excel draws a cell."""
    with pymupdf.open() as document:
        page = document.new_page(width=size[0], height=size[1])
        for point, text, rotate in texts:
            page.insert_text(point, text, fontsize=10, rotate=rotate)
        page.clean_contents()
        xref = page.get_contents()[0]
        x0, y0, x1, y1 = clip
        prefix = f"q {x0} {size[1] - y1} {x1 - x0} {y1 - y0} re W n\n".encode()
        document.update_stream(xref, prefix + document.xref_stream(xref) + b"\nQ\n")
        document.save(path)
    return path


def _clipped_chars(path):
    with compare.ClipLoss(path) as loss, pymupdf.open(path) as document:
        rows = compare.pdf_glyph_rows(document[0], loss.page(0))
    return {glyph.char: glyph.clipped for row in rows for glyph in row}


def test_glyphs_whose_ink_excels_clip_cut_are_marked_along_the_line(tmp_path):
    # The clip ends inside "D": A-C are whole, D is cut, glyphs wholly outside are not even extracted.
    path = _clipped_pdf(tmp_path / "a.pdf", [((20, 40), "ABCDEFGH", 0)], (15, 20, 46, 60))
    chars = _clipped_chars(path)
    assert {char: chars[char] for char in "ABC"} == dict.fromkeys("ABC", False)
    assert chars["D"] is True


def test_glyphs_whose_ink_excels_clip_cut_are_marked_across_the_line(tmp_path):
    # A row too low for the text: the clip cuts the tops of every glyph (the line's cross direction).
    path = _clipped_pdf(tmp_path / "b.pdf", [((20, 40), "HHH", 0)], (15, 35, 200, 60))
    assert set(_clipped_chars(path).values()) == {True}


def test_turned_glyphs_cut_across_their_line_are_marked(tmp_path):
    # Text reading upward whose glyph tops (to the left) pass the clip: measured on p30-37 up to 1.53 pt.
    path = _clipped_pdf(tmp_path / "c.pdf", [((40, 150), "HH", 90)], (35, 20, 80, 190))
    assert set(_clipped_chars(path).values()) == {True}
    whole = _clipped_pdf(tmp_path / "d.pdf", [((40, 150), "HH", 90)], (20, 20, 80, 190))
    assert set(_clipped_chars(whole).values()) == {False}


def test_removing_clipping_leaves_strings_and_other_operators_alone():
    content = (b"q 0 0 10 10 re W n BT /F1 10 Tf (W n) Tj <574e> Tj [(a\\) W* n)] TJ ET Q q 1 1 2 2 re W* n Q "
               b"/Span <</ActualText (a>b W n)>> BDC EMC % W n\n")
    assert compare.without_clipping(content) == (
        b"q 0 0 10 10 re   n BT /F1 10 Tf (W n) Tj <574e> Tj [(a\\) W* n)] TJ ET Q q 1 1 2 2 re    n Q "
        b"/Span <</ActualText (a>b W n)>> BDC EMC % W n\n")


def test_an_inline_image_in_a_content_stream_is_refused():
    with pytest.raises(ValueError, match="inline image"):
        compare.without_clipping(b"q BI /W 1 /H 1 ID \x00\xff EI Q")


def test_pieces_of_turned_text_are_given_to_their_commands_both_ways(tmp_path):
    from gridlayoutpdf import A4_PORTRAIT_REFERENCE, GridSpec, QuarterTurn, render_pdf

    grid = A4_PORTRAIT_REFERENCE.grid
    local = GridSpec(0, 0, grid.pitch_y, grid.pitch_x, 4, 1)
    pages = Document(QuarterTurn(Paragraph("上へ"), turns=1, local_grid=local),
                     QuarterTurn(Paragraph("下へ"), turns=-1, local_grid=local)).layout()
    render_pdf(pages, tmp_path / "turned.pdf")
    with pymupdf.open(tmp_path / "turned.pdf") as pdf:
        pieces = {piece.text: piece for piece in compare.pdf_pieces(pdf[0], pages[0])}
    assert pieces["上へ"].vertical and pieces["下へ"].vertical
    assert {pieces[text].category for text in ("上へ", "下へ")} == {"standard"}


def test_figure_and_decoration_pieces_are_known_from_their_tags(tmp_path):
    from gridlayoutpdf import Figure, Label, PageDecoration, render_pdf

    pages = Document(Figure(4, 3, Label(10, 10, "図の字")), Paragraph("本文"),
                     footer=PageDecoration("頁 {page}", align="right")).layout()
    render_pdf(pages, tmp_path / "tags.pdf")
    with pymupdf.open(tmp_path / "tags.pdf") as pdf:
        pieces = {piece.text: piece for piece in compare.pdf_pieces(pdf[0], pages[0])}
    assert pieces["図の字"].figure and pieces["図の字"].align is None
    assert pieces["頁1"].align == "right" and not pieces["頁1"].figure
    assert not pieces["本文"].figure and pieces["本文"].align is None


def test_a_footer_in_no_cell_where_the_pdf_has_it_passes():
    footer = P("3/5", Box(500, 800, 520, 810), align="right")
    (item,) = match([footer], [glyphs("3/5", 500, 800, width=20 / 3)]).matches
    assert item.within and item.rule == "ordinary"


def test_figure_text_is_compared_at_its_nearest_anchors_even_over_a_cell():
    label = P("図", Box(100, 50, 110, 60), figure=True)
    area = cell(90, 45, 140, 61, "left", address="B3")
    (item,) = match([label], [glyphs("図", 104, 50, width=10)], [area]).matches
    assert item.anchor == ("nearest", "nearest") and item.cell is None


def test_only_figure_text_is_counted_as_compared_at_nearest_anchors():
    figure = P("a", Box(100, 50, 105, 60), figure=True)
    stray = P("b", Box(200, 50, 205, 60))  # in no cell and not figure text: no_cell
    summary = compare.page_summary(1, "S", match([figure, stray], [glyphs("a", 100, 50), glyphs("b", 200, 50)]))
    assert summary["nearest"] == 1 and summary["nearest_texts"] == ["a"]
    (outside,) = summary["outside"]
    assert outside["rule"] == "no_cell" and outside["dx_pt"] is None and outside["dy_pt"] is None


def _package(path, files):
    with zipfile.ZipFile(path, "w") as package:
        for name, data in files.items():
            package.writestr(name, data)
    return path


_RELS = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{}</Relationships>'
_REL = '<Relationship Id="{}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/{}" Target="{}"{}/>'
_ANCHOR = '<xdr:twoCellAnchor/>'


def test_workbook_shapes_follow_the_relationships_by_type(tmp_path):
    drawing = ('<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">'
               + _ANCHOR * 2 + '</xdr:wsDr>')
    path = _package(tmp_path / "book.xlsx", {
        # Another relationship first; the officeDocument one is found by its type.
        "_rels/.rels": _RELS.format(
            '<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/package/2006/relationships/'
            'metadata/core-properties" Target="docProps/core.xml"/>'
            + _REL.format("rId1", "officeDocument", "/xl/workbook.xml", "")),
        "xl/workbook.xml": ('<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                            '<sheets><sheet name="S" sheetId="1" r:id="rId1"/></sheets></workbook>'),
        "xl/_rels/workbook.xml.rels": _RELS.format(_REL.format("rId1", "worksheet", "worksheets/sheet1.xml", "")),
        "xl/worksheets/sheet1.xml": "<worksheet/>",
        "xl/worksheets/_rels/sheet1.xml.rels": _RELS.format(
            _REL.format("rId2", "vmlDrawing", "../drawings/vmlDrawing1.vml", "")  # not a drawing, comes first
            + _REL.format("rId1", "drawing", "/xl/drawings/drawing1.xml", "")  # an absolute target
            + _REL.format("rId3", "hyperlink", "../../../outside", ' TargetMode="External"')),  # never resolved
        "xl/drawings/drawing1.xml": drawing,
        "xl/drawings/vmlDrawing1.vml": "<xml/>",
    })
    assert compare.workbook_shapes(path) == {"S": 2}


def test_a_package_without_a_workbook_is_named_in_the_error(tmp_path):
    path = _package(tmp_path / "empty.xlsx", {"_rels/.rels": _RELS.format("")})
    with pytest.raises(ValueError, match="workbook"):
        compare.workbook_shapes(path)


@pytest.mark.parametrize("field", [{"shapes": True}, {"number_cells": False}])
def test_a_boolean_where_excel_reports_a_count_is_an_excel_error(tmp_path, field):
    code, _ = compare.run(_target(), tmp_path, excel=_copying_excel(sheet_field=field), zoom=0.5)
    assert code == compare.EXIT_EXCEL_ERROR


def test_excel_opening_the_workbook_without_sheets_is_an_excel_error():
    with pytest.raises(compare.ExcelError, match="no sheet"):
        compare.validate_excel_report({"version": "16.0", "build": "1", "printer": "P", "open_error": None,
                                       "repair_logs": [], "sheets": [], "temp": "C:/Temp"})


def test_every_result_says_whether_its_checks_are_complete(tmp_path):
    def broken():
        raise XlsxStrictError("1 shift(s)")

    _, library = compare.run(compare.Target("broken", "S", broken), tmp_path / "a", excel=_copying_excel())
    _, excel = compare.run(_target(), tmp_path / "b", excel=_copying_excel(drop_key="sheets"), zoom=0.5)
    _, passed = compare.run(_target(), tmp_path / "c", excel=_copying_excel(), zoom=0.5)
    assert (library["checks_complete"], excel["checks_complete"], passed["checks_complete"]) == (False, False, True)


def test_an_output_file_the_library_cannot_write_is_not_reported_as_the_librarys_error(tmp_path, monkeypatch):
    def locked(*args, **kwargs):
        raise PermissionError("workbook.xlsx is open in another program")

    monkeypatch.setattr(compare, "render_xlsx", locked)
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel())
    assert code == compare.EXIT_INTERNAL_ERROR and "PermissionError" == report["internal_error"]["type"]
    assert report["checks_complete"] is False


def test_pieces_from_one_fraction_lines_rhs_outside_the_fraction_stay_standard(tmp_path):
    from gridlayoutpdf import Equation, EquationAnchors, EquationGroup, Fraction, Operator, Row, render_pdf

    document = Document(EquationGroup(Equation("y", Row("7", Operator("×"), Fraction("1", "3"))),
                                      anchors=EquationAnchors(4, 10)))
    pages = document.layout()
    render_pdf(pages, tmp_path / "rhs.pdf")
    with pymupdf.open(tmp_path / "rhs.pdf") as pdf:
        categories = {piece.text: piece.category for piece in compare.pdf_pieces(pdf[0], pages[0])}
    assert categories["y"] == FR and categories["7×"] == "standard"  # PyMuPDF joins "7 ×" into one span


@pytest.mark.parametrize(("owners", "expected"), [
    ((FR, FR), FR), ((FR, "standard"), "standard"), ((FR, None), "standard"), ((), "standard")])
def test_a_piece_is_fr021_only_when_every_glyph_came_from_an_fr021_command(owners, expected):
    commands = [None if owner is None else compare.CommandInfo(0.0, 0.0, owner, None, False, 0, 1) for owner in owners]
    assert compare.piece_category(commands) == expected


def test_glyphs_of_a_page_not_drawn_by_its_scene_are_an_error(tmp_path):
    from gridlayoutpdf import render_pdf

    pages = Document(Paragraph("本文の段落")).layout()
    other = Document(Paragraph("別の段落です")).layout()
    render_pdf(pages, tmp_path / "a.pdf")
    with pymupdf.open(tmp_path / "a.pdf") as pdf, pytest.raises(ValueError, match="glyphs"):
        compare.pdf_pieces(pdf[0], other[0])


def test_a_faithful_print_of_a_fraction_line_fails_only_the_fr021_pieces(tmp_path):
    from gridlayoutpdf import Equation, EquationAnchors, EquationGroup, Fraction

    # Excel merges the line over its rows and centres it; a print where the parts stay on the library's
    # line is off the range's middle by about half a row: the FR021 pieces fail, the others pass.
    document = Document(EquationGroup(Equation("I", Fraction("4.50", "12"), label="二次", result="0.375",
                                               unit="m"), anchors=EquationAnchors(4, 10, 16, 17)))
    target = compare.Target("eq", "S", lambda: (document.layout(), None))
    code, report = compare.run(target, tmp_path, excel=_copying_excel(), zoom=0.5)
    assert code == compare.EXIT_FAILED and report["checks"]["positions"] is False
    assert report["summary"]["outside"] == 0 and report["summary"][FR]["outside"] >= 1
    page = report["pages"][0]
    assert page[FR]["outside"] == sum(item["category"] == FR for item in page["outside"]) >= 1
    assert page["outside_count"] == 0 and page[FR]["missing"] == 0
    assert all("range_dy" in item["failures"] or "no_cell" in item["failures"] for item in page["outside"])


def test_excel_text_no_piece_took_fails(tmp_path):
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(extra_text="ZZZ"), zoom=0.5)
    assert code == compare.EXIT_FAILED and report["checks"]["no_extra_text"] is False
    assert report["pages"][0]["extra"] == ["ZZZ"]


def test_the_maps_clipped_texts_and_dropped_shapes_are_reported(tmp_path, monkeypatch):
    import dataclasses

    from gridlayoutpdf.renderers.xlsx import XlsxFallback

    original = compare.render_xlsx

    def clipping(*args, **kwargs):
        result = original(*args, **kwargs)
        return dataclasses.replace(result, texts_clipped=(XlsxFallback("S", 1, "text_clipped", "A1: 長い"),),
                                   dropped=("page 1: a zero-length line",))

    monkeypatch.setattr(compare, "render_xlsx", clipping)
    _, report = compare.run(_target(), tmp_path, excel=_copying_excel(), zoom=0.5)
    assert report["xlsx"]["texts_clipped"] == [{"page": 1, "detail": "A1: 長い"}]
    assert report["xlsx"]["dropped"] == 1


def test_no_excel_page_to_compare_fails_the_page_size(tmp_path):
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(no_pdf=True), zoom=0.5)
    assert code == compare.EXIT_FAILED and report["checks"]["page_size"] is False


def test_a_page_with_text_commands_but_no_piece_fails_the_positions(tmp_path, monkeypatch):
    original = compare.pdf_pieces

    def blind(page, scene=None):
        return [] if page.number == 1 else original(page, scene)

    monkeypatch.setattr(compare, "pdf_pieces", blind)
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(), zoom=0.5)
    assert code == compare.EXIT_FAILED and report["checks"]["positions"] is False
    assert report["pages"][1]["text_commands"] > 0 and report["pages"][1]["pieces"] == 0


@pytest.mark.parametrize("field", [{"number_cells": "1"}, {"shapes": -1}, {"hashes": ["E2"]},
                                   {"hashes": [{"address": "E2"}]}])
def test_an_odd_value_in_excels_sheet_report_is_an_excel_error(tmp_path, field):
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(sheet_field=field), zoom=0.5)
    assert code == compare.EXIT_EXCEL_ERROR


def test_an_internal_error_marks_the_checks_incomplete(tmp_path):
    _, report = compare.run(_target(), tmp_path, excel=_copying_excel(garbage=True), zoom=0.5)
    assert report["checks_complete"] is False


def test_an_output_folder_that_cannot_be_written_is_an_internal_error(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(compare, "excel_missing_reason", lambda: None)
    blocked = tmp_path / "file"
    blocked.write_text("not a folder", encoding="utf-8")
    code = compare.main(["--example", "cantilever", "--pages", "1", "--out", str(blocked)])
    assert code == compare.EXIT_INTERNAL_ERROR
    assert "could not" in capsys.readouterr().err


@pytest.mark.parametrize("make", [
    lambda: compare.Match(P("a", Box(0, 0, 1, 1)), Box(0, 0, 1, 1), 0, 0, (1.0,), ("left", "center"), "A1"),
    lambda: compare.Match(P("a", Box(0, 0, 1, 1)), Box(0, 0, 1, 1), 0, 0, (1.0, 1.0), ("left", "center"), "A1",
                          "merged_range", 0.0, True),
    lambda: compare.Match(P("a", Box(0, 0, 1, 1)), Box(0, 0, 1, 1), 0, 0, (1.0, 1.0), ("west", "center"), "A1"),
    lambda: compare.Match(P("a", Box(0, 0, 1, 1)), Box(0, 0, 1, 1), 0, 0, (1.0, 1.0), ("left", "center"), "A1",
                          size_ratio=-1.0),
], ids=["tolerance-pair", "range-rule-for-standard", "anchor", "size-ratio"])
def test_a_match_rejects_inconsistent_values(make):
    with pytest.raises(ValueError):
        make()


# --- review round 4 -------------------------------------------------------------------------------------

def _ink_top(path, box):
    """The top of the ink inside ``box`` (top-left pt), measured at 64 px per pt."""
    with pymupdf.open(path) as document:
        pixmap = document[0].get_pixmap(matrix=pymupdf.Matrix(64, 64), clip=pymupdf.Rect(*box), alpha=False,
                                        colorspace=pymupdf.csGRAY)
    samples, stride = pixmap.samples, pixmap.stride  # one copy of the buffer
    for row in range(pixmap.height):
        if min(samples[row * stride:row * stride + pixmap.width]) < 128:
            return box[1] + row / 64
    raise AssertionError("no ink")


@pytest.mark.parametrize("phase", [0.0, 0.125, 0.0625])
def test_a_cut_of_a_third_of_a_point_is_found_at_any_pixel_phase(tmp_path, phase):
    whole = _clipped_pdf(tmp_path / "whole.pdf", [((20, 40), "HHH", 0)], (0, 0, 300, 200))
    top = _ink_top(whole, (15, 25, 60, 45))
    cut = _clipped_pdf(tmp_path / "cut.pdf", [((20, 40), "HHH", 0)], (15, top + 0.3 + phase, 200, 60))
    assert set(_clipped_chars(cut).values()) == {True}
    # A tenth of a point too (review round 5: missed at 4 px per pt, or with a level of 128 at 8 px per pt).
    shallow = _clipped_pdf(tmp_path / "shallow.pdf", [((20, 40), "HHH", 0)], (15, top + 0.1 + phase, 200, 60))
    assert set(_clipped_chars(shallow).values()) == {True}
    above = _clipped_pdf(tmp_path / "above.pdf", [((20, 40), "HHH", 0)], (15, top - 0.3 - phase, 200, 60))
    assert set(_clipped_chars(above).values()) == {False}


def _streams_pdf(path, clipped_text, free_text, clip, size=(300, 200)):
    """A page of three content streams: a clip, the clipped text, and the end of the clip; then free text."""
    with pymupdf.open() as document:
        page = document.new_page(width=size[0], height=size[1])
        page.insert_text(*clipped_text, fontsize=10)
        page.clean_contents()
        text = page.get_contents()[0]
        x0, y0, x1, y1 = clip
        streams = []
        for data in (f"q {x0} {size[1] - y1} {x1 - x0} {y1 - y0} re W n\n".encode(), b"\nQ\n"):
            xref = document.get_new_xref()
            document.update_object(xref, "<<>>")
            document.update_stream(xref, data)
            streams.append(xref)
        document.xref_set_key(page.xref, "Contents", f"[{streams[0]} 0 R {text} 0 R {streams[1]} 0 R]")
        if free_text is not None:
            page.insert_text(*free_text, fontsize=10)
        document.save(path)
    return path


def _glyph_boxes(path):
    with pymupdf.open(path) as document:
        return {char["c"]: Box.of(char["bbox"]) for block in document[0].get_text("rawdict")["blocks"]
                for line in block.get("lines", ()) for span in line["spans"] for char in span["chars"]}


def test_a_clip_in_its_own_content_stream_is_removed_too(tmp_path):
    boxes = _glyph_boxes(_clipped_pdf(tmp_path / "plain.pdf", [((20, 40), "ABCD", 0)], (0, 0, 300, 200)))
    edge = boxes["C"].x0 + 0.7 * (boxes["C"].x1 - boxes["C"].x0)  # inside C, right of its centre
    path = _streams_pdf(tmp_path / "streams.pdf", ((20, 40), "ABCD"), None, (15, 20, edge, 60))
    chars = _clipped_chars(path)
    assert (chars["A"], chars["B"], chars["C"]) == (False, False, True)


def test_glyphs_the_clip_dropped_are_listed_and_their_ink_is_not_blamed_on_a_neighbour(tmp_path):
    # A glyph Excel's clip takes so far that it is not extracted (in Excel, review round 4 measured one
    # whose centre was outside; here F-H lie wholly outside); its lost ink lies under the next cell's text,
    # which Excel did not cut.
    boxes = _glyph_boxes(_clipped_pdf(tmp_path / "plain.pdf", [((20, 40), "ABCDEFGH", 0)], (0, 0, 300, 200)))
    edge = boxes["F"].x0 - 0.2
    path = _streams_pdf(tmp_path / "next.pdf", ((20, 40), "ABCDEFGH"), ((boxes["G"].x0, 40), "x"),
                        (15, 20, edge, 60))
    with compare.ClipLoss(path) as loss, pymupdf.open(path) as document:
        clip_page = loss.page(0)
        rows = compare.pdf_glyph_rows(document[0], clip_page)
    chars = {glyph.char: glyph.clipped for row in rows for glyph in row}
    assert "F" not in chars and chars["x"] is False
    assert [char for char, _ in clip_page.dropped] == list("FGH")


def test_a_print_excel_clipped_fails_through_the_run(tmp_path):
    # The stand-in Excel shows the page through a clip that cuts the last glyph of the first paragraph.
    pages = _target().build()[0]
    from gridlayoutpdf import render_pdf

    render_pdf(pages, tmp_path / "probe.pdf")
    with pymupdf.open(tmp_path / "probe.pdf") as probe:
        last = next(Box.of(char["bbox"]) for block in probe[0].get_text("rawdict")["blocks"]
                    for line in block.get("lines", ()) for span in line["spans"] for char in span["chars"]
                    if char["c"] == "元")
    clip = (0, 0, last.x0 + 0.7 * (last.x1 - last.x0), 842)
    code, report = compare.run(_target(), tmp_path / "run", excel=_copying_excel(clip=clip), zoom=0.5)
    assert code == compare.EXIT_FAILED and report["checks"]["positions"] is False
    (item,) = [item for item in report["pages"][0]["outside"] if item["text"] == "断面の諸元"]
    assert item["failures"] == ["clipped"] and item["clipped_glyphs"] == ["元"]


def test_offsets_of_pieces_in_no_cell_are_not_the_page_largest():
    stray = P("b", Box(200, 50, 205, 60))  # no_cell: its nearest-anchor offset understates it
    summary = compare.page_summary(1, "S", match([stray], [glyphs("b", 203, 50)]))
    assert summary["max_dx_pt"] is None and summary["outside"][0]["rule"] == "no_cell"


# --- review round 5 -------------------------------------------------------------------------------------

def _first_paragraph_boxes():
    from gridlayoutpdf import render_pdf

    pages = _target().build()[0]
    with tempfile.TemporaryDirectory() as folder:
        render_pdf(pages, Path(folder) / "probe.pdf")
        with pymupdf.open(Path(folder) / "probe.pdf") as probe:
            return {char["c"]: Box.of(char["bbox"]) for block in probe[0].get_text("rawdict")["blocks"]
                    for line in block.get("lines", ()) for span in line["spans"] for char in span["chars"]
                    if 50 < char["bbox"][1] < 70}


def test_glyphs_the_clip_took_away_fail_through_the_run(tmp_path):
    boxes = _first_paragraph_boxes()
    clip = (0, 0, boxes["諸"].x0 - 0.2, 842)  # "諸元" wholly outside
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(clip=clip), zoom=0.5)
    page = report["pages"][0]  # the clip runs down the page, so the table's texts right of it go too
    chars = [item["char"] for item in page["clipped_out"]]
    assert {"諸", "元"} <= set(chars) and page["clipped_out_count"] == len(chars) > 2
    assert report["summary"]["clipped_out"] == page["clipped_out_count"] and report["checks"]["positions"] is False


def test_a_glyph_hidden_by_a_clip_alone_fails_the_positions(tmp_path):
    code, report = compare.run(_target(), tmp_path, excel=_copying_excel(hidden_text="Z"), zoom=0.5)
    page = report["pages"][0]
    assert [item["char"] for item in page["clipped_out"]] == ["Z"]
    assert report["summary"]["missing"] == 0 and report["summary"]["extra"] == 0
    assert code == compare.EXIT_FAILED and report["checks"]["positions"] is False


def _abcd_with_edge(tmp_path):
    boxes = _glyph_boxes(_clipped_pdf(tmp_path / "plain.pdf", [((20, 40), "ABCD", 0)], (0, 0, 300, 200)))
    return boxes["C"].x0 + 0.7 * (boxes["C"].x1 - boxes["C"].x0), boxes  # inside C; D wholly outside


def _judged(path):
    with compare.ClipLoss(path) as loss, pymupdf.open(path) as document:
        clip_page = loss.page(0)
        rows = compare.pdf_glyph_rows(document[0], clip_page)
    return {glyph.char: glyph.clipped for row in rows for glyph in row}, [char for char, _ in clip_page.dropped]


def test_a_clip_in_a_later_content_stream_is_removed(tmp_path):
    edge, _ = _abcd_with_edge(tmp_path)
    path = _streams_pdf(tmp_path / "later.pdf", ((20, 40), "ABCD"), None, (15, 20, edge, 60))
    with pymupdf.open(path) as document:  # an empty stream first, so the clip is in the second one
        page = document[0]
        empty = document.get_new_xref()
        document.update_object(empty, "<<>>")
        document.update_stream(empty, b" ")
        document.xref_set_key(page.xref, "Contents",
                              "[" + " ".join(f"{xref} 0 R" for xref in [empty, *page.get_contents()]) + "]")
        document.save(tmp_path / "later2.pdf")
    chars, dropped = _judged(tmp_path / "later2.pdf")
    assert (chars["A"], chars["B"], chars["C"], dropped) == (False, False, True, ["D"])


def _shown(source, target, nested=0):
    """``source``'s page shown as a form XObject (its clip inside the form), ``nested`` times more."""
    for level in range(nested + 1):
        with pymupdf.open(source) as original, pymupdf.open() as document:
            page = document.new_page(width=original[0].rect.width, height=original[0].rect.height)
            page.show_pdf_page(page.rect, original, 0)
            out = target.with_name(f"{target.stem}-{level}.pdf")
            document.save(out)
        source = out
    return source


@pytest.mark.parametrize("nested", [0, 1], ids=["form", "form-in-form"])
def test_a_clip_inside_a_form_xobject_is_removed(tmp_path, nested):
    edge, _ = _abcd_with_edge(tmp_path)
    clipped = _clipped_pdf(tmp_path / "inner.pdf", [((20, 40), "ABCD", 0)], (15, 20, edge, 60))
    chars, dropped = _judged(_shown(clipped, tmp_path / "shown.pdf", nested))
    assert (chars["A"], chars["B"], chars["C"], dropped) == (False, False, True, ["D"])


def test_repeated_glyphs_the_clip_took_away_are_the_later_ones(tmp_path):
    boxes = _glyph_boxes(_clipped_pdf(tmp_path / "plain.pdf", [((20, 40), "AB", 0)], (0, 0, 300, 200)))
    width = boxes["A"].x1 - boxes["A"].x0
    path = _clipped_pdf(tmp_path / "aaaa.pdf", [((20, 40), "AAAAAA", 0)], (15, 20, 20 + 3 * width - 0.1, 60))
    with compare.ClipLoss(path) as loss:
        dropped = loss.page(0).dropped
    assert len(dropped) >= 3 and all(box.x0 >= 20 + 3 * width - 0.5 for _, box in dropped)


def test_a_failed_clip_removal_closes_the_pdfs(tmp_path, monkeypatch):
    path = _clipped_pdf(tmp_path / "a.pdf", [((20, 40), "ABC", 0)], (0, 0, 300, 200))

    def broken(content):
        raise ValueError("cannot remove the clipping")

    monkeypatch.setattr(compare, "without_clipping", broken)
    with pytest.raises(ValueError, match="cannot remove"):
        compare.ClipLoss(path)
    path.unlink()  # Windows refuses while a document holds the file


def test_the_text_flags_leave_out_glyphs_wholly_outside_a_clip(tmp_path):
    # ClipPage.dropped relies on this bit: with TEXT_CLIP the glyph D outside the clip is not extracted.
    edge, boxes = _abcd_with_edge(tmp_path)
    path = _clipped_pdf(tmp_path / "flags.pdf", [((20, 40), "ABCD", 0)], (15, 20, boxes["D"].x0 - 0.2, 60))
    assert compare.TEXT_FLAGS & pymupdf.TEXT_CLIP
    with pymupdf.open(path) as document:
        kept = document[0].get_text("text", flags=compare.TEXT_FLAGS)
        everything = document[0].get_text("text", flags=compare.TEXT_FLAGS & ~pymupdf.TEXT_CLIP)
    assert "D" not in kept and "D" in everything


def test_a_glyph_box_of_no_area_is_an_error():
    with pytest.raises(ValueError, match="no area"):
        compare.ClipPage.clipped(None, Box(10, 10, 10, 20))


# --- review round 6 (LOW) -------------------------------------------------------------------------------

def test_a_check_that_was_not_run_does_not_pass():
    assert compare.verdict({"opens_without_repair": True, "positions": None}) == compare.EXIT_FAILED


def test_after_an_internal_error_the_checks_not_run_are_null(tmp_path, monkeypatch):
    def unknown(*args):
        raise ValueError("the workbook has a horizontal alignment the comparison does not know: 'fill'")

    monkeypatch.setattr(compare, "workbook_cells", unknown)
    _, report = compare.run(_target(), tmp_path, excel=_copying_excel(), zoom=0.5)
    assert report["checks_complete"] is False
    assert {name for name, value in report["checks"].items() if value is None} == {
        "readback", "no_hashes", "page_count", "page_size", "positions", "no_extra_text"}


def test_a_glyph_box_of_no_area_names_its_glyph_and_page(tmp_path, monkeypatch):
    path = _clipped_pdf(tmp_path / "a.pdf", [((20, 40), "AB", 0)], (0, 0, 300, 200))

    def flat(bbox):
        x0, y0, x1, y1 = bbox
        return Box(x0, y0, x0, y1)

    with compare.ClipLoss(path) as loss, pymupdf.open(path) as document:
        clip_page = loss.page(0)
        monkeypatch.setattr(compare.Box, "of", staticmethod(flat))
        with pytest.raises(ValueError, match=r"'A'.*page 1"):
            compare.pdf_glyph_rows(document[0], clip_page)


def test_a_glyph_box_past_the_page_edge_counts_as_clipped(tmp_path):
    path = _clipped_pdf(tmp_path / "edge.pdf", [((290, 40), "AB", 0)], (0, 0, 300, 200))  # B ends past x 300
    chars = _clipped_chars(path)
    assert chars["B"] is True
