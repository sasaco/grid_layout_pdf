"""FR-021: mapping tagged ScenePage sequences to a grid Excel workbook model (Step 8).

The pages are laid out with a font registry of fixed advances (0.6 em per
character) so cell ranges can be computed by hand. Cells are 0-based (row,
col) of the sheet; the A4 reference grid is 30 columns x 47 rows of
16.08 x 15.48 pt from (50.88, 54.598).
"""

from io import BytesIO
import math
from dataclasses import replace
from pathlib import Path

from PIL import Image as PILImage
import pytest

from gridlayoutpdf.blocks import Paragraph, Rule
from gridlayoutpdf.containers import Columns, Indent, SideRule
from gridlayoutpdf.document import Document, PageDecoration, PageSection
from gridlayoutpdf.equations import Equation, EquationAnchors, EquationGroup, MathBlock
from gridlayoutpdf.figures import Figure, Image, Label, TextBox
from gridlayoutpdf.fonts import FontRegistry, TextMetrics
from gridlayoutpdf.geometry import GridSpec, Rect
from gridlayoutpdf.math import Fraction, Operator, Row, Script, TermStop
from gridlayoutpdf.profiles import A3_LANDSCAPE_REFERENCE, A4_PORTRAIT_REFERENCE, PageProfile
from gridlayoutpdf.renderers.xlsx import XlsxError, XlsxResult, XlsxStrictError, render_xlsx
from gridlayoutpdf.renderers.xlsx import _model as xl
from gridlayoutpdf.renderers.xlsx._map import map_reports
from gridlayoutpdf.rotation import QuarterTurn
from gridlayoutpdf.scene import Line, RasterImage, Rectangle, ScenePage, Text
from gridlayoutpdf.semantics import Kind, Role, Scope, Tag
from gridlayoutpdf.styles import ReportStyle, TextStyle
from gridlayoutpdf.tables import Cell, Table


class FixedFonts(FontRegistry):
    def measure(self, text, style):
        return TextMetrics(len(text) * style.size * 0.6, style.size * 0.8, style.size * 0.2)


FONTS = FixedFonts()
STYLE = TextStyle(size=10)
REPORT_STYLE = ReportStyle(body=STYLE, heading=TextStyle(size=12))
A4 = A4_PORTRAIT_REFERENCE
OX, OY, PX, PY = 50.88, 54.598, 16.08, 15.48
PADDING = 2.4
CELL = 15.75  # the A4 grid's cells are 21 Excel pixels (printed at 16.08 x 15.48 pt)


def padding(size):
    return PADDING


def doc(*blocks, **options):
    options.setdefault("profile", A4)
    return Document(*blocks, fonts=FONTS, style=REPORT_STYLE, **options)


def mapped(*blocks, strict=False, name="計算書", **options):
    pages = doc(*blocks, **options).layout()
    return map_reports([(name, pages)], FONTS, strict=strict, cell_padding_pt=padding)


def first_sheet(result):
    workbook, _ = result
    return workbook.sheets[0]


def cell_at(sheet, row, col):
    return next(cell for cell in sheet.cells if (cell.row, cell.col) == (row, col))


def values(sheet):
    return {(cell.row, cell.col): cell.value for cell in sheet.cells if cell.value is not None}


def text_cell(sheet, value):
    """The cell showing ``value``; a number cell matches the text it shows as well."""
    return next(cell for cell in sheet.cells if _plain(cell.value) == value
                or (isinstance(cell.value, float) and isinstance(value, str) and _number(value) == cell.value))


def _number(text):
    try:
        return float(text.replace("−", "-"))
    except ValueError:
        return None


def _plain(value):
    if isinstance(value, xl.RichText):
        return "".join(run.text for run in value.runs)
    return value


def merge_of(sheet, row, col):
    return next((merge for merge in sheet.merges
                 if merge.first_row <= row <= merge.last_row and merge.first_col <= col <= merge.last_col), None)


def paragraph_tag(align="left", line=0):
    return Tag(Role.TEXT, (Scope(Kind.PARAGRAPH, index=line, data={"align": align, "heading": False}),))


def raw_page(*commands, profile=A4, metadata=None):
    return ScenePage(profile, commands, metadata or {})


def baseline(row):
    return OY + row * PY + 11


def x_of(col, offset=0.0):
    return OX + col * PX + offset


# 1. Text pieces, merges and alignment --------------------------------------------

def test_a_paragraph_line_is_one_cell_with_its_font_size_and_alignment():
    sheet = first_sheet(mapped(Paragraph("a")))
    cell = text_cell(sheet, "a")
    assert (cell.row, cell.col) == (0, 0)
    assert cell.font == xl.Font("IPAex明朝", 10)
    assert cell.h_align == "left" and cell.v_align == "center"
    assert merge_of(sheet, 0, 0) is None  # 6 pt (8 px) and its padding fit one 15.75 pt column


def test_a_piece_wider_than_one_column_merges_the_columns_it_spans():
    sheet = first_sheet(mapped(Paragraph("abcdef")))  # 36 pt: columns 0-2, room in Excel for 36 + 4.8 pt
    assert merge_of(sheet, 0, 0) == xl.CellRange(0, 0, 0, 2)
    assert values(sheet) == {(0, 0): "abcdef"}


@pytest.mark.parametrize(("align", "expected"), [("center", "center"), ("right", "right"), ("justify", "left")])
def test_paragraph_alignment_becomes_the_cell_alignment(align, expected):
    sheet = first_sheet(mapped(Paragraph("ab", align=align)))
    assert text_cell(sheet, "ab").h_align == expected


def test_pages_are_stacked_with_a_page_break_between_them():
    pages = doc(Paragraph("a"), Paragraph("b")).layout()
    second = ScenePage(pages[0].profile, pages[0].commands, {**pages[0].metadata, "page_number": 2})
    workbook, result = map_reports([("R", (pages[0], second))], FONTS, cell_padding_pt=padding)
    sheet = workbook.sheets[0]
    assert len(sheet.row_heights_pt) == 2 * 47 and len(sheet.column_widths_pt) == 30
    assert sheet.row_breaks == (47,)
    assert set(sheet.column_widths_pt) == set(sheet.row_heights_pt) == {CELL}  # 21 px, printed 16.08 x 15.48 pt
    assert (47, 0) in values(sheet) and (48, 0) in values(sheet)
    assert result.sheets[0].first_page == 1 and result.sheets[0].last_page == 2


def test_page_setup_keeps_the_grid_origin_and_prints_at_100_percent():
    sheet = first_sheet(mapped(Paragraph("a")))
    page = sheet.page
    assert (page.paper, page.orientation) == ("A4", "portrait")
    assert page.margins.left == pytest.approx(OX) and page.margins.top == pytest.approx(OY)
    assert page.scale == 100 and sheet.show_gridlines and not page.print_gridlines


@pytest.mark.parametrize(("profile", "right", "bottom"), [
    # A4: 30 x 16.08 x 1.006 + 4.5 pt wide, 47 x 15.48 x 1.006 + 4.5 pt tall (_units.full_scale_size_pt)
    (A4_PORTRAIT_REFERENCE, 54.6, 50.86),
    (A3_LANDSCAPE_REFERENCE, 51.34, 50.86),
])
def test_right_and_bottom_margins_shrink_only_to_the_measured_full_scale_size(profile, right, bottom):
    """Excel aligns its header and footer sections with the margins, so they shrink as little as 100 % allows."""
    from gridlayoutpdf.renderers.xlsx._map import DEFAULT_FONT
    from gridlayoutpdf.renderers.xlsx._units import print_fit

    sheet = first_sheet(mapped(Paragraph("a"), profile=profile))
    margins = sheet.page.margins
    assert (margins.right, margins.bottom) == (right, bottom)
    assert print_fit(sheet, DEFAULT_FONT).scale == sheet.page.scale == 100
    for wider in (replace(margins, right=right + 0.01), replace(margins, bottom=bottom + 0.01)):
        assert print_fit(replace(sheet, page=replace(sheet.page, margins=wider)), DEFAULT_FONT).scale < 100


def test_a_margin_on_an_exact_hundredth_still_prints_at_100_percent():
    # 29 columns of 18 px (13.8 pt printed): 29 x 13.8 x 1.006 + 4.5 = 407.1012 pt, so the margin
    # 1190.5512 - 50.88 - 407.1012 = 732.57 pt is a whole hundredth and the print area, computed back,
    # comes out 1 ulp short of the need
    grid = GridSpec(50.88, 54.598, 13.8, 15.48, 29, 47)
    page = raw_page(Text(x_of(0), baseline(0), "a", STYLE, tag=paragraph_tag()),
                    profile=PageProfile("A3 landscape", 1191, 842, grid))
    workbook, result = map_reports([("R", (page,))], FONTS, strict=True, cell_padding_pt=padding)
    sheet = workbook.sheets[0]
    assert sheet.column_widths_pt == (13.5,) * 29  # 18 px
    assert sheet.page.scale == 100 and result.fallbacks == ()
    assert sheet.page.margins.right == 732.56


# 2. Overlaps ------------------------------------------------------------------------

def overlapping_page():
    return raw_page(Text(x_of(2), baseline(3), "ab", STYLE, tag=paragraph_tag(line=0)),
                    Text(x_of(2, 8), baseline(3), "cd", STYLE, tag=paragraph_tag(line=1)))


def test_an_overlapping_piece_moves_right_and_the_move_is_recorded():
    workbook, result = map_reports([("R", (overlapping_page(),))], FONTS, cell_padding_pt=padding)
    assert values(workbook.sheets[0]) == {(3, 2): "ab", (3, 3): "cd"}
    (shift,) = result.shifts
    assert (shift.sheet, shift.page, shift.text) == ("R", 1, "cd")
    assert (shift.from_cell, shift.to_cell, shift.reason) == ((3, 2), (3, 3), "overlap")
    assert (shift.from_ref, shift.to_ref) == ("C4", "D4")


def test_strict_refuses_a_move(tmp_path):
    with pytest.raises(XlsxStrictError, match="1 shift") as caught:
        map_reports([("R", (overlapping_page(),))], FONTS, strict=True, cell_padding_pt=padding)
    assert len(caught.value.shifts) == 1
    target = tmp_path / "strict.xlsx"
    with pytest.raises(XlsxStrictError):
        render_xlsx([("R", (overlapping_page(),))], target, FONTS, strict=True, cell_padding_pt=padding)
    assert not target.exists()


def test_a_move_past_the_right_edge_is_always_an_error():
    page = raw_page(Text(x_of(29), baseline(0), "a", STYLE, tag=paragraph_tag(line=0)),
                    Text(x_of(29, 4), baseline(0), "b", STYLE, tag=paragraph_tag(line=1)))
    with pytest.raises(XlsxError, match="right edge"):
        map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)


def test_text_beyond_the_grid_is_an_error():
    page = raw_page(Text(x_of(29), baseline(0), "abcdef", STYLE, tag=paragraph_tag()))
    with pytest.raises(XlsxError, match="right edge"):
        map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)


# 3. Scripts --------------------------------------------------------------------------

@pytest.mark.parametrize(("node", "vert"), [(Script("x", superscript="2"), "superscript"),
                                            (Script("x", subscript="i"), "subscript")])
def test_a_script_is_rich_text_in_its_base_cell_at_the_base_size(node, vert):
    sheet = first_sheet(mapped(MathBlock(node)))
    (cell,) = [cell for cell in sheet.cells if cell.value is not None]
    assert isinstance(cell.value, xl.RichText)
    base, script = cell.value.runs
    assert base.text == "x" and base.vert_align is None
    assert script.vert_align == vert and (script.font or cell.font).size == 10


# 4. Numbers --------------------------------------------------------------------------

def test_a_number_piece_is_a_number_cell_with_the_pdf_display():
    sheet = first_sheet(mapped(Paragraph("−2.000")))
    cell = text_cell(sheet, -2.0)
    assert cell.number_format == '0.000;"−"0.000'


@pytest.mark.parametrize("text", ["0500", "−0.000", "1.5 m"])
def test_texts_that_are_not_plain_numbers_stay_text(text):
    sheet = first_sheet(mapped(Paragraph(text)))
    assert text_cell(sheet, text).number_format is None


def test_a_number_without_room_for_the_cell_padding_stays_text():
    # "12.5" is 4 glyphs of 8 px (24 pt) in a two-column table cell (31.5 pt in Excel, its borders keep it
    # from growing): 2 x 3.8 pt padding does not fit.
    pages = doc(Table([2], [[Cell("12.5")]])).layout()
    workbook, result = map_reports([("R", pages)], FONTS, cell_padding_pt=lambda size: 3.8)
    assert text_cell(workbook.sheets[0], "12.5").number_format is None
    assert [item.kind for item in result.numbers_as_text] == ["number_as_text"] and not result.fallbacks
    map_reports([("R", pages)], FONTS, strict=True, cell_padding_pt=lambda size: 3.8)  # not a strict failure
    workbook, result = map_reports([("R", pages)], FONTS, cell_padding_pt=lambda size: 3.7)
    assert text_cell(workbook.sheets[0], 12.5).number_format == "0.0" and not result.numbers_as_text


def test_numbers_in_figures_stay_text():
    workbook, _ = mapped(Figure(4, 2, Label(5, 12, "12.5")))
    sheet = workbook.sheets[0]
    assert not sheet.cells
    (group,) = sheet.drawings
    (box,) = group.children
    assert isinstance(box, xl.TextBox) and box.text.paragraphs[0][0].text == "12.5"


# 5. Fractions and equations ------------------------------------------------------------

def equation():
    return EquationGroup(
        Equation("I", Fraction(Row("4.50", Operator("×"), "2.70"), "12"), result="7.381", unit="m"),
        anchors=EquationAnchors(2, 12, 16, 17))


def test_numerator_and_denominator_are_cells_in_consecutive_rows():
    sheet = first_sheet(mapped(equation()))
    numerator, denominator = text_cell(sheet, "4.50×2.70"), text_cell(sheet, "12")
    assert denominator.row == numerator.row + 1
    assert numerator.h_align == denominator.h_align == "center"
    # both lie high in their rows (the numerator above the bar, the denominator right under it)
    assert numerator.v_align == denominator.v_align == "top"


def test_the_fraction_bar_is_a_shape_on_the_row_boundary_between_them():
    sheet = first_sheet(mapped(equation()))
    numerator = text_cell(sheet, "4.50×2.70")
    (bar,) = [drawing for drawing in sheet.drawings if isinstance(drawing, xl.Line)]
    assert bar.y1 == bar.y2 == pytest.approx(sum(sheet.row_heights_pt[:numerator.row + 1]))
    assert not sheet.borders


def test_lhs_relations_result_and_unit_are_separate_cells_merged_over_the_equation_rows():
    sheet = first_sheet(mapped(equation()))
    numerator = text_cell(sheet, "4.50×2.70")
    pieces = {value: text_cell(sheet, value) for value in ("I", "=", "7.381", "m")}
    assert len({(cell.row, cell.col) for cell in pieces.values()}) == 4
    for cell in pieces.values():
        merge = merge_of(sheet, cell.row, cell.col)
        assert (merge.first_row, merge.last_row) == (numerator.row, numerator.row + 1)
        assert cell.v_align == "center"
    assert pieces["I"].h_align == "right"


# 6. Borders and shapes -------------------------------------------------------------------

def test_table_edges_are_cell_borders_and_cell_alignment_is_kept():
    table = Table([2, 3], [[Cell("a"), Cell("1.0")]], numeric_columns=(1,))
    sheet = first_sheet(mapped(table))
    edges = {border.edge_key(): border.style for border in sheet.borders}
    rows = max(row for kind, row, _ in edges if kind == "horizontal")
    expected = {("horizontal", row, col) for row in (0, rows) for col in range(5)}
    expected |= {("vertical", row, col) for row in range(rows) for col in (0, 2, 5)}
    assert set(edges) == expected and set(edges.values()) == {"thin"}
    assert text_cell(sheet, 1.0).h_align == "right"
    assert not sheet.drawings


def test_a_user_line_on_a_cell_boundary_is_a_shape():
    figure = Figure(4, 2, Line(1, PY, 4 * PX - 1, PY, 0.5, (1, 0, 0)))
    sheet = first_sheet(mapped(figure))
    assert not sheet.borders
    (group,) = sheet.drawings
    (line,) = group.children
    assert isinstance(line, xl.Line) and line.color == (1, 0, 0)


def test_a_rule_is_the_top_border_of_its_row():
    sheet = first_sheet(mapped(Paragraph("a"), Rule("medium", columns=3)))
    assert {(border.row, border.col, border.side, border.style) for border in sheet.borders} == {
        (1, col, "top", "medium") for col in range(3)}


def test_a_side_rule_on_a_column_boundary_is_a_left_border():
    sheet = first_sheet(mapped(SideRule(Paragraph("ab"), PX)))
    assert {(border.row, border.col, border.side) for border in sheet.borders} == {(0, 1, "left")}


def test_a_side_rule_off_the_column_boundaries_is_a_shape_and_recorded():
    blocks = (SideRule(Paragraph("ab"), 6),)
    workbook, result = mapped(*blocks)
    assert not workbook.sheets[0].borders
    assert [fallback.kind for fallback in result.fallbacks] == ["border_off_grid"]
    with pytest.raises(XlsxStrictError, match="1 fallback"):
        mapped(*blocks, strict=True)


def test_a_border_inside_a_merged_piece_is_a_shape_and_recorded():
    border = Line(x_of(1), OY, x_of(1), OY + PY, 0.5,
                  tag=Tag(Role.TABLE_BORDER, (), {"style": "thin", "stroke": 0}))
    page = raw_page(Text(x_of(0), baseline(0), "abcd", STYLE, tag=paragraph_tag()), border)
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    sheet = workbook.sheets[0]
    assert not sheet.borders and isinstance(sheet.drawings[0], xl.Line)
    assert [fallback.kind for fallback in result.fallbacks] == ["border_in_merge"]


def test_shape_and_border_colors_are_kept():
    figure = Figure(4, 2, Rectangle(2, 2, 20, 10, 0.5, (0, 0, 1), fill=(1, 1, 0)))
    table = Table([2], [[Cell("a")]], border_color=(0, 0.5, 0))
    sheet = first_sheet(mapped(figure, table))
    (group,) = sheet.drawings
    (rect,) = group.children
    assert (rect.color, rect.fill) == ((0, 0, 1), (1, 1, 0))
    assert {border.color for border in sheet.borders} == {(0, 0.5, 0)}


# 7. Rotated tables -------------------------------------------------------------------------

def turned(local_grid):
    return QuarterTurn(Paragraph("abc"), turns=1, local_grid=local_grid, block_id="turned")


def test_a_quarter_turn_on_the_swapped_page_grid_becomes_turned_cells():
    sheet = first_sheet(mapped(turned(GridSpec(0, 0, PY, PX, 4, 2))))
    cell = text_cell(sheet, "abc")
    assert cell.rotation == 90
    merge = merge_of(sheet, cell.row, cell.col)
    assert merge.first_col == merge.last_col and merge.last_row - merge.first_row == 1


def test_a_quarter_turn_on_another_grid_is_a_shape_and_recorded():
    block = turned(GridSpec(0, 0, 10, 10, 6, 2))
    workbook, result = mapped(block)
    sheet = workbook.sheets[0]
    assert not sheet.cells
    (group,) = sheet.drawings
    (box,) = group.children
    assert isinstance(box, xl.TextBox) and box.text.rotation == 90
    assert [(fallback.kind, fallback.page) for fallback in result.fallbacks] == [("quarter_turn", 1)]
    with pytest.raises(XlsxStrictError):
        mapped(block, strict=True)


# 8. Figures and images ----------------------------------------------------------------------

def png_bytes(mode="RGB"):
    buffer = BytesIO()
    PILImage.new(mode, (4, 4), "red").save(buffer, "PNG")
    return buffer.getvalue()


def test_a_figure_is_one_group_of_shapes_with_text_inside_its_boxes():
    figure = Figure(8, 4, TextBox(2, 2, 60, 30, ("box",)), Label(80, 20, "lab"),
                    Line(1, 1, 10, 10), block_id="fig")
    sheet = first_sheet(mapped(figure))
    (group,) = sheet.drawings
    rect, label, line = group.children
    assert isinstance(rect, xl.Rect) and rect.text.paragraphs == ((xl.Run("box"),),)
    assert isinstance(label, xl.TextBox) and isinstance(line, xl.Line)


def test_an_image_is_a_picture(tmp_path):
    path = tmp_path / "a.png"
    path.write_bytes(png_bytes())
    sheet = first_sheet(mapped(Image(path, col_span=2, rows=2)))
    (picture,) = sheet.drawings
    assert isinstance(picture, xl.Picture) and picture.format == "png"
    # the 2 x 2 cell frame holds a 30.96 pt square image; Excel's cells keep it at the same fractions
    assert (picture.width, picture.height) == pytest.approx((30.96 / PX * CELL, 2 * CELL))


def test_a_raster_image_of_another_format_is_written_as_png():
    buffer = BytesIO()
    PILImage.new("RGB", (4, 4)).save(buffer, "BMP")
    page = raw_page(RasterImage(OX, OY, 10, 10, buffer.getvalue(), tag=Tag(Role.IMAGE)))
    workbook, _ = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    (picture,) = workbook.sheets[0].drawings
    assert picture.format == "png"


# 9. Headers and footers ------------------------------------------------------------------------

A4_BODY = PageProfile("A4 body", 595, 842, A4.grid, 2, 45)
A3_BODY = PageProfile("A3 body", 1191, 842, A3_LANDSCAPE_REFERENCE.grid, 2, 45)


def test_page_fields_become_the_page_code_and_the_report_total():
    blocks = [Paragraph("a"), PageSection(A3_BODY, Paragraph("b")), Paragraph("c")]
    workbook, result = mapped(*blocks, profile=A4_BODY, header=PageDecoration("{page}/{pages}", align="right"))
    first, second, third = workbook.sheets
    assert [sheet.name for sheet in workbook.sheets] == ["計算書", "計算書_2", "計算書_3"]
    assert first.header.right.parts == (xl.PageNumber(), "/3")
    assert first.header.right.font == xl.Font("IPAex明朝", 9)
    assert [sheet.page.first_page_number for sheet in workbook.sheets] == [1, 2, 3]
    assert second.page.paper == "A3" and second.page.orientation == "landscape"
    assert [(info.name, info.first_page, info.last_page) for info in result.sheets] == [
        ("計算書", 1, 1), ("計算書_2", 2, 2), ("計算書_3", 3, 3)]
    assert not any(cell.value == "1/3" for cell in first.cells)


def test_a_formatted_page_field_is_an_error():
    with pytest.raises(XlsxError, match="format"):
        mapped(Paragraph("a"), profile=A4_BODY, footer=PageDecoration("p.{page:03}"))


def test_a_decoration_size_that_excel_headers_cannot_show_is_an_error():
    with pytest.raises(XlsxError, match="whole pt"):
        mapped(Paragraph("a"), profile=A4_BODY, footer=PageDecoration("{page}", style=TextStyle(size=9.5)))


# 10. Sheets and limits ----------------------------------------------------------------------------

@pytest.mark.parametrize("name", ["a/b", "x" * 32, ""])
def test_an_invalid_sheet_name_is_an_error(name):
    with pytest.raises(XlsxError, match="sheet name"):
        mapped(Paragraph("a"), name=name)


def test_a_generated_sheet_name_that_repeats_another_report_is_an_error():
    pages = doc(Paragraph("a"), PageSection(A3_LANDSCAPE_REFERENCE, Paragraph("b"))).layout()
    with pytest.raises(XlsxError, match="used twice"):
        map_reports([("R", pages), ("r_2", pages[:1])], FONTS, cell_padding_pt=padding)


def test_more_columns_than_excel_has_is_an_error():
    profile = PageProfile("wide", 8300, 20, GridSpec(0, 0, 0.5, 10, 16_385, 1))
    page = raw_page(Text(1, 8, "a", STYLE, tag=paragraph_tag()), profile=profile)
    with pytest.raises(XlsxError, match="16384 columns"):
        map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)


def test_more_rows_than_excel_has_is_an_error():
    profile = PageProfile("tall", 20, 1_048_577 * 0.5 + 1, GridSpec(0, 0, 10, 0.5, 1, 1_048_577))
    page = raw_page(profile=profile)
    with pytest.raises(XlsxError, match="1048576 rows"):
        map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)


def test_reports_must_be_name_and_page_pairs():
    with pytest.raises(ValueError, match="nonempty sequence of ScenePage"):
        map_reports([("R", ())], FONTS, cell_padding_pt=padding)
    with pytest.raises(ValueError, match="nonempty sequence of ScenePage"):
        map_reports([("R", ("page",))], FONTS, cell_padding_pt=padding)
    with pytest.raises(ValueError, match="at least one report"):
        map_reports([], FONTS, cell_padding_pt=padding)
    with pytest.raises(ValueError, match="pair"):
        map_reports([("R",)], FONTS, cell_padding_pt=padding)
    with pytest.raises(ValueError, match="FontRegistry"):
        map_reports([("R", (raw_page(),))], object(), cell_padding_pt=padding)
    with pytest.raises(ValueError, match="strict"):
        map_reports([("R", (raw_page(),))], FONTS, strict=1, cell_padding_pt=padding)
    with pytest.raises(ValueError, match="function"):
        map_reports([("R", (raw_page(),))], FONTS, cell_padding_pt=2.4)


# 11. Untagged commands --------------------------------------------------------------------------

def test_untagged_commands_are_counted():
    page = raw_page(Text(x_of(0), baseline(0), "a", STYLE), Line(OX, OY + 40, OX + 30, OY + 40))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    assert dict(result.untagged) == {"Line": 1, "Text": 1}
    assert values(workbook.sheets[0]) == {(0, 0): "a"}


# 12. The file -------------------------------------------------------------------------------------

def test_render_xlsx_writes_the_same_bytes_and_returns_the_result(tmp_path):
    pages = doc(equation(), Paragraph("abc")).layout()
    first, second = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    result = render_xlsx([("R", pages)], first, FONTS, cell_padding_pt=padding)
    render_xlsx([("R", pages)], second, FONTS, cell_padding_pt=padding)
    assert isinstance(result, XlsxResult) and result.path == first
    assert first.read_bytes() == second.read_bytes()


def test_document_render_xlsx(tmp_path):
    target = tmp_path / "doc.xlsx"
    result = doc(Paragraph("abc")).render_xlsx(target, "R", cell_padding_pt=padding)
    assert target.exists() and result.sheets[0].name == "R" and result.path == Path(target)


# 13. More rules ------------------------------------------------------------------------------------

def fraction_tags():
    line = (Scope(Kind.EQUATION, index=0), Scope(Kind.EQUATION_LINE, index=0, data={"continuation": False,
                                                                                    "result_line": False}))
    rhs = (*line, Scope(Kind.EQUATION_PART, "rhs"))
    numerator = Tag(Role.MATH, (*rhs, Scope(Kind.FRACTION, "numerator")), {"atom": "literal"})
    denominator = Tag(Role.MATH, (*rhs, Scope(Kind.FRACTION, "denominator")), {"atom": "literal"})
    bar = Tag(Role.FRACTION_BAR, (*rhs, Scope(Kind.FRACTION, "bar")))
    lhs = Tag(Role.MATH, (*line, Scope(Kind.EQUATION_PART, "lhs")), {"atom": "literal"})
    return numerator, denominator, bar, lhs


def test_a_denominator_sharing_the_numerator_row_moves_down_and_is_recorded():
    numerator, denominator, bar, lhs = fraction_tags()
    page = raw_page(Text(x_of(2), baseline(3) - 6, "a", STYLE, tag=numerator),
                    Line(x_of(2), baseline(3) - 4, x_of(3), baseline(3) - 4, 0.5, tag=bar),
                    Text(x_of(2), baseline(3) + 3.5, "b", STYLE, tag=denominator),
                    Text(x_of(0), baseline(3), "I", STYLE, tag=lhs))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    sheet = workbook.sheets[0]
    assert values(sheet)[(3, 2)] == "a" and values(sheet)[(4, 2)] == "b"
    assert [(shift.text, shift.from_cell, shift.to_cell, shift.reason) for shift in result.shifts] == [
        ("b", (3, 2), (4, 2), "fraction")]
    assert merge_of(sheet, 3, 0) == xl.CellRange(3, 0, 4, 0)
    (line,) = sheet.drawings
    assert line.y1 == pytest.approx(4 * CELL)


def test_a_fraction_in_a_superscript_is_recorded_and_refused_by_strict():
    block = MathBlock(Script("x", superscript=Fraction("1", "2")))
    workbook, result = mapped(block)
    assert [item.kind for item in result.fallbacks] == ["fraction_in_script"]
    assert {_plain(cell.value) for cell in workbook.sheets[0].cells} == {"x", 1.0, 2.0}
    with pytest.raises(XlsxStrictError):
        mapped(block, strict=True)


def test_a_fraction_in_a_script_of_an_unstepped_size_is_written_at_the_nearest_step():
    # 12 pt x 0.65 = 7.8 pt: a fraction's cells inside a superscript have no base size to take.
    block = MathBlock(Script("x", superscript=Fraction("1", "2")), style=TextStyle(size=12))
    workbook, result = mapped(block)
    sizes = {cell.font.size for cell in workbook.sheets[0].cells if cell.value != "x"}
    assert sizes == {8.0} and [item.kind for item in result.fallbacks] == ["fraction_in_script"]


def test_a_denominator_moved_inside_a_script_names_the_script():
    numerator, denominator, bar, _ = fraction_tags()
    script = Scope(Kind.SCRIPT, "superscript")
    numerator, denominator, bar = (Tag(tag.role, (*tag.path[:3], script, *tag.path[3:]), tag.data)
                                   for tag in (numerator, denominator, bar))
    page = raw_page(Text(x_of(2), baseline(3) - 6, "1", STYLE, tag=numerator),
                    Line(x_of(2), baseline(3) - 4, x_of(3), baseline(3) - 4, 0.5, tag=bar),
                    Text(x_of(2), baseline(3) - 1, "2", STYLE, tag=denominator))
    _, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    assert [(shift.text, shift.reason) for shift in result.shifts] == [("2", "fraction_in_script")]


def test_turned_cells_move_along_their_reading_direction():
    def tag(line):
        turn = Scope(Kind.QUARTER_TURN, data={"turns": 1, "block_id": "t", "local_grid": {
            "origin_x": 0, "origin_y": 0, "pitch_x": PY, "pitch_y": PX, "columns": 4, "rows": 1}})
        return Tag(Role.TEXT, (turn, Scope(Kind.PARAGRAPH, index=line, data={"align": "left"})))

    # reading bottom to top, "ab" fills row 9 of column 1 and "cd" rows 8-9, so "cd" moves up to rows 7-8
    page = raw_page(Text(x_of(1, 12), OY + 10 * PY, "ab", STYLE, 1, tag=tag(0)),
                    Text(x_of(1, 12), OY + 10 * PY - 4, "cd", STYLE, 1, tag=tag(1)))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    sheet = workbook.sheets[0]
    assert values(sheet) == {(9, 1): "ab", (7, 1): "cd"}
    assert merge_of(sheet, 7, 1) == xl.CellRange(7, 1, 8, 1)
    (shift,) = result.shifts
    assert (shift.from_cell, shift.to_cell) == ((8, 1), (7, 1))
    assert {cell.rotation for cell in workbook.sheets[0].cells} == {90}


def test_a_clockwise_quarter_turn_is_270_degrees():
    block = QuarterTurn(Paragraph("abc"), turns=-1, local_grid=GridSpec(0, 0, PY, PX, 4, 2))
    cell = text_cell(first_sheet(mapped(block)), "abc")
    assert cell.rotation == 270 and cell.v_align == "top"


def test_a_table_inside_a_swapped_quarter_turn_keeps_cell_borders():
    table = Table([2, 2], [[Cell("a"), Cell("b")]])
    block = QuarterTurn(table, turns=1, local_grid=GridSpec(0, 0, PY, PX, 4, 3))
    workbook, result = mapped(block)
    assert workbook.sheets[0].borders and not result.fallbacks
    assert not workbook.sheets[0].drawings


def test_the_thicker_of_two_borders_on_one_edge_is_kept():
    table = Table([2], [[Cell("a")]])
    workbook, result = mapped(table, Rule("thick", columns=2))
    sheet = workbook.sheets[0]
    bottom = max(border.edge_key()[1] for border in sheet.borders if border.edge_key()[0] == "horizontal")
    styles = {border.style for border in sheet.borders if border.edge_key()[:2] == ("horizontal", bottom)}
    assert styles == {"thick"} and not result.fallbacks  # a thick solid line covers a thin one


def test_a_double_border_is_one_edge():
    sheet = first_sheet(mapped(Rule("double", columns=2)))
    assert [(border.row, border.col, border.side, border.style) for border in sheet.borders] == [
        (0, 0, "top", "double"), (0, 1, "top", "double")]


def test_text_left_of_the_grid_is_an_error():
    page = raw_page(Text(OX - 5, baseline(0), "a", STYLE, tag=paragraph_tag()))
    with pytest.raises(XlsxError, match="left edge"):
        map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)


def test_a_paper_other_than_a4_or_a3_is_an_error():
    profile = PageProfile("letter", 612, 792, GridSpec(50, 50, 16, 16, 30, 40))
    with pytest.raises(XlsxError, match="A4 or A3"):
        map_reports([("R", (raw_page(profile=profile),))], FONTS, cell_padding_pt=padding)


def test_a_decoration_that_changes_between_pages_starts_the_next_sheet():
    pages = doc(Paragraph("a"), profile=A4_BODY, footer=PageDecoration("{page}")).layout()
    other = doc(Paragraph("b"), profile=A4_BODY, footer=PageDecoration("p{page}")).layout()
    moved = doc(Paragraph("c"), profile=PageProfile("A4 lower", 595, 842, A4.grid, 3, 44),
                footer=PageDecoration("p{page}")).layout()
    workbook, result = map_reports([("R", (*pages, *other, *moved))], FONTS, cell_padding_pt=padding)
    assert [sheet.name for sheet in workbook.sheets] == ["R", "R_2", "R_3"]
    assert workbook.sheets[1].footer.left.parts == ("p", xl.PageNumber())
    assert [(info.first_page, info.last_page) for info in result.sheets] == [(1, 1), (2, 2), (3, 3)]


def test_header_and_footer_print_where_the_pdf_draws_them():
    workbook, _ = mapped(Paragraph("a"), profile=A4_BODY, header=PageDecoration("見出し", align="left"),
                         footer=PageDecoration("{page}", align="center"))
    sheet = workbook.sheets[0]
    assert sheet.header.left.parts == ("見出し",) and sheet.footer.center.parts == (xl.PageNumber(),)
    assert sheet.page.margins.header == pytest.approx(OY)  # the header frame is grid row 0
    assert sheet.page.margins.footer == pytest.approx(842 - OY - 47 * PY)  # the footer frame is the last row
    assert all(cell.value != "見出し" for cell in sheet.cells)


def test_render_xlsx_uses_the_library_cell_padding_when_none_is_given(tmp_path, monkeypatch):
    import openpyxl

    from gridlayoutpdf import equations

    pages = doc(Table([1], [[Cell("12", padding=1)]])).layout()  # 12 pt in one Excel column of 15.75 pt
    for padding_pt, expected in ((1.0, 12), (2.4, "12")):
        monkeypatch.setattr(equations, "excel_cell_padding_pt", lambda size, value=padding_pt: value)
        target = tmp_path / f"{padding_pt}.xlsx"
        render_xlsx([("R", pages)], target, FONTS)
        sheet = openpyxl.load_workbook(target)["R"]
        assert [cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None] == [expected]


def test_a_term_column_keeps_its_alignment():
    term = Tag(Role.MATH, (Scope(Kind.EQUATION_PART, "term", 1, {"align": "right"}),), {"atom": "literal"})
    page = raw_page(Text(x_of(4), baseline(0), "ab", STYLE, tag=term))
    workbook, _ = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    assert workbook.sheets[0].cells[0].h_align == "right"


def test_separate_placements_with_the_same_path_are_separate_pieces():
    sheet = first_sheet(mapped(Paragraph("ab"), Paragraph("cd")))
    assert values(sheet) == {(0, 0): "ab", (1, 0): "cd"}


def test_upside_down_text_outside_a_figure_is_a_shape_and_recorded():
    page = raw_page(Text(x_of(5), baseline(2), "ab", STYLE, 2, tag=paragraph_tag()))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    (box,) = workbook.sheets[0].drawings
    assert box.text.rotation == 180 and [item.kind for item in result.fallbacks] == ["rotated_text"]


def test_shapes_of_no_extent_are_dropped_and_listed():
    other = Tag(Role.OTHER)
    page = raw_page(Line(OX + 5, OY + 5, OX + 5, OY + 5, tag=other),
                    Rectangle(OX + 5, OY + 5, 0, 10, tag=other))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    (line,) = workbook.sheets[0].drawings
    assert isinstance(line, xl.Line) and len(result.dropped) == 1


def test_a_relation_centred_on_its_anchor_takes_the_column_right_of_it():
    sheet = first_sheet(mapped(equation()))  # first_equal is column boundary 2
    relations = [cell for cell in sheet.cells if cell.value == "="]
    assert relations[0].col == 2 and relations[0].h_align == "center"
    lhs = text_cell(sheet, "I")
    assert (merge_of(sheet, lhs.row, lhs.col) or xl.CellRange(lhs.row, lhs.col, lhs.row, lhs.col + 1)).last_col <= 1
    assert relations[1].col == 12  # second_equal


def grid_group(*, padding_pt=None):
    """ "15" in a right-aligned term column, then "2": terms of whole columns (term_columns="grid")."""
    return EquationGroup(Equation("A", Row("15", TermStop("right"), Operator("+"), TermStop("left"), "2", gap=0)),
                         anchors=EquationAnchors(2), term_columns="grid")


def term_columns(pages, text):
    """(first page column, column count) of the term holding ``text``, from its tag and the block's column."""
    command = next(command for command in pages[0].commands if isinstance(command, Text) and command.text == text)
    term = command.tag.scope(Kind.EQUATION_PART)
    return term.data["first_column"], term.data["columns"]


def term_cells(sheet, text):
    cell = text_cell(sheet, text)
    merge = merge_of(sheet, cell.row, cell.col)
    return (cell.col, cell.col) if merge is None else (merge.first_col, merge.last_col)


@pytest.mark.parametrize("before", [0, 3])
def test_a_grid_term_takes_the_columns_its_tag_names_from_its_block(before):
    group = grid_group()
    blocks = (Columns(Paragraph("x"), group, col_spans=(before, 26)),) if before else (group,)
    pages = doc(*blocks).layout()
    workbook, result = map_reports([("R", pages)], FONTS, cell_padding_pt=padding)
    first, count = term_columns(pages, "15")
    if before:  # the second child of the Columns starts after the first and the gutter
        before = round(pages[0].metadata["placements"][0]["parts"][0]["metadata"]["children"][1]["x_pt"] / PX)
    assert term_cells(workbook.sheets[0], "15") == (before + first, before + first + count - 1)
    command = next(command for command in pages[0].commands if isinstance(command, Text) and command.text == "15")
    align = command.tag.scope(Kind.EQUATION_PART).data["align"]
    assert text_cell(workbook.sheets[0], "15").h_align == align and not result.fallbacks and not result.shifts


@pytest.mark.parametrize(("block", "reason"), [
    (lambda group: Indent(group, 5), "not on a page column boundary"),
    (lambda group: Table([30], [[Cell(group)]]), "not known"),
])
def test_a_grid_term_off_the_page_columns_is_recorded(block, reason):
    pages = doc(block(grid_group())).layout()
    _, result = map_reports([("R", pages)], FONTS, cell_padding_pt=padding)
    kinds = {item.kind for item in result.fallbacks}
    assert kinds == {"term_off_grid"} and all(reason in item.detail for item in result.fallbacks)
    with pytest.raises(XlsxStrictError):
        map_reports([("R", pages)], FONTS, strict=True, cell_padding_pt=padding)


def test_a_grid_term_laid_out_with_another_cell_padding_is_recorded():
    pages = doc(grid_group()).layout()  # laid out with the library's padding
    _, result = map_reports([("R", pages)], FONTS, cell_padding_pt=lambda size: 2.0)
    assert {item.kind for item in result.fallbacks} == {"cell_padding_mismatch"}


def test_grid_term_glyphs_outside_their_columns_are_recorded():
    term = Tag(Role.MATH, (Scope(Kind.EQUATION, index=0), Scope(Kind.EQUATION_LINE, index=0),
                           Scope(Kind.EQUATION_PART, "term", 0, {"align": "left", "first_column": 0, "columns": 1,
                                                                  "cell_padding_pt": PADDING})),
               {"atom": "literal"})
    metadata = {"placements": [{"col": 4, "commands": (0, 1), "parts": [
        {"row": 0, "rows": 1, "metadata": {"equation_index": 0, "line_index": 0}}]}]}
    page = raw_page(Text(x_of(4, PADDING), baseline(0), "abc", STYLE, tag=term), metadata=metadata)
    _, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    assert [item.kind for item in result.fallbacks] == ["term_off_grid"]
    assert "leave its term columns" in result.fallbacks[0].detail


def test_the_number_room_check_allows_the_shared_roundoff():
    pages = doc(Paragraph("12.5")).layout()  # 24 pt on Excel's screen in two columns of 31.5 pt
    almost = (2 * CELL - 24 + 0.5e-7) / 2
    workbook, _ = map_reports([("R", pages)], FONTS, cell_padding_pt=lambda size: almost)
    assert text_cell(workbook.sheets[0], "12.5").number_format == "0.0"


def test_a_wrapped_equation_label_keeps_each_line_in_its_own_row():
    group = EquationGroup(Equation("I", Fraction("a", "b"), label="abcd efgh"), label_columns=2,
                          anchors=EquationAnchors(6))
    workbook, result = mapped(group)
    sheet = workbook.sheets[0]
    first, second = text_cell(sheet, "abcd"), text_cell(sheet, "efgh")
    assert second.row == first.row + 1 and not result.shifts
    assert merge_of(sheet, first.row, first.col) is None or merge_of(sheet, first.row, first.col).last_row == first.row


class InkFonts(FixedFonts):
    """Glyph ink 1 pt inside each end of the advance, as real fonts' side bearings."""

    def measure(self, text, style):
        metrics = super().measure(text, style)
        return TextMetrics(metrics.width, metrics.ascent, metrics.descent,
                           ink_bounds=Rect(1, -metrics.ascent, max(0.0, metrics.width - 2), metrics.ascent))


def test_a_piece_takes_the_columns_its_ink_touches_not_its_side_bearings():
    # "ab" ends at column 3's left boundary + 0.5 pt, but its ink ends 0.5 pt before that boundary
    page = raw_page(Text(x_of(3, 0.5) - 12, baseline(0), "ab", STYLE, tag=paragraph_tag(line=0)),
                    Text(x_of(3, 0.2), baseline(0), "cd", STYLE, tag=paragraph_tag(line=1)))
    workbook, result = map_reports([("R", (page,))], InkFonts(), cell_padding_pt=padding)
    assert values(workbook.sheets[0]) == {(0, 2): "ab", (0, 3): "cd"} and not result.shifts


# 14. Review round 1 ---------------------------------------------------------------------------------

def border_tag(role=Role.TABLE_BORDER, style="thin", **data):
    return Tag(role, (), {"style": style, "stroke": 0, **data})


def test_a_table_split_at_a_page_break_keeps_the_next_page_top_border():
    rows = [[Cell(f"r{index}")] for index in range(60)]
    workbook, result = mapped(Paragraph("x"), Table([2], rows))  # page 1's part of the table ends at row 47
    sheet = workbook.sheets[0]
    sides = {(border.row, border.side) for border in sheet.borders if border.col == 0}
    assert (47, "top") in sides and (46, "bottom") in sides  # the PDF draws both: page 1 closes, page 2 opens
    assert not result.fallbacks


def test_the_model_accepts_a_bottom_and_a_top_border_only_at_a_page_break():
    common = dict(column_widths_pt=(PX,), row_heights_pt=(PY,) * 4)
    pair = (xl.EdgeBorder(1, 0, "bottom", "thin"), xl.EdgeBorder(2, 0, "top", "thin"))
    xl.Sheet("S", **common, borders=pair, row_breaks=(2,))
    with pytest.raises(ValueError, match="same edge"):
        xl.Sheet("S", **common, borders=pair)


def test_a_turned_side_rule_is_a_horizontal_border():
    block = QuarterTurn(SideRule(Paragraph("ab"), PY), turns=1, local_grid=GridSpec(0, 0, PY, PX, 4, 2))
    workbook, result = mapped(block)
    sheet = workbook.sheets[0]
    assert sheet.borders and {border.side for border in sheet.borders} <= {"top", "bottom"}
    assert not result.fallbacks


def test_a_border_shorter_than_its_cells_is_a_recorded_shape():
    workbook, result = mapped(SideRule(Paragraph("ab"), PX, top=8))
    assert not workbook.sheets[0].borders and len(workbook.sheets[0].drawings) == 1
    assert [item.kind for item in result.fallbacks] == ["border_off_grid"]


def test_a_border_whose_ends_are_off_the_boundaries_is_a_recorded_shape():
    line = Line(x_of(1, 6), OY + PY, x_of(4, -6), OY + PY, 0.5, tag=border_tag(Role.RULE))
    workbook, result = map_reports([("R", (raw_page(line),))], FONTS, cell_padding_pt=padding)
    assert not workbook.sheets[0].borders and [item.kind for item in result.fallbacks] == ["border_off_grid"]


def test_a_table_seam_bridge_is_not_a_border_or_a_record():
    bridge = Line(x_of(1), OY + PY - 0.25, x_of(1), OY + PY, 0.5, tag=border_tag(drawn="continuing"))
    workbook, result = map_reports([("R", (raw_page(bridge),))], FONTS, cell_padding_pt=padding)
    assert not workbook.sheets[0].borders and not workbook.sheets[0].drawings and not result.fallbacks


def test_a_side_rule_border_takes_its_fill_color():
    sheet = first_sheet(mapped(SideRule(Paragraph("ab"), PX, color=(1, 0, 0))))
    assert {border.color for border in sheet.borders} == {(1, 0, 0)}


def test_borders_on_one_edge_that_differ_in_more_than_thickness_are_recorded():
    lines = (Line(OX, OY + PY, x_of(2), OY + PY, 0.5, tag=border_tag(Role.RULE)),
             Line(OX, OY + PY, x_of(2), OY + PY, 0.5, (1, 0, 0), tag=border_tag(Role.RULE)))
    workbook, result = map_reports([("R", (raw_page(*lines),))], FONTS, cell_padding_pt=padding)
    assert {border.color for border in workbook.sheets[0].borders} == {(0, 0, 0)}
    assert [item.kind for item in result.fallbacks] == ["border_conflict"]


def test_a_dashed_border_gives_way_to_a_solid_one_of_the_same_thickness():
    lines = (Line(OX, OY + PY, x_of(2), OY + PY, 0.5, dash=(3, 1.5), tag=border_tag(Role.RULE, "dashed")),
             Line(OX, OY + PY, x_of(2), OY + PY, 0.5, tag=border_tag(Role.RULE)))
    workbook, result = map_reports([("R", (raw_page(*lines),))], FONTS, cell_padding_pt=padding)
    assert {border.style for border in workbook.sheets[0].borders} == {"thin"}
    assert [item.kind for item in result.fallbacks] == ["border_conflict"]


def test_atoms_more_than_half_a_cell_apart_are_separate_pieces():
    sheet = first_sheet(mapped(MathBlock(Row("a", "b", "c", gap=40))))
    assert len([cell for cell in sheet.cells if cell.value is not None]) == 3


def test_a_drawing_in_the_margin_of_a_later_page_is_an_error():
    inside = raw_page(Text(x_of(0), baseline(0), "a", STYLE, tag=paragraph_tag()))
    outside = raw_page(Line(OX, OY - 20, OX + 30, OY - 20, tag=Tag(Role.OTHER)))
    with pytest.raises(XlsxError, match="page 2"):
        map_reports([("R", (inside, outside))], FONTS, cell_padding_pt=padding)


def test_a_two_point_closed_polyline_is_a_line():
    from gridlayoutpdf.scene import Polyline

    shape = Polyline(((OX + 5, OY + 5), (OX + 25, OY + 5), (OX + 5, OY + 5)), closed=True, tag=Tag(Role.OTHER))
    workbook, result = map_reports([("R", (raw_page(shape),))], FONTS, cell_padding_pt=padding)
    assert isinstance(workbook.sheets[0].drawings[0], xl.Line) and not result.dropped


def test_too_many_pages_for_excel_page_breaks_is_an_xlsx_error():
    pages = (raw_page(),) * 1028
    with pytest.raises(XlsxError, match="row breaks"):
        map_reports([("R", pages)], FONTS, cell_padding_pt=padding)


def test_a_font_name_excel_cannot_hold_is_an_xlsx_error():
    page = raw_page(Text(x_of(0), baseline(0), "a", TextStyle(font="x" * 32), tag=paragraph_tag()))
    with pytest.raises(XlsxError, match="page 1"):
        map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)


def test_a_broken_image_is_an_xlsx_error():
    page = raw_page(RasterImage(OX, OY, 10, 10, b"not an image", tag=Tag(Role.IMAGE)))
    with pytest.raises(XlsxError, match="image"):
        map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)


def test_a_bad_cell_padding_is_an_error():
    pages = doc(Paragraph("12")).layout()
    with pytest.raises(ValueError, match="cell_padding_pt"):
        map_reports([("R", pages)], FONTS, cell_padding_pt=lambda size: -1)


def test_records_check_their_values_and_are_immutable():
    import copy
    import pickle

    from gridlayoutpdf.renderers.xlsx import XlsxFallback, XlsxSheet, XlsxShift

    with pytest.raises(ValueError):
        XlsxShift("R", 1, "a", (0, 0), (0, 1), "bogus")
    with pytest.raises(ValueError):
        XlsxShift("R", 1, "a", (-1, 0), (0, 1), "overlap")
    with pytest.raises(ValueError):
        XlsxFallback("R", 1, "nope", "")
    with pytest.raises(ValueError):
        XlsxSheet("R", "R", 5, 2, 47, 30)
    counts = {"Text": 1}
    result = XlsxResult(None, (), untagged=counts)
    counts["Text"] = 5
    assert dict(result.untagged) == {"Text": 1} and hash(result) == hash(XlsxResult(None, (), untagged={"Text": 1}))
    _, result = map_reports([("R", (overlapping_page(),))], FONTS, cell_padding_pt=padding)
    error = XlsxStrictError("message", result.shifts, ())
    for clone in (pickle.loads(pickle.dumps(error)), copy.copy(error)):
        assert clone.shifts == result.shifts and str(clone) == "message"


def test_a_number_that_excel_draws_wider_on_screen_stays_text():
    # 9 pt glyphs of 5.4 pt (7.2 px) take 8 px each on Excel's screen: 43.2 pt of text is 48 pt there.
    nine = TextStyle(size=9)
    pages = doc(Table([3], [[Cell("88888888", style=nine, padding=1)]])).layout()  # 3 columns: Excel 47.25 pt
    workbook, _ = map_reports([("R", pages)], FONTS, cell_padding_pt=padding)
    assert text_cell(workbook.sheets[0], "88888888").number_format is None
    pages = doc(Table([3], [[Cell("8888888", style=nine, padding=1)]])).layout()  # 56 px = 42 pt + 4.8 pt fits
    workbook, _ = map_reports([("R", pages)], FONTS, cell_padding_pt=padding)
    assert text_cell(workbook.sheets[0], "8888888").number_format == "0"


def test_a_number_grows_into_free_cells_away_from_its_alignment():
    # "8888888" (7 glyphs of 8 px = 42 pt) right-aligned in a 4-column table cell: its ink touches 3 columns
    # (47.25 pt), 42 + 2 x 4 pt padding needs a fourth, which it takes on the left, inside the cell.
    table = Table([1, 4], [[Cell("a"), Cell("8888888")]], numeric_columns=(1,))
    pages = doc(table).layout()
    workbook, result = map_reports([("R", pages)], FONTS, cell_padding_pt=lambda size: 4.0)
    cell = text_cell(workbook.sheets[0], "8888888")
    merge = merge_of(workbook.sheets[0], cell.row, cell.col)
    assert cell.number_format == "0" and (merge.first_col, merge.last_col) == (1, 4)
    assert cell.h_align == "right" and not result.numbers_as_text


def test_a_number_does_not_grow_across_a_border():
    table = Table([1, 3], [[Cell("a"), Cell("8888888")]], numeric_columns=(1,))  # the cell has 3 columns only
    pages = doc(table).layout()
    workbook, result = map_reports([("R", pages)], FONTS, cell_padding_pt=lambda size: 4.0)
    assert text_cell(workbook.sheets[0], "8888888").number_format is None
    assert [item.kind for item in result.numbers_as_text] == ["number_as_text"]


# 15. Review round 2 ---------------------------------------------------------------------------------

def numerator_tag():
    numerator, _, _, _ = fraction_tags()
    return numerator


def test_a_centred_number_grows_on_both_sides_at_once():
    # 42 pt of glyphs (56 px) in columns 5-7 (47.25 pt): 2 x 4 pt padding needs 50 pt, one cell more on each side
    page = raw_page(Text(x_of(5, 3), baseline(3), "8888888", STYLE, tag=numerator_tag()))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=lambda size: 4.0)
    assert merge_of(workbook.sheets[0], 3, 5) == xl.CellRange(3, 4, 3, 8) and not result.numbers_as_text


def test_a_centred_number_with_one_side_taken_stays_text():
    other = Tag(Role.TEXT, (Scope(Kind.PARAGRAPH, index=9, data={"align": "left"}),))
    page = raw_page(Text(x_of(5, 3), baseline(3), "8888888", STYLE, tag=numerator_tag()),
                    Text(x_of(8, 1), baseline(3), "a", STYLE, tag=other))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=lambda size: 4.0)
    assert text_cell(workbook.sheets[0], "8888888").number_format is None
    assert [item.kind for item in result.numbers_as_text] == ["number_as_text"]


def test_a_number_does_not_grow_over_a_border_between_its_rows():
    numerator, denominator, bar, lhs = fraction_tags()
    page = raw_page(Text(x_of(3, 3), baseline(3), "88888", STYLE, tag=lhs),  # cols 3-4, merged over rows 3-4
                    Text(x_of(8), baseline(3) - 6, "a", STYLE, tag=numerator),
                    Line(x_of(8), baseline(3) - 4, x_of(9), baseline(3) - 4, 0.5, tag=bar),
                    Text(x_of(8), baseline(3) + 8, "b", STYLE, tag=denominator),
                    Line(x_of(2), OY + 4 * PY, x_of(3), OY + 4 * PY, 0.5, tag=border_tag(Role.RULE)))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=lambda size: 9.0)
    assert text_cell(workbook.sheets[0], "88888").number_format is None  # growing left would cover the rule
    assert not result.fallbacks and len(result.numbers_as_text) == 1


def test_a_header_longer_than_excel_holds_is_an_xlsx_error():
    with pytest.raises(XlsxError, match="255"):
        mapped(Paragraph("a"), profile=A4_BODY, header=PageDecoration("&" * 200, style=TextStyle(size=4)))


def test_result_sequences_are_not_strings():
    with pytest.raises(ValueError):
        XlsxResult(None, (), dropped="abc")


def test_nested_indents_place_a_left_aligned_grid_term_of_several_columns():
    group = EquationGroup(Equation("A", Row("1", TermStop("left"), "123456", gap=0)), anchors=EquationAnchors(2),
                          term_columns="grid")
    pages = doc(Indent(Indent(group, PX), PX)).layout()
    workbook, result = map_reports([("R", pages)], FONTS, cell_padding_pt=padding)
    first, count = term_columns(pages, "123456")
    assert count >= 2 and term_cells(workbook.sheets[0], "123456") == (2 + first, 2 + first + count - 1)
    assert not result.fallbacks and not result.shifts


def test_a_grid_term_split_into_several_pieces_is_recorded():
    group = EquationGroup(Equation("A", Row("2", Fraction("a", "b"), TermStop("left"), "3", gap=0)),
                          anchors=EquationAnchors(2),
                          term_columns="grid")
    _, result = map_reports([("R", doc(group).layout())], FONTS, cell_padding_pt=padding)
    assert any(item.kind == "term_off_grid" and "split" in item.detail for item in result.fallbacks)


def axis(pitch, count, column=True):
    from gridlayoutpdf.renderers.xlsx import _map

    table, ratio = ((_map.PRINTED_COLUMN_PT, _map.PRINT_COLUMN_PT_PER_PX) if column
                    else (_map.PRINTED_ROW_PT, _map.PRINT_ROW_PT_PER_PX))
    return _map._Axis.build(pitch, count, table, ratio, "column" if column else "row")


def test_the_a4_reference_grid_keeps_its_21_pixel_cells():
    sheet = first_sheet(mapped(Paragraph("a")))
    assert set(sheet.column_widths_pt) == set(sheet.row_heights_pt) == {CELL}


@pytest.mark.parametrize(("pitch", "count", "column", "pixels"),
                         [(10.8, 46, True, {14, 15}), (14.52, 51, False, {19, 20}), (16.08, 30, True, {21})])
def test_cells_take_whole_pixels_whose_printed_boundaries_stay_near_the_pdf(pitch, count, column, pixels):
    built = axis(pitch, count, column)
    assert {round(size / 0.75) for size in built.sizes} == pixels
    assert max(abs(printed - k * pitch) for k, printed in enumerate(built.printed_starts)) <= 0.5


def test_a_pitch_under_one_pixel_is_an_error():
    with pytest.raises(XlsxError, match="half a cell"):
        axis(0.25, 40)


def test_shapes_keep_their_fraction_of_their_cell_on_a_distributed_grid():
    built = axis(10.8, 46)
    for offset in (0.0, 5.4, 10.8 * 7 + 2.7, 10.8 * 46):
        index = min(int(offset // 10.8), 45)
        fraction = (offset - index * 10.8) / 10.8
        assert built.at(offset) == pytest.approx(built.starts[index] + fraction * built.sizes[index])


def test_pages_that_print_under_100_percent_are_recorded():
    # fits the PDF's A4, not the 100 % rule: 54 x 15.48 x 1.006 + 4.5 = 845.4 pt (full_scale_size_pt)
    grid = GridSpec(50.88, 2.0, 16.08, 15.48, 30, 54)
    page = raw_page(Text(x_of(0), baseline(0), "a", STYLE, tag=paragraph_tag()),
                    profile=PageProfile("tall", 595, 842, grid))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    assert workbook.sheets[0].page.scale < 100 and result.sheets[0].scale == workbook.sheets[0].page.scale
    assert [item.kind for item in result.fallbacks] == ["print_scale"]
    with pytest.raises(XlsxStrictError):
        map_reports([("R", (page,))], FONTS, strict=True, cell_padding_pt=padding)


def test_a_fraction_in_a_script_of_a_turned_table_is_recorded():
    block = QuarterTurn(MathBlock(Row("a", Script("b", superscript=Fraction("1", "2")))), turns=1,
                        local_grid=GridSpec(0, 0, PY, PX, 8, 4))
    _, result = mapped(*(Paragraph(f"p{index}") for index in range(20)), block)  # room above to read upwards
    assert "fraction_in_script" in {item.kind for item in result.fallbacks}


def test_a_number_written_in_rich_text_is_recorded():
    other = TextStyle(font="IPAexゴシック", size=10)
    page = raw_page(Text(x_of(0), baseline(0), "12", STYLE, tag=paragraph_tag()),
                    Text(x_of(0, 12), baseline(0), ".5", other, tag=paragraph_tag()))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    assert isinstance(workbook.sheets[0].cells[0].value, xl.RichText)
    assert [item.kind for item in result.numbers_as_text] == ["number_as_text"]


def test_a_seam_bridge_is_listed_as_dropped():
    bridge = Line(x_of(1), OY + PY - 0.25, x_of(1), OY + PY, 0.5, tag=border_tag(drawn="continuing"))
    _, result = map_reports([("R", (raw_page(bridge),))], FONTS, cell_padding_pt=padding)
    assert len(result.dropped) == 1 and "page 1" in result.dropped[0]


def test_strict_errors_and_shifts_check_their_records():
    from gridlayoutpdf.renderers.xlsx import XlsxShift

    with pytest.raises(ValueError):
        XlsxStrictError("m", ["not a shift"])
    with pytest.raises(ValueError, match="move"):
        XlsxShift("R", 1, "a", (0, 1), (0, 1), "overlap")


def test_incomplete_page_metadata_is_an_xlsx_error():
    page = raw_page(Text(x_of(0), baseline(0), "a", STYLE, tag=paragraph_tag()),
                    metadata={"decorations": {"header": {"template": "x"}}})
    with pytest.raises(XlsxError, match="metadata"):
        map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)


# 16. Review round 3 ---------------------------------------------------------------------------------

@pytest.mark.parametrize("other", ["thin", "medium", "thick"])
@pytest.mark.parametrize("double_first", [True, False])
def test_a_double_border_and_a_single_one_on_one_edge_are_a_conflict(other, double_first):
    from gridlayoutpdf.borders import BorderStyle

    def lines(style):
        return [replace(line, tag=border_tag(Role.RULE, style, stroke=index))
                for index, line in enumerate(BorderStyle(style).lines(OX, OY + PY, x_of(2), OY + PY, (0, 0, 0)))]

    first, second = (lines("double"), lines(other)) if double_first else (lines(other), lines("double"))
    _, result = map_reports([("R", (raw_page(*first, *second),))], FONTS, cell_padding_pt=padding)
    assert [item.kind for item in result.fallbacks] == ["border_conflict"]


def test_a_grid_term_number_does_not_grow_past_its_term_columns():
    # 9 pt: "88888888" fits its term columns in the PDF but needs a pixel more on Excel's screen
    nine = TextStyle(size=9)
    group = EquationGroup(Equation("A", Row("x", TermStop("left"), "88888888", gap=0)), anchors=EquationAnchors(2),
                          term_columns="grid", style=nine)
    pages = doc(group).layout()
    workbook, result = map_reports([("R", pages)], FONTS, cell_padding_pt=padding)
    first, count = term_columns(pages, "88888888")
    assert term_cells(workbook.sheets[0], "88888888") == (first, first + count - 1)
    cell = text_cell(workbook.sheets[0], "88888888")
    assert cell.number_format is None and len(result.numbers_as_text) == 1
    assert "grid term" in result.numbers_as_text[0].detail


def test_a_grid_term_of_a_fraction_alone_is_recorded():
    group = EquationGroup(Equation("A", Row("1", TermStop("left"), Fraction("a", "b"), gap=0)),
                          anchors=EquationAnchors(2), term_columns="grid")
    _, result = map_reports([("R", doc(group).layout())], FONTS, cell_padding_pt=padding)
    details = [item.detail for item in result.fallbacks if item.kind == "term_off_grid"]
    assert len(details) == 1 and "fraction" in details[0]


@pytest.mark.parametrize("metadata", [
    {"placements": [{"col": 0, "commands": (0, 5)}]},
    {"placements": [{"col": 0, "commands": None}]},
    {"placements": ["x"]},
    {"decorations": ["header"]},
])
def test_malformed_page_metadata_is_an_xlsx_error(metadata):
    page = raw_page(Text(x_of(0), baseline(0), "a", STYLE, tag=paragraph_tag()), metadata=metadata)
    with pytest.raises(XlsxError, match="metadata"):
        map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)


@pytest.mark.parametrize("node", [Script("10", superscript="3"), Script("2.5", subscript="1")])
def test_a_scripted_number_is_not_a_number_kept_as_text(node):
    _, result = mapped(MathBlock(node))  # "103" and "2.51" read as numbers, but they are 10³ and 2.5₁
    assert not result.numbers_as_text


def test_records_follow_the_page_order_whatever_the_hash_seed():
    import os
    import subprocess
    import sys

    script = ("import sys; sys.path.insert(0, 'tests/unit'); import test_xlsx_map as t; "
              "print(t.shift_texts())")
    outputs = set()
    for seed in ("1", "2", "3"):
        completed = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True,
                                   env={**os.environ, "PYTHONHASHSEED": seed}, cwd=Path(__file__).parents[2])
        assert completed.returncode == 0, completed.stderr
        outputs.add(completed.stdout)
    assert len(outputs) == 1


def shift_texts():
    """Shifts of four sibling fractions whose denominators all move down (used by the hash seed test)."""
    numerator, denominator, bar, _ = fraction_tags()
    commands = []
    for index in range(4):
        row = Scope(Kind.ROW, index=index)
        tags = [Tag(tag.role, (*tag.path[:3], row, *tag.path[3:]), tag.data) for tag in (numerator, denominator, bar)]
        x = x_of(2 + 3 * index)
        commands += [Text(x, baseline(3) - 6, f"n{index}", STYLE, tag=tags[0]),
                     Line(x, baseline(3) - 4, x + PX, baseline(3) - 4, 0.5, tag=tags[2]),
                     Text(x, baseline(3) - 1, f"d{index}", STYLE, tag=tags[1])]
    _, result = map_reports([("R", (raw_page(*commands),))], FONTS, cell_padding_pt=padding)
    return [shift.text for shift in result.shifts]


def turned_tag(turns, align):
    turn = Scope(Kind.QUARTER_TURN, data={"turns": turns, "block_id": "t", "local_grid": {
        "origin_x": 0, "origin_y": 0, "pitch_x": PY, "pitch_y": PX, "columns": 4, "rows": 1}})
    return Tag(Role.TEXT, (turn, Scope(Kind.PARAGRAPH, index=0, data={"align": align})))


@pytest.mark.parametrize(("quarter_turns", "align", "rows", "v_align"), [
    (1, "left", (8, 11), "bottom"), (1, "right", (9, 12), "top"), (1, "center", (8, 12), "center"),
    (3, "left", (9, 12), "top"), (3, "right", (8, 11), "bottom"), (3, "center", (8, 12), "center"),
])
def test_a_turned_number_grows_away_from_its_reading_alignment(quarter_turns, align, rows, v_align):
    # 42 pt of glyphs in rows 9-11 (47.25 pt) need 50 pt with 2 x 4 pt padding: one row more
    if quarter_turns == 1:  # reads upwards from the bottom of row 11
        text = Text(x_of(1, 12), OY + 12 * PY, "8888888", STYLE, 1, tag=turned_tag(1, align))
    else:  # reads downwards from the top of row 9
        text = Text(x_of(1, 4), OY + 9 * PY + 0.5, "8888888", STYLE, 3, tag=turned_tag(-1, align))
    workbook, result = map_reports([("R", (raw_page(text),))], FONTS, cell_padding_pt=lambda size: 4.0)
    cell = text_cell(workbook.sheets[0], "8888888")
    merge = merge_of(workbook.sheets[0], cell.row, cell.col)
    assert (merge.first_row, merge.last_row) == rows and cell.v_align == v_align
    assert cell.number_format == "0" and not result.numbers_as_text


def test_a_turned_number_does_not_grow_across_a_border():
    text = Text(x_of(1, 12), OY + 12 * PY, "8888888", STYLE, 1, tag=turned_tag(1, "left"))
    rule = Line(x_of(1), OY + 9 * PY, x_of(2), OY + 9 * PY, 0.5, tag=border_tag(Role.RULE))  # above row 9
    workbook, result = map_reports([("R", (raw_page(text, rule),))], FONTS, cell_padding_pt=lambda size: 4.0)
    assert text_cell(workbook.sheets[0], "8888888").number_format is None and result.numbers_as_text


def test_an_unknown_term_alignment_is_left():
    term = Tag(Role.MATH, (Scope(Kind.EQUATION_PART, "term", 0, {"align": "justify"}),), {"atom": "literal"})
    workbook, _ = map_reports([("R", (raw_page(Text(x_of(4), baseline(0), "ab", STYLE, tag=term)),))], FONTS,
                              cell_padding_pt=padding)
    assert workbook.sheets[0].cells[0].h_align == "left"


@pytest.mark.parametrize("values", [{"untagged": 5}, {"dropped": 5}, {"dropped": None}, {"untagged": [1]},
                                    {"untagged": [("a",)]}])
def test_result_counts_must_be_sequences_of_pairs(values):
    with pytest.raises(ValueError):
        XlsxResult(None, (), **values)


def test_a_strict_error_does_not_hold_numbers_kept_as_text():
    from gridlayoutpdf.renderers.xlsx import XlsxFallback

    with pytest.raises(ValueError):
        XlsxStrictError("m", (), (XlsxFallback("R", 1, "number_as_text", ""),))


TURNED_GRID = PageProfile("turned", 595, 842, GridSpec(50.88, 54.598, 10.8, 14.52, 46, 51))


def test_bars_and_border_shapes_lie_on_the_boundaries_of_a_distributed_grid():
    grid = TURNED_GRID.grid
    numerator, denominator, bar, _ = fraction_tags()
    y = grid.y(5)
    inside = Line(grid.x(3), grid.y(2), grid.x(3), grid.y(3), 0.25, tag=border_tag(Role.TABLE_BORDER, "hair"))
    page = raw_page(Text(grid.x(10), y - 6, "a", STYLE, tag=numerator),
                    Line(grid.x(10), y - 4, grid.x(11), y - 4, 0.5, tag=bar),
                    Text(grid.x(10), y + 8, "b", STYLE, tag=denominator),
                    Text(grid.x(2), grid.y(2) + 10, "abcd", STYLE, tag=paragraph_tag()),  # merges columns 2-3
                    inside, profile=TURNED_GRID)
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    sheet = workbook.sheets[0]
    tops = [math.fsum(sheet.row_heights_pt[:row]) for row in range(len(sheet.row_heights_pt) + 1)]
    lefts = [math.fsum(sheet.column_widths_pt[:col]) for col in range(len(sheet.column_widths_pt) + 1)]
    assert {round(size / 0.75) for size in sheet.row_heights_pt} == {19, 20}
    assert {round(size / 0.75) for size in sheet.column_widths_pt} == {14, 15}
    lines = [drawing for drawing in sheet.drawings if isinstance(drawing, xl.Line)]
    bar_line = next(line for line in lines if line.y1 == line.y2)
    assert (bar_line.x1, bar_line.y1) == pytest.approx((lefts[10], tops[5]))  # above the denominator's row
    edge = next(line for line in lines if line.x1 == line.x2)  # the border inside "abcd"'s merge
    assert (edge.x1, edge.y1, edge.y2) == pytest.approx((lefts[3], tops[2], tops[3]))
    assert [item.kind for item in result.fallbacks] == ["border_in_merge"]


def test_a_denominator_pushed_past_the_bottom_is_an_error():
    numerator, denominator, bar, _ = fraction_tags()
    y = baseline(46)
    page = raw_page(Text(x_of(2), y - 6, "a", STYLE, tag=numerator),
                    Line(x_of(2), y - 4, x_of(3), y - 4, 0.5, tag=bar),
                    Text(x_of(2), y - 1, "b", STYLE, tag=denominator))
    with pytest.raises(XlsxError, match="bottom edge"):
        map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)


def test_diamonds_and_circles_are_shapes():
    from gridlayoutpdf.figures import TextDiamond
    from gridlayoutpdf.scene import Circle

    sheet = first_sheet(mapped(Figure(10, 6, TextDiamond(2, 2, 100, 60, ("a", "b")), Circle(130, 40, 10))))
    (group,) = sheet.drawings
    diamond, circle = group.children
    assert isinstance(diamond, xl.Diamond) and len(diamond.text.paragraphs) == 2
    assert isinstance(circle, xl.Ellipse)


@pytest.mark.parametrize("line", [
    Line(OX, OY + PY, OX + 30, OY + PY, 0.5, tag=Tag(Role.RULE, (), {"style": "bogus"})),
    Line(OX, OY, OX + 30, OY + 30, 0.5, tag=Tag(Role.RULE, (), {"style": "thin"})),
])
def test_a_border_without_a_style_or_slanted_is_a_recorded_shape(line):
    workbook, result = map_reports([("R", (raw_page(line),))], FONTS, cell_padding_pt=padding)
    assert not workbook.sheets[0].borders and [item.kind for item in result.fallbacks] == ["border_off_grid"]


def test_many_fractions_map_in_linear_time():
    import time

    group = EquationGroup(*(Equation(f"x{index}", Fraction("a", "b")) for index in range(15)),
                          anchors=EquationAnchors(4))
    pages = doc(*(group,) * 40).layout()
    start = time.perf_counter()
    map_reports([("R", pages)], FONTS, cell_padding_pt=padding)
    assert time.perf_counter() - start < 10


# 17. Review round 4 ---------------------------------------------------------------------------------

def fraction_page(row, script=False):
    numerator, denominator, bar, _ = fraction_tags()
    if script:
        scope = Scope(Kind.SCRIPT, "superscript")
        numerator, denominator, bar = (Tag(tag.role, (*tag.path[:3], scope, *tag.path[3:]), tag.data)
                                       for tag in (numerator, denominator, bar))
    y = baseline(row)
    return raw_page(Text(x_of(4), y - 6, "a", STYLE, tag=numerator),
                    Line(x_of(4), y - 4, x_of(5), y - 4, 0.5, tag=bar),
                    Text(x_of(4), y + 8, "b", STYLE, tag=denominator))


def test_the_same_fraction_on_two_pages_of_a_sheet_stays_on_its_own_page():
    workbook, result = map_reports([("R", (fraction_page(2), fraction_page(30)))], FONTS, cell_padding_pt=padding)
    sheet = workbook.sheets[0]
    assert values(sheet) == {(2, 4): "a", (3, 4): "b", (47 + 30, 4): "a", (47 + 31, 4): "b"}
    assert not result.shifts
    bars = sorted(drawing.y1 for drawing in sheet.drawings)
    tops = [math.fsum(sheet.row_heights_pt[:row]) for row in range(len(sheet.row_heights_pt) + 1)]
    assert bars == pytest.approx([tops[3], tops[47 + 31]])


def test_a_fraction_in_a_script_on_each_of_two_pages_is_recorded_twice():
    _, result = map_reports([("R", (fraction_page(2, True), fraction_page(30, True)))], FONTS,
                            cell_padding_pt=padding)
    assert [(item.kind, item.page) for item in result.fallbacks] == [("fraction_in_script", 1),
                                                                     ("fraction_in_script", 2)]


def test_two_borders_giving_way_to_one_are_two_conflicts():
    long = Line(OX, OY + PY, x_of(6), OY + PY, 1.5, tag=border_tag(Role.RULE, "thick"))
    dashed = [Line(x_of(col), OY + PY, x_of(col + 1), OY + PY, 0.5, dash=(3, 1.5),
                   tag=border_tag(Role.RULE, "dashed")) for col in (0, 3)]
    _, result = map_reports([("R", (raw_page(long, *dashed),))], FONTS, cell_padding_pt=padding)
    assert [item.kind for item in result.fallbacks] == ["border_conflict"] * 2


def test_border_colors_are_compared_as_excel_writes_them():
    lines = (Line(OX, OY + PY, x_of(2), OY + PY, 0.5, (0.2, 0.2, 0.2), tag=border_tag(Role.RULE)),
             Line(OX, OY + PY, x_of(2), OY + PY, 0.5, (0.2000001, 0.2, 0.2), tag=border_tag(Role.RULE)))
    _, result = map_reports([("R", (raw_page(*lines),))], FONTS, cell_padding_pt=padding)
    assert not result.fallbacks


# 18. Review round 5 ---------------------------------------------------------------------------------

def test_nested_fractions_move_from_the_inside_out():
    # the inner numerator a, inner denominator b and outer denominator c all start on row 3
    line = (Scope(Kind.EQUATION, index=0), Scope(Kind.EQUATION_LINE, index=0), Scope(Kind.EQUATION_PART, "rhs"))
    outer_num = (*line, Scope(Kind.FRACTION, "numerator"))
    a = Tag(Role.MATH, (*outer_num, Scope(Kind.FRACTION, "numerator")))
    b = Tag(Role.MATH, (*outer_num, Scope(Kind.FRACTION, "denominator")))
    c = Tag(Role.MATH, (*line, Scope(Kind.FRACTION, "denominator")))
    page = raw_page(*(Text(x_of(2), baseline(3), text, STYLE, tag=tag) for text, tag in (("a", a), ("b", b), ("c", c))))
    _, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    assert [(shift.text, shift.from_ref, shift.to_ref) for shift in result.shifts] == [
        ("b", "C4", "C5"), ("c", "C4", "C6")]


def test_a_small_nested_fraction_stacks_in_three_rows():
    sheet = first_sheet(mapped(MathBlock(Fraction(Fraction("a", "b"), "c"), style=TextStyle(size=4))))
    a, b, c = (text_cell(sheet, value) for value in "abc")
    assert (b.row, c.row) == (a.row + 1, a.row + 2) and c.col == a.col


def equation_page(row):
    numerator, denominator, bar, lhs = fraction_tags()
    y = baseline(row)
    return raw_page(Text(x_of(1), y, "x", STYLE, tag=lhs),
                    Text(x_of(4), y - 6, "a", STYLE, tag=numerator),
                    Line(x_of(4), y - 4, x_of(5), y - 4, 0.5, tag=bar),
                    Text(x_of(4), y + 8, "b", STYLE, tag=denominator))


def test_equation_lines_with_the_same_tags_on_two_pages_merge_on_their_own_page():
    workbook, _ = map_reports([("R", (equation_page(2), equation_page(30)))], FONTS, cell_padding_pt=padding)
    sheet = workbook.sheets[0]
    first, second = (merge_of(sheet, row, 1) for row in (2, 47 + 30))
    assert (first.first_row, first.last_row) == (2, 3) and (second.first_row, second.last_row) == (77, 78)


def test_split_grid_terms_are_recorded_on_each_page():
    group = EquationGroup(Equation("A", Row("2", Fraction("a", "b"), TermStop("left"), "3", gap=0)),
                          anchors=EquationAnchors(2), term_columns="grid")
    from gridlayoutpdf.blocks import PageBreak

    _, result = map_reports([("R", doc(group, PageBreak(), group).layout())], FONTS, cell_padding_pt=padding)
    assert [item.page for item in result.fallbacks if item.kind == "term_off_grid"] == [1, 2]


@pytest.mark.parametrize(("first", "second"), [(0, 1), (-1, 0), (0, True)])
def test_a_stroke_number_that_does_not_fit_its_style_names_its_own_line(first, second):
    long = Line(OX, OY + PY, x_of(6), OY + PY, 1.5, tag=border_tag(Role.RULE, "thick"))
    # a single dashed line has only stroke 0: another stroke number must not make it a neighbour's line
    dashed = [Line(x_of(col), OY + PY, x_of(col + 1), OY + PY, 0.5, dash=(3, 1.5),
                   tag=border_tag(Role.RULE, "dashed", stroke=stroke)) for col, stroke in ((0, first), (3, second))]
    _, result = map_reports([("R", (raw_page(long, *dashed),))], FONTS, cell_padding_pt=padding)
    assert [item.kind for item in result.fallbacks] == ["border_conflict"] * 2


def test_a_fraction_bar_whose_tag_names_no_fraction_is_an_error():
    bar = Line(x_of(2), baseline(3), x_of(3), baseline(3), 0.5, tag=Tag(Role.FRACTION_BAR))
    with pytest.raises(XlsxError, match="fraction bar"):
        map_reports([("R", (raw_page(bar),))], FONTS, cell_padding_pt=padding)


# 19. Step 10 comparison (lead, 2026-10-20) ------------------------------------------------------------

@pytest.mark.parametrize(("valign", "expected"), [("top", "top"), ("middle", "center"), ("bottom", "bottom")])
def test_a_one_line_table_cell_fills_its_whole_cell_with_its_alignment(valign, expected):
    table = Table([4], [[Cell("abc", valign=valign, align="right")]], row_heights=(2,))
    sheet = first_sheet(mapped(table))
    cell = text_cell(sheet, "abc")
    assert merge_of(sheet, cell.row, cell.col) == xl.CellRange(0, 0, 1, 3)
    assert (cell.h_align, cell.v_align) == ("right", expected)


def test_a_spanning_table_cell_fills_its_rows_and_columns():
    table = Table([2, 2], [[Cell("ab", colspan=2, rowspan=2)], []], row_heights=(1, 1))
    sheet = first_sheet(mapped(table))
    cell = text_cell(sheet, "ab")
    assert merge_of(sheet, cell.row, cell.col) == xl.CellRange(0, 0, 1, 3)


def test_a_table_cell_of_several_lines_keeps_a_cell_per_line():
    sheet = first_sheet(mapped(Table([3], [[Cell("ab\ncd")]])))
    first, second = text_cell(sheet, "ab"), text_cell(sheet, "cd")
    assert second.row == first.row + 1
    assert all((merge_of(sheet, cell.row, cell.col) or xl.CellRange(cell.row, 0, cell.row, 1)).last_row == cell.row
               for cell in (first, second))


def test_a_table_cell_inside_columns_fills_its_cell():
    table = Table([3], [[Cell("ab")]], row_heights=(2,))
    pages = doc(Columns(Paragraph("x"), table, col_spans=(2, 3))).layout()
    workbook, _ = map_reports([("R", pages)], FONTS, cell_padding_pt=padding)
    left = round(pages[0].metadata["placements"][0]["parts"][0]["metadata"]["children"][1]["x_pt"] / PX)
    cell = text_cell(workbook.sheets[0], "ab")
    assert merge_of(workbook.sheets[0], cell.row, cell.col) == xl.CellRange(0, left, 1, left + 2)


def test_a_one_line_cell_of_a_turned_table_fills_its_turned_cell():
    table = Table([2], [[Cell("ab", valign="top")]], row_heights=(3,))
    block = QuarterTurn(table, turns=1, local_grid=GridSpec(0, 0, PY, PX, 4, 3))
    sheet = first_sheet(mapped(*(Paragraph(f"p{index}") for index in range(3)), block))
    cell = text_cell(sheet, "ab")
    merge = merge_of(sheet, cell.row, cell.col)
    # the cell is 2 local columns by 3 local rows: 3 page columns by 2 page rows, read upwards
    assert (merge.last_col - merge.first_col + 1, merge.last_row - merge.first_row + 1) == (3, 2)
    assert cell.rotation == 90 and cell.h_align == "left"  # the local top is the page's left


def test_text_grows_into_free_cells_until_excel_has_room_for_it():
    # "abcdefgh" (8 glyphs of 8 px = 48 pt) touches 3 columns (47.25 pt) and needs 48 + 2 x 2.4 pt
    sheet = first_sheet(mapped(Paragraph("abcdefgh")))
    assert merge_of(sheet, 0, 0) == xl.CellRange(0, 0, 0, 3)


def test_text_that_cannot_grow_enough_is_recorded_as_clipped():
    # 60 pt of glyphs fit the 4-column cell in the PDF; Excel needs 60 + 4.8 pt of its 63 pt, borders both sides
    table = Table([4], [[Cell("abcdefghij", padding=1)]])
    workbook, result = mapped(table)
    assert [item.kind for item in result.texts_clipped] == ["text_clipped"]
    assert not result.fallbacks
    mapped(table, strict=True)  # recorded only, not refused


def wide_relation_page(text="===="):
    relation = Tag(Role.MATH, (Scope(Kind.EQUATION, index=0), Scope(Kind.EQUATION_LINE, index=0),
                               Scope(Kind.EQUATION_PART, "relation")), {"atom": "relation"})
    # 24 pt centred on column boundary 5
    return raw_page(Text(x_of(5) - 12, baseline(0), text, STYLE, tag=relation))


def test_a_wide_relation_stays_centred_on_its_anchor():
    workbook, _ = map_reports([("R", (wide_relation_page(),))], FONTS, cell_padding_pt=padding)
    cell = text_cell(workbook.sheets[0], "====")
    assert merge_of(workbook.sheets[0], cell.row, cell.col) == xl.CellRange(0, 4, 0, 5)


def test_a_wide_relation_keeps_off_the_left_hand_side():
    lhs = Tag(Role.MATH, (Scope(Kind.EQUATION, index=0), Scope(Kind.EQUATION_LINE, index=0),
                          Scope(Kind.EQUATION_PART, "lhs")), {"atom": "literal"})
    page = raw_page(*wide_relation_page().commands, Text(x_of(4, 1), baseline(0), "x", STYLE, tag=lhs))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    cell = text_cell(workbook.sheets[0], "====")
    assert merge_of(workbook.sheets[0], cell.row, cell.col) == xl.CellRange(0, 5, 0, 6) and not result.shifts


def test_the_left_hand_side_reaches_the_column_before_its_relation_s_anchor():
    # "B′ ：" in WebPier's definitions: the prime's ink sits at the left of its advance, so the ink of the
    # right-aligned left-hand side ends a column short of the anchor; its cells still reach the anchor (api.md).
    lhs = Tag(Role.MATH, (Scope(Kind.EQUATION, index=0), Scope(Kind.EQUATION_LINE, index=0),
                          Scope(Kind.EQUATION_PART, "lhs")), {"atom": "literal"})
    relation = Tag(Role.MATH, (Scope(Kind.EQUATION, index=0), Scope(Kind.EQUATION_LINE, index=0),
                               Scope(Kind.EQUATION_PART, "relation")), {"atom": "relation"})
    page = raw_page(Text(x_of(3, 2), baseline(0), "x", STYLE, tag=lhs),
                    Text(x_of(5) - 3, baseline(0), "=", STYLE, tag=relation))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    sheet = workbook.sheets[0]
    cell = text_cell(sheet, "x")
    assert merge_of(sheet, cell.row, cell.col) == xl.CellRange(0, 3, 0, 4) and cell.h_align == "right"
    assert text_cell(sheet, "=").col == 5 and not result.shifts


def test_only_the_last_piece_of_a_split_left_hand_side_reaches_the_anchor():
    # atoms a cell apart are two pieces (the cantilever's "T  L1"): the first keeps its cells
    lhs = Tag(Role.MATH, (Scope(Kind.EQUATION, index=0), Scope(Kind.EQUATION_LINE, index=0),
                          Scope(Kind.EQUATION_PART, "lhs")), {"atom": "literal"})
    relation = Tag(Role.MATH, (Scope(Kind.EQUATION, index=0), Scope(Kind.EQUATION_LINE, index=0),
                               Scope(Kind.EQUATION_PART, "relation")), {"atom": "relation"})
    page = raw_page(Text(x_of(1, 2), baseline(0), "a", STYLE, tag=lhs), Text(x_of(3, 2), baseline(0), "b", STYLE, tag=lhs),
                    Text(x_of(6) - 3, baseline(0), "=", STYLE, tag=relation))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    sheet = workbook.sheets[0]
    first, second = text_cell(sheet, "a"), text_cell(sheet, "b")
    assert merge_of(sheet, first.row, first.col) is None and first.col == 1
    assert merge_of(sheet, second.row, second.col) == xl.CellRange(0, 3, 0, 5) and not result.shifts


@pytest.mark.parametrize(("gap", "joined"), [(1.5, "ab"), (4.0, "a b")])
def test_atoms_join_with_a_space_only_where_the_gap_is_half_a_space_or_more(gap, joined):
    # FixedFonts' space is 6 pt at 10 pt: a gap of 3 pt or more becomes a space
    tag = Tag(Role.MATH, (Scope(Kind.MATH_BLOCK),), {"atom": "literal"})
    page = raw_page(Text(x_of(2), baseline(0), "a", STYLE, tag=tag), Text(x_of(2, 6 + gap), baseline(0), "b", STYLE, tag=tag))
    workbook, _ = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    assert values(workbook.sheets[0]) == {(0, 2): joined}


@pytest.mark.parametrize(("offset", "expected"), [(11, "center"), (7.5, "top"), (10, "center"), (13, "center")])
def test_a_line_takes_the_vertical_alignment_nearest_its_place_in_its_row(offset, expected):
    # FixedFonts' line box is 8 pt above and 2 pt below the baseline; the row is 15.48 pt
    page = raw_page(Text(x_of(2), OY + 3 * PY + offset, "ab", STYLE, tag=paragraph_tag()))
    workbook, _ = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    assert text_cell(workbook.sheets[0], "ab").v_align == expected


def test_a_turned_line_overhanging_its_column_stays_in_the_middle():
    # On 10.8 pt columns (the cantilever's turned tables) a line reading upwards has its box 8 pt left of the
    # baseline and 2 pt right of it: from 4 pt left of column 3 to 6 pt into it. Its right (the text's
    # bottom) is nearer Excel's bottom-aligned place than the middle, by less than CROSS_ALIGN_MARGIN_PT:
    # it stays in the middle, where Excel cuts no glyph.
    grid = TURNED_GRID.grid
    tag = Tag(Role.TEXT, (Scope(Kind.QUARTER_TURN, data={"turns": 1, "block_id": "t", "local_grid": {
        "origin_x": 0, "origin_y": 0, "pitch_x": grid.pitch_y, "pitch_y": grid.pitch_x, "columns": 4, "rows": 1}}),
        Scope(Kind.PARAGRAPH, index=0, data={"align": "left"})))
    page = raw_page(Text(grid.x(3) + 4, grid.y(10), "ab", STYLE, 1, tag=tag), profile=TURNED_GRID)
    workbook, _ = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    cell = text_cell(workbook.sheets[0], "ab")
    assert cell.col == 3 and cell.h_align == "center"


# 20. Review of the comparison fixes -------------------------------------------------------------------

def test_a_script_over_a_fraction_maps():
    workbook, _ = mapped(MathBlock(Script(Fraction("a", "b"), "2")))
    (two,) = [cell for cell in workbook.sheets[0].cells if isinstance(cell.value, xl.RichText)]
    assert two.value.runs[0].vert_align == "superscript"


def test_a_wide_relation_leaves_the_cells_of_a_fraction_in_the_left_hand_side():
    # the relation would straddle its anchor over columns 4-5; the denominator below takes columns 2-4 later
    group = EquationGroup(Equation(Fraction("M", "ZZZZ"), "1.25", relation="==="), anchors=EquationAnchors(5))
    workbook, result = mapped(group)
    denominator = text_cell(workbook.sheets[0], "ZZZZ")
    assert denominator.col == 2 and "ZZZZ" not in {shift.text for shift in result.shifts}


def test_one_line_equation_pieces_stay_centred_in_their_rows():
    group = EquationGroup(Equation("A", "3 + 9.15", result="12.15", unit="m2"),
                          anchors=EquationAnchors(3, 12, 17, 18), style=TextStyle(size=12))
    sheet = first_sheet(mapped(group))
    assert {cell.v_align for cell in sheet.cells} == {"center"}


def test_a_top_padding_far_from_excels_keeps_the_text_in_its_own_cells():
    from gridlayoutpdf import Insets

    table = Table([4], [[Cell("ab", valign="top", padding=Insets(left=3, right=3, top=20, bottom=3))]],
                  row_heights=(3,))
    sheet = first_sheet(mapped(table))
    cell = text_cell(sheet, "ab")
    merge = merge_of(sheet, cell.row, cell.col)
    assert merge is None or merge.first_row == merge.last_row  # not the whole three-row frame


@pytest.mark.parametrize(("valign", "insets"), [("bottom", (3, 3, 3, 20)), ("middle", (3, 3, 40, 3))])
def test_a_vertical_padding_far_from_excels_keeps_the_text_in_its_own_rows(valign, insets):
    from gridlayoutpdf import Insets

    left, right, top, bottom = insets
    table = Table([4], [[Cell("ab", valign=valign, padding=Insets(left=left, right=right, top=top, bottom=bottom))]],
                  row_heights=(4,))
    sheet = first_sheet(mapped(table))
    cell = text_cell(sheet, "ab")
    merge = merge_of(sheet, cell.row, cell.col)
    assert merge is None or merge.first_row == merge.last_row  # not the whole four-row frame


def test_a_centred_cell_with_even_padding_fills_its_cell():
    table = Table([6], [[Cell("ab", align="center")]], row_heights=(2,))
    sheet = first_sheet(mapped(table))
    cell = text_cell(sheet, "ab")
    assert merge_of(sheet, cell.row, cell.col) == xl.CellRange(0, 0, 1, 5) and cell.h_align == "center"


def test_a_centred_cell_with_uneven_side_padding_keeps_the_text_in_its_own_cells():
    from gridlayoutpdf import Insets

    table = Table([6], [[Cell("ab", align="center", padding=Insets(left=40, right=3, top=3, bottom=3))]],
                  row_heights=(2,))
    sheet = first_sheet(mapped(table))
    cell = text_cell(sheet, "ab")
    merge = merge_of(sheet, cell.row, cell.col)
    assert merge is None or (merge.first_col, merge.last_col) != (0, 5)


def test_a_cell_padding_far_from_excels_keeps_the_text_in_its_own_cells():
    from gridlayoutpdf import Insets

    table = Table([4], [[Cell("ab", align="left", padding=Insets(left=24, right=3, top=3, bottom=3))]],
                  row_heights=(2,))
    sheet = first_sheet(mapped(table))
    cell = text_cell(sheet, "ab")
    assert cell.col == 1  # the text starts 24 pt into the cell, in its second column


def test_a_number_kept_as_text_is_not_also_recorded_as_clipped():
    pages = doc(Table([2], [[Cell("12.5")]])).layout()
    _, result = map_reports([("R", pages)], FONTS, cell_padding_pt=lambda size: 3.8)
    assert len(result.numbers_as_text) == 1 and not result.texts_clipped


def test_a_one_line_cell_of_a_clockwise_table_takes_its_top_to_the_right():
    table = Table([2], [[Cell("ab", valign="top")]], row_heights=(3,))
    block = QuarterTurn(table, turns=-1, local_grid=GridSpec(0, 0, PY, PX, 4, 3))
    cell = text_cell(first_sheet(mapped(block)), "ab")
    assert cell.rotation == 270 and cell.h_align == "right"


def test_a_table_cell_in_an_indented_stack_fills_its_cell():
    from gridlayoutpdf.containers import KeepTogether, Stack

    table = Table([3], [[Cell("ab")]], row_heights=(2,))
    sheet = first_sheet(mapped(KeepTogether(Stack(Paragraph("x"), Indent(table, 2 * PX)))))
    cell = text_cell(sheet, "ab")
    assert merge_of(sheet, cell.row, cell.col) == xl.CellRange(1, 2, 2, 4)


def test_the_cells_of_a_split_table_fill_their_cells_on_the_next_page_too():
    rows = [[Cell("head")]] + [[Cell(f"r{index}")] for index in range(30)]
    table = Table([3], rows, header_rows=1, row_heights=(2,) * 31)
    sheet = first_sheet(mapped(table))
    for value in ("r0", "r29"):
        cell = text_cell(sheet, value)
        merge = merge_of(sheet, cell.row, cell.col)
        assert (merge.last_row - merge.first_row, merge.first_col, merge.last_col) == (1, 0, 2)
    heads = [cell for cell in sheet.cells if cell.value == "head"]
    assert len(heads) == 2 and all(merge_of(sheet, cell.row, cell.col).last_col == 2 for cell in heads)


def test_a_table_in_a_fractional_indent_keeps_its_own_cells():
    table = Table([3], [[Cell("ab")]], row_heights=(2,))
    sheet = first_sheet(mapped(Indent(table, 5)))
    cell = text_cell(sheet, "ab")
    merge = merge_of(sheet, cell.row, cell.col)
    assert merge is None or merge.last_row == merge.first_row


def test_an_indented_stack_s_table_cell_fills_its_cell():
    table = Table([3], [[Cell("ab")]], row_heights=(2,))
    from gridlayoutpdf.containers import KeepTogether, Stack

    sheet = first_sheet(mapped(Indent(KeepTogether(Stack(Paragraph("x"), table)), 2 * PX)))
    cell = text_cell(sheet, "ab")
    assert merge_of(sheet, cell.row, cell.col) == xl.CellRange(1, 2, 2, 4)


def test_note_records_stay_in_their_own_lists():
    from gridlayoutpdf.renderers.xlsx import XlsxFallback

    number, clipped = XlsxFallback("R", 1, "number_as_text", ""), XlsxFallback("R", 1, "text_clipped", "")
    for values in ({"texts_clipped": (number,)}, {"numbers_as_text": (clipped,)}, {"fallbacks": (clipped,)}):
        with pytest.raises(ValueError):
            XlsxResult(None, (), **values)
    with pytest.raises(ValueError):
        XlsxStrictError("m", (), (clipped,))


def test_a_relation_equally_near_two_places_takes_the_right_one():
    # On a grid of 16 pt from the origin (exact in binary), 3 columns of "=" centred on boundary 5 straddle it
    # either way: columns 3-5 or 4-6 are equally near
    profile = PageProfile("exact", 595, 842, GridSpec(0, 0, 16, 16, 30, 47))
    relation = Tag(Role.MATH, (Scope(Kind.EQUATION, index=0), Scope(Kind.EQUATION_LINE, index=0),
                               Scope(Kind.EQUATION_PART, "relation")), {"atom": "relation"})
    page = raw_page(Text(5 * 16 - 18, 11, "======", STYLE, tag=relation), profile=profile)  # 36 pt: 3 columns
    workbook, _ = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    cell = text_cell(workbook.sheets[0], "======")
    assert merge_of(workbook.sheets[0], cell.row, cell.col).first_col == 4


def test_a_number_in_mixed_fonts_grows_into_a_free_cell():
    other = TextStyle(font="IPAexゴシック", size=10)
    page = raw_page(Text(x_of(0), baseline(0), "123456", STYLE, tag=paragraph_tag()),
                    Text(x_of(0, 36), baseline(0), ".5", other, tag=paragraph_tag()))
    workbook, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    assert merge_of(workbook.sheets[0], 0, 0) == xl.CellRange(0, 0, 0, 3)  # 48 + 4.8 pt need a fourth column
    assert len(result.numbers_as_text) == 1 and "cut" not in result.numbers_as_text[0].detail


def test_a_number_in_mixed_fonts_is_recorded_once():
    other = TextStyle(font="IPAexゴシック", size=10)
    page = raw_page(Text(x_of(0), baseline(0), "123456", STYLE, tag=paragraph_tag()),
                    Text(x_of(0, 36), baseline(0), ".5", other, tag=paragraph_tag()),
                    Text(x_of(3, 2), baseline(0), "x", STYLE, tag=paragraph_tag(line=1)))  # the next cell is taken
    _, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=padding)
    assert len(result.numbers_as_text) == 1 and not result.texts_clipped
    assert "may be cut" in result.numbers_as_text[0].detail


def test_script_runs_are_measured_at_their_own_size():
    # "MMMM" + superscript "MMMM" (6.5 pt): 32 + 24 px = 42 pt, + 2 x 1 pt fits three columns (47.25 pt);
    # measured at the base size the superscript would take 32 px and need 50 pt
    tag = Tag(Role.MATH, (Scope(Kind.MATH_BLOCK),), {"atom": "literal"})
    script = Tag(Role.MATH, (Scope(Kind.MATH_BLOCK), Scope(Kind.SCRIPT, "superscript")), {"atom": "literal"})
    page = raw_page(Text(x_of(0), baseline(0), "MMMM", STYLE, tag=tag),
                    Text(x_of(0, 25), baseline(0) - 4, "MMMM", TextStyle(size=6.5), tag=script),
                    Text(x_of(3, 1), baseline(0), "x", STYLE, tag=paragraph_tag()))  # column 3 is taken
    _, result = map_reports([("R", (page,))], FONTS, cell_padding_pt=lambda size: 1.0)
    assert not result.texts_clipped


def test_records_are_exported_from_the_package():
    import gridlayoutpdf

    for name in ("render_xlsx", "XlsxResult", "XlsxSheet", "XlsxShift", "XlsxFallback", "XlsxError",
                 "XlsxStrictError"):
        assert name in gridlayoutpdf.__all__ and hasattr(gridlayoutpdf, name)
