"""Cells, rich text, formats, merges and borders read back with openpyxl."""

import io
import re
import zipfile

import openpyxl
import pytest
from openpyxl.cell.rich_text import CellRichText

from gridlayoutpdf.renderers.xlsx._model import (
    BORDER_STYLES,
    CellRange,
    CellValue,
    EdgeBorder,
    Font,
    RichText,
    Run,
    Sheet,
    Workbook,
    range_edges,
)
from gridlayoutpdf.renderers.xlsx._package import write_workbook

MINCHO = Font("IPAex明朝", 10.5)


def _book(cells=(), merges=(), borders=(), columns=6, rows=6):
    return Workbook((Sheet("S", (16.08,) * columns, (15.48,) * rows, cells=cells, merges=merges, borders=borders),))


def _read(workbook, rich_text=False):
    return openpyxl.load_workbook(io.BytesIO(write_workbook(workbook)), rich_text=rich_text)["S"]


def _part(workbook, name):
    with zipfile.ZipFile(io.BytesIO(write_workbook(workbook))) as archive:
        return archive.read(name).decode("utf-8")


def test_strings_numbers_and_blank_formatted_cells():
    sheet = _read(_book((
        CellValue(0, 0, "Ｌ＝"), CellValue(0, 1, -2.0), CellValue(1, 1, 12345.678),
        CellValue(2, 2, None, fill=(1, 1, 0)), CellValue(3, 0, "  前後の空白  "), CellValue(3, 1, "1行目\n2行目"),
    )))
    assert sheet["A1"].value == "Ｌ＝"
    assert sheet["B1"].value == -2 and sheet["B1"].data_type == "n"
    assert sheet["B2"].value == 12345.678
    assert sheet["C3"].value is None
    assert sheet["C3"].fill.fgColor.rgb == "FFFFFF00" and sheet["C3"].fill.fill_type == "solid"
    assert sheet["A4"].value == "  前後の空白  "
    assert sheet["B4"].value == "1行目\n2行目"


@pytest.mark.parametrize("number_format", ['0.000;"−"0.000', "#,##0.0", '"+"0.00', "0.00", "General", "@"])
def test_number_formats_are_kept(number_format):
    sheet = _read(_book((CellValue(0, 0, -2.0, number_format=number_format),)))
    assert sheet["A1"].number_format == number_format


def test_number_format_ids_are_deduplicated_in_order_of_first_appearance():
    styles = _part(_book((
        CellValue(0, 0, 1.0, number_format="#,##0.0"), CellValue(0, 1, 2.0, number_format='"+"0.00'),
        CellValue(1, 0, 3.0, number_format="#,##0.0"), CellValue(1, 1, 4.0, number_format="0.00"),
    )), "xl/styles.xml")
    assert re.findall(r'<numFmt numFmtId="(\d+)" formatCode="([^"]*)"/>', styles) == [
        ("164", "#,##0.0"), ("165", "&quot;+&quot;0.00"),
    ]


def test_rich_text_runs_with_superscript_and_subscript():
    rich = RichText((
        Run("σ", MINCHO), Run("ck", Font("IPAex明朝", 7), vert_align="subscript"),
        # Only checks writing: Step 8 gives scripts the base size because Excel shrinks them itself.
        Run("2", Font("IPAex明朝", 7), vert_align="superscript"), Run(" N/mm", MINCHO),
    ))
    cell = _read(_book((CellValue(0, 0, rich, font=MINCHO),)), rich_text=True)["A1"]
    assert isinstance(cell.value, CellRichText)
    assert str(cell.value) == "σck2 N/mm"
    runs = list(cell.value)
    assert [run.text for run in runs] == ["σ", "ck", "2", " N/mm"]
    assert [run.font.vertAlign for run in runs] == [None, "subscript", "superscript", None]
    assert [run.font.sz for run in runs] == [10.5, 7, 7, 10.5]
    assert {run.font.rFont for run in runs} == {"IPAex明朝"}


def test_rich_text_run_without_font_uses_the_cell_font():
    rich = RichText((Run("x"), Run("2", vert_align="superscript")))
    runs = list(_read(_book((CellValue(0, 0, rich, font=MINCHO),)), rich_text=True)["A1"].value)
    assert [(run.font.rFont, run.font.sz) for run in runs] == [("IPAex明朝", 10.5)] * 2


@pytest.mark.parametrize(("h_align", "v_align", "rotation", "excel"), [
    ("left", "top", 0, 0), ("center", "center", 90, 90), ("right", "bottom", 270, 180),
])
def test_alignment_and_counterclockwise_rotation_become_excel_text_rotation(h_align, v_align, rotation, excel):
    # Excel's textRotation 90 is 90° counterclockwise and 180 is 90° clockwise (270° counterclockwise).
    cell = _read(_book((CellValue(0, 0, "x", h_align=h_align, v_align=v_align, rotation=rotation),)))["A1"]
    assert (cell.alignment.horizontal, cell.alignment.vertical) == (h_align, v_align)
    assert cell.alignment.textRotation == excel


def test_cell_font_name_size_color_bold_italic():
    font = Font("IPAexゴシック", 9.5, color=(1, 0, 0.5), bold=True, italic=True)
    cell = _read(_book((CellValue(0, 0, "x", font=font),)))["A1"]
    assert (cell.font.name, cell.font.sz, cell.font.b, cell.font.i) == ("IPAexゴシック", 9.5, True, True)
    assert cell.font.color.rgb == "FFFF0080"


def test_styles_are_deduplicated_and_numbered_in_order_of_first_appearance():
    first = Font("IPAex明朝", 10)
    second = Font("IPAexゴシック", 10)
    book = _book((
        CellValue(0, 0, "a", font=first), CellValue(0, 1, "b", font=second),
        CellValue(1, 0, "c", font=first), CellValue(1, 1, "d", font=second, h_align="right"),
    ))
    sheet_xml = _part(book, "xl/worksheets/sheet1.xml")
    assert re.findall(r'<c r="(\w+)" s="(\d+)"', sheet_xml) == [("A1", "1"), ("B1", "2"), ("A2", "1"), ("B2", "3")]
    styles = _part(book, "xl/styles.xml")
    assert styles.count("<font>") == 3  # default + two cell fonts


def test_shared_strings_are_deduplicated():
    book = _book((CellValue(0, 0, "同じ"), CellValue(0, 1, "同じ"), CellValue(1, 0, "別")))
    strings = _part(book, "xl/sharedStrings.xml")
    assert strings.count("<si>") == 2
    assert 'count="3" uniqueCount="2"' in strings


def test_merges_are_written():
    sheet = _read(_book((CellValue(0, 0, "結合"),), merges=(CellRange(0, 0, 1, 2), CellRange(3, 3, 5, 3))))
    assert sorted(str(item) for item in sheet.merged_cells.ranges) == ["A1:C2", "D4:D6"]


@pytest.mark.parametrize("style", BORDER_STYLES)
def test_every_excel_border_style_is_written(style):
    cell = _read(_book(borders=(EdgeBorder(1, 1, "top", style, color=(0, 0, 1)),)))["B2"]
    assert cell.border.top.style == style
    assert cell.border.top.color.rgb == "FF0000FF"
    assert cell.border.bottom.style is None


def test_border_styles_are_the_thirteen_excel_styles():
    assert BORDER_STYLES == (
        "hair", "thin", "medium", "thick", "double", "dotted", "dashed", "dashDot", "dashDotDot",
        "mediumDashed", "mediumDashDot", "mediumDashDotDot", "slantDashDot",
    )


def test_each_side_goes_to_the_named_cell_edge():
    sheet = _read(_book(borders=(
        EdgeBorder(1, 1, "left", "thin"), EdgeBorder(1, 1, "right", "medium"),
        EdgeBorder(1, 1, "top", "thick"), EdgeBorder(1, 1, "bottom", "double"),
    )))
    border = sheet["B2"].border
    assert (border.left.style, border.right.style, border.top.style, border.bottom.style) == (
        "thin", "medium", "thick", "double")


def test_merged_range_sides_are_written_on_every_cell_along_the_side():
    merge = CellRange(1, 1, 2, 3)
    edges = range_edges(merge, "top", "thin") + range_edges(merge, "bottom", "medium") \
        + range_edges(merge, "left", "hair") + range_edges(merge, "right", "double")
    sheet = _read(_book((CellValue(1, 1, "結合"),), merges=(merge,), borders=edges))
    assert [sheet.cell(2, col).border.top.style for col in (2, 3, 4)] == ["thin"] * 3
    assert [sheet.cell(3, col).border.bottom.style for col in (2, 3, 4)] == ["medium"] * 3
    assert [sheet.cell(row, 2).border.left.style for row in (2, 3)] == ["hair"] * 2
    assert [sheet.cell(row, 4).border.right.style for row in (2, 3)] == ["double"] * 2


def test_range_edges_lists_one_edge_per_cell_along_the_side():
    assert range_edges(CellRange(1, 1, 2, 3), "left", "thin", color=(1, 0, 0)) == (
        EdgeBorder(1, 1, "left", "thin", (1, 0, 0)), EdgeBorder(2, 1, "left", "thin", (1, 0, 0)))


# --- validation -----------------------------------------------------------


@pytest.mark.parametrize("kwargs", [
    {"h_align": "justify"}, {"v_align": "middle"}, {"rotation": 45}, {"rotation": 180}, {"rotation": 90.0},
    {"number_format": ""}, {"number_format": "x" * 256}, {"fill": (2, 0, 0)}, {"font": "IPAex明朝"},
])
def test_cell_rejects_invalid_formats(kwargs):
    with pytest.raises(ValueError):
        CellValue(0, 0, "x", **kwargs)


def test_rich_text_and_runs_are_validated():
    with pytest.raises(ValueError, match="run"):
        RichText(())
    with pytest.raises(ValueError, match="Run"):
        RichText(("x",))
    with pytest.raises(ValueError):
        Run("")
    with pytest.raises(ValueError, match="vert_align"):
        Run("x", vert_align="super")
    with pytest.raises(ValueError):
        Run("x", font="IPAex明朝")


@pytest.mark.parametrize("bounds", [(1, 0, 0, 0), (0, 1, 0, 0), (-1, 0, 0, 0), (0, 0, 0, 0)])
def test_cell_ranges_must_be_ordered_and_span_more_than_one_cell(bounds):
    with pytest.raises(ValueError):
        CellRange(*bounds)


def test_overlapping_merges_are_errors():
    with pytest.raises(ValueError, match="overlap"):
        _book(merges=(CellRange(0, 0, 1, 1), CellRange(1, 1, 2, 2)))


def test_merge_outside_the_sheet_is_an_error():
    with pytest.raises(ValueError, match="outside"):
        _book(merges=(CellRange(0, 0, 0, 6),))


def test_value_inside_a_merge_but_not_at_its_top_left_is_an_error():
    with pytest.raises(ValueError, match="top-left"):
        _book((CellValue(0, 1, "x"),), merges=(CellRange(0, 0, 1, 1),))


def test_formatted_blank_cell_inside_a_merge_is_allowed():
    assert _book((CellValue(1, 1, None),), merges=(CellRange(0, 0, 1, 1),))


@pytest.mark.parametrize("kwargs", [{"side": "middle"}, {"style": "thinner"}, {"style": "none"}, {"color": (0, 0)}])
def test_edge_border_rejects_unknown_side_style_or_color(kwargs):
    arguments = {"row": 0, "col": 0, "side": "top", "style": "thin"} | kwargs
    with pytest.raises(ValueError):
        EdgeBorder(**arguments)


def test_the_same_physical_edge_given_twice_is_an_error():
    with pytest.raises(ValueError, match="same edge"):
        _book(borders=(EdgeBorder(0, 0, "bottom", "thin"), EdgeBorder(1, 0, "top", "thin")))
    with pytest.raises(ValueError, match="same edge"):
        _book(borders=(EdgeBorder(0, 0, "right", "thin"), EdgeBorder(0, 1, "left", "medium")))
    with pytest.raises(ValueError, match="same edge"):
        _book(borders=(EdgeBorder(0, 0, "top", "thin"), EdgeBorder(0, 0, "top", "thin")))


def test_border_on_an_edge_inside_a_merge_is_an_error():
    with pytest.raises(ValueError, match="inside the merged"):
        _book(merges=(CellRange(0, 0, 1, 1),), borders=(EdgeBorder(0, 0, "bottom", "thin"),))
    with pytest.raises(ValueError, match="inside the merged"):
        _book(merges=(CellRange(0, 0, 1, 1),), borders=(EdgeBorder(1, 1, "left", "thin"),))


def test_border_outside_the_sheet_is_an_error():
    with pytest.raises(ValueError, match="outside"):
        _book(borders=(EdgeBorder(6, 0, "top", "thin"),))


@pytest.mark.parametrize("code", ['0.0"', "[Red", '"a";"b";"c";"d";"e"', "0.0\\", "0_", "0*", "[Red]]"])
def test_number_formats_excel_refuses_to_open_are_errors(code):
    # Excel 16.0 would not open a workbook with '0.0"' or '[Red' (it needs repair).
    with pytest.raises(ValueError, match="number_format"):
        CellValue(0, 0, 1.0, number_format=code)


@pytest.mark.parametrize("code", ['0.000;"−"0.000', '[Red]0.0;[Blue]-0.0;0;@', '\\"0', '0_)', '0*-', '"[";"]"',
                                  "[$-411]0", '#,##0.0"kN"'])
def test_well_formed_number_formats_are_accepted(code):
    assert CellValue(0, 0, 1.0, number_format=code).number_format == code


# --- review round 1 ----------------------------------------------------------


@pytest.mark.parametrize("text", ["_x0041_", "a_x000D_b", "a\rb", "a\r\nb", "_x005F_"])
def test_cell_text_that_spreadsheetml_would_change_is_escaped(text):
    book = _book((CellValue(0, 0, text), CellValue(0, 1, RichText((Run(text), Run("2", vert_align="superscript"))))))
    xml = _part(book, "xl/sharedStrings.xml")
    assert "\r" not in xml
    if "_x" in text:
        assert "_x005F_x" in xml
    if "\r" in text:
        assert "_x000D_" in xml


def test_spreadsheetml_text_escape():
    from gridlayoutpdf.renderers.xlsx._xml import escape_spreadsheet_text
    assert escape_spreadsheet_text("a_x0041_b") == "a_x005F_x0041_b"
    assert escape_spreadsheet_text("a\r\nb&<") == "a_x000D_\nb&amp;&lt;"
    assert escape_spreadsheet_text("_x12_ _xZZZZ_") == "_x12_ _xZZZZ_"


@pytest.mark.parametrize("code", ["[Foo]0", "[]0", "[Color0]0", "[Color57]0", "0" * 254 + "😀"])
def test_number_formats_with_unknown_brackets_or_too_long_are_errors(code):
    with pytest.raises(ValueError, match="number_format"):
        CellValue(0, 0, 1.0, number_format=code)


@pytest.mark.parametrize("code", ["[Red]0", "[Color10]0", "[Color56]0", "[>=100]0", "[>.5]0", "[>=1E3]0", "[$-411]0",
                                  "[$€-407]0", "[h]:mm", "[DBNum1]0", "[dbnum1]0"])
def test_number_formats_with_known_brackets_are_accepted(code):
    assert CellValue(0, 0, 1.0, number_format=code).number_format == code


# --- review round 3 ----------------------------------------------------------


@pytest.mark.parametrize("text", ["_x0041_", "a\rb", "_x0041_x0042_"])
def test_rich_text_runs_are_escaped_on_their_own(text):
    import re

    book = _book((CellValue(0, 0, RichText((Run(text), Run("2", vert_align="superscript")))),))
    xml = _part(book, "xl/sharedStrings.xml")
    runs = re.findall(r"<t(?: [^>]*)?>([^<]*)</t>", xml)
    decoded = [re.sub("_x([0-9A-Fa-f]{4})_", lambda match: chr(int(match.group(1), 16)), run) for run in runs]
    assert decoded == [text, "2"]
