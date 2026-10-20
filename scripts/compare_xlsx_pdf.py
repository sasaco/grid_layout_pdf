"""Compare the library PDF with Excel's PDF of the same pages' grid Excel workbook (Step 10, second stage).

The pages are laid out once, written as the library PDF (``render_pdf``) and as
a grid Excel workbook (``render_xlsx``, ``strict=True`` unless --no-strict).
A hidden Excel (PowerShell COM, the same driver as the Excel tests: the
started process is followed by its PID and stopped if it stays) opens the
workbook read-only, exports every sheet to PDF with its active printer (never
changed, only recorded), counts each sheet's shapes and reads the screen text
of every number cell. The checks are:

1. ``opens_without_repair``: ``Workbooks.Open`` with alerts off fails on a
   workbook Excel would repair (measured in Step 7), and no new repair log
   (``%TEMP%\\error*.xml``) appears (none was ever seen; a secondary check).
2. ``readback``: Excel has as many shapes on each sheet as the file's drawing
   (Excel may open a workbook and silently drop drawing parts).
3. ``strict``: the workbook was written without shifts or shape fallbacks.
4. ``no_hashes``: no number cell shows ``#`` (Excel's ``###`` for a too
   narrow cell) or nothing (a cell too narrow for one ``#``, measured at
   11 px) in its screen text (``Range.Text``), and Excel counts as many number
   cells as the workbook holds (the top-left cell of a merged range).
5. ``page_count`` and ``page_size``: Excel's PDF has the library's pages,
   sheet by sheet, on paper within 1 pt of the library's.
6. ``positions``: every text piece of the library PDF (a PyMuPDF span,
   whitespace removed) is printed on the same page of Excel's PDF, in the same
   direction, at about its size (a lost superscript prints at the base size)
   and within one grid cell along its reading direction and half a cell
   across it (the page grid's pitch). Pieces are matched by text: Excel's
   glyphs of one direction and baseline are read in order as one row and
   searched for the piece's text, longer texts first, equal texts nearest
   first, one glyph for one piece. Excel prints the bundled IPAex faces with
   another face (measured: ＭＳ ゴシック), whose glyph widths differ, so
   positions are compared at the anchor the workbook cell's alignment gives:
   the left edge of a left-aligned cell, the right edge of a right-aligned
   one, the centre of a centred one, and in the other axis the top, middle or
   bottom by the vertical alignment. The cell is the one holding the centre of
   Excel's text, else the one holding the library piece's baseline; a header
   or footer piece takes its decoration's alignment; only the other pieces in
   no cell (figure text) take, per axis, the nearest of the three anchors.
   Pieces whose FR-021 vertical merge is user-confirmed (category
   ``fr021_vertical_merge``, from the tags) are judged vertically against
   their merged range (lead decision 2026-10-20, see ``_rule``). No glyph
   of a piece may have lost ink to Excel's clip (``clipped``: each glyph's
   box rendered as printed and with the clipping removed, ClipLoss), and no
   glyph may be clipped away entirely (``clipped_out``). A report without
   any text piece, or a page that draws text without a piece, does not
   pass.
7. ``no_extra_text``: every glyph Excel printed belongs to a piece.

The result is ``result.json`` (per page: counts, the largest offsets, the
pieces not found or out of tolerance, Excel's unused glyphs; the cells
showing ``#``; the printer and Excel's version) beside the library PDF, the
workbook, Excel's PDFs and, per page, a PNG of the library page and Excel's
page side by side. A check that could not run is null in ``checks`` (when
Excel could not open the workbook, all but ``opens_without_repair`` and
``strict``), ``checks_complete`` is then false, and it never passes. Exit
codes: 0 passed, 1 a check failed (or Excel could not open the workbook), 2
bad arguments,
3 no Excel here, 4 the library could not write the report (a strict error is
reported as it is), 5 Excel did not run to the end, 6 the comparison itself
failed (an internal error).

    uv run --locked --extra dev python -X utf8 scripts/compare_xlsx_pdf.py --example representative --strict
    uv run --locked --extra dev python -X utf8 scripts/compare_xlsx_pdf.py --example cantilever
"""

import argparse
import bisect
from collections import Counter
from collections.abc import Callable, Iterable, Sequence
from contextlib import ExitStack
from dataclasses import dataclass, fields, replace
import importlib.util
import json
import math
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import time
import traceback
from typing import NamedTuple, Protocol
import xml.etree.ElementTree as ElementTree
import zipfile

import pymupdf
from PIL import Image, ImageChops, ImageDraw

from gridlayoutpdf import FontRegistry, ScenePage, Text, render_pdf, render_xlsx
from gridlayoutpdf.renderers.xlsx import XlsxResult, XlsxStrictError
from gridlayoutpdf.semantics import Kind

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
REPRESENTATIVE = REPOSITORY_ROOT / "examples" / "representative_report.py"
REPRESENTATIVE_SHEET = "代表例"  # the example's SHEET_NAME, checked when it is loaded
CANTILEVER_BUILDER = REPOSITORY_ROOT / "tests" / "fixtures" / "cantilever_p01_37" / "builder.py"
CANTILEVER_SHEET = "片持ち梁"
CANTILEVER_PAGES = range(1, 38)  # FR-014: the fixed 37 original pages
DEFAULT_FONT = "IPAex明朝"

EXIT_PASSED = 0
EXIT_FAILED = 1
EXIT_USAGE = 2  # argparse's exit code
EXIT_NO_EXCEL = 3
EXIT_LIBRARY_ERROR = 4
EXIT_EXCEL_ERROR = 5
EXIT_INTERNAL_ERROR = 6

EPSILON_PT = 1e-6  # offsets this close to the tolerance are within it
BASELINE_TOLERANCE_PT = 1.0  # Excel's glyphs whose baselines are this close are one row
ORIGIN_TOLERANCE_PT = 0.05  # a PDF glyph origin this close to a Text command's baseline lies on it
LOOKUP_NUDGE_PT = 0.01  # the library piece's cell is looked up this far from its baseline, on its glyphs' side
# Excel prints a piece at about its size (9.96 pt for 10, 6.60 pt for a 6.5 pt script, measured); a
# superscript that lost its format prints at the base size, about 1.5 times the script size.
SIZE_RATIO_LIMITS = (0.75, 1 / 0.75)
PAGE_SIZE_TOLERANCE_PT = 1.0  # Excel's A4 is 595.32 x 841.92 pt against the library's 595 x 842 (measured)
# Excel's clipping is found by rendering each glyph's box of its pages as printed and with the clip
# operators removed (ClipLoss): a glyph is clipped when the renders differ inside its box. CLIP_ZOOM px per pt
# (1/16 pt pixels, each glyph rendered on its own); a pixel differs when its grey levels differ by more than
# CLIP_LEVEL. An edge pixel whose ink lies wholly inside the clip changes by at most 255 c (1 - c) <= 63.75
# (c: the part of the pixel inside the clip), so 64 never counts a clip edge that took no ink (the largest
# such difference measured was 46). Review round 4 measured at 4 px per pt cuts of 0.08-0.19 pt missed;
# a cut of a third of a point is found at any pixel phase (unit test).
CLIP_ZOOM = 16
CLIP_LEVEL = 64
RECORD_LIMIT = 50  # shifts and fallbacks listed for a strict error; unused Excel texts listed per page
# Explicit flags. They hold TEXT_CLIP (64, the same bit as TEXT_MEDIABOX_CLIP in PyMuPDF 1.28; unit
# test): with it MuPDF leaves out a glyph outside the clip (in Excel's PDF one whose box's centre is
# outside, review round 4). A glyph cut but extracted is marked by ClipLoss; the ones left out are found in
# the unclipped copy and listed as clipped out (ClipPage.dropped).
TEXT_FLAGS = pymupdf.TEXTFLAGS_RAWDICT
# Piece categories. FR021: an equation part that FR-021 (user-confirmed) merges over its equation line's
# rows and centres vertically (lead decision 2026-10-20): the label, lhs, relations, result and unit of an
# equation line with a fraction. Its vertical position is judged against its merged range.
STANDARD = "standard"
FR021 = "fr021_vertical_merge"
CATEGORIES = (STANDARD, FR021)
FR021_PARTS = frozenset({"label", "lhs", "relation", "result_relation", "result", "unit"})
HORIZONTAL_ALIGNMENTS = ("left", "center", "right")
VERTICAL_ALIGNMENTS = ("top", "center", "bottom")
RULES = ("ordinary", "merged_range", "no_cell")
NEAREST = "nearest"
_SAFE_NAME = re.compile(r"[^\w.-]")


class ExcelError(RuntimeError):
    """Excel did not run the export to the end (the script failed, timed out, Excel stayed or answered oddly)."""


# --- geometry and matching (Excel-independent) -------------------------------------------------------

def _finite(value, name: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"{name} must be a finite number, not {value!r}")
    return float(value)


@dataclass(frozen=True)
class Box:
    """A box in PDF points, top-left origin, y down (PyMuPDF's coordinates)."""

    x0: float
    y0: float
    x1: float
    y1: float

    def __post_init__(self) -> None:
        for name in ("x0", "y0", "x1", "y1"):
            _finite(getattr(self, name), name)
        if self.x0 > self.x1 or self.y0 > self.y1:
            raise ValueError(f"a box needs x0 <= x1 and y0 <= y1, not {self}")

    @classmethod
    def around(cls, boxes: Iterable["Box"]) -> "Box":
        items = list(boxes)
        if not items:
            raise ValueError("a box needs at least one box to enclose")
        return cls(min(b.x0 for b in items), min(b.y0 for b in items), max(b.x1 for b in items),
                   max(b.y1 for b in items))

    @classmethod
    def of(cls, bbox: Sequence[float]) -> "Box":
        """A PyMuPDF bbox (x0, y0, x1, y1)."""
        x0, y0, x1, y1 = bbox
        return cls(x0, y0, x1, y1)

    @property
    def centre(self) -> tuple[float, float]:
        return (self.x0 + self.x1) / 2, (self.y0 + self.y1) / 2

    def holds(self, point: tuple[float, float]) -> bool:
        return self.x0 <= point[0] <= self.x1 and self.y0 <= point[1] <= self.y1

    def as_list(self) -> list[float]:
        return [round(self.x0, 3), round(self.y0, 3), round(self.x1, 3), round(self.y1, 3)]


@dataclass(frozen=True)
class Piece:
    """A text piece of the library PDF: a span's text without whitespace and its glyphs' box.

    ``lookup`` is the point whose cell the piece is in by the map's rule (on
    its baseline, nudged to its glyphs' side), ``align`` the horizontal
    alignment of its header or footer, ``figure`` whether a figure drew it
    (both from the tags), ``size`` its font size (None: unknown, which fails).
    """

    text: str
    box: Box
    vertical: bool = False  # read along y (a quarter-turned cell)
    category: str = STANDARD  # STANDARD or FR021, from the tags of the commands that drew it
    size: float | None = None
    lookup: tuple[float, float] | None = None
    align: str | None = None
    figure: bool = False

    def __post_init__(self) -> None:
        if not isinstance(self.text, str) or not self.text or any(char.isspace() for char in self.text):
            raise ValueError(f"a piece's text must be non-empty without whitespace, not {self.text!r}")
        if self.category not in CATEGORIES:
            raise ValueError(f"category must be one of {', '.join(CATEGORIES)}, not {self.category!r}")
        if self.size is not None and _finite(self.size, "size") <= 0:
            raise ValueError("size must be positive")
        if self.lookup is not None:
            if len(self.lookup) != 2:
                raise ValueError(f"lookup is an (x, y) point, not {self.lookup!r}")
            object.__setattr__(self, "lookup", (_finite(self.lookup[0], "lookup x"), _finite(self.lookup[1], "lookup y")))
        if self.align is not None and self.align not in HORIZONTAL_ALIGNMENTS:
            raise ValueError(f"align must be one of {', '.join(HORIZONTAL_ALIGNMENTS)}, not {self.align!r}")


@dataclass(frozen=True)
class Glyph:
    """One glyph Excel printed (whitespace is never a glyph here)."""

    char: str
    box: Box
    size: float | None = None
    vertical: bool = False
    clipped: bool = False  # Excel's clip took ink of this glyph (ClipLoss)

    def __post_init__(self) -> None:
        if not isinstance(self.char, str) or len(self.char) != 1 or self.char.isspace():
            raise ValueError(f"a glyph is one non-space character, not {self.char!r}")
        if self.size is not None and _finite(self.size, "size") <= 0:
            raise ValueError("size must be positive")


@dataclass(frozen=True)
class CellArea:
    """A workbook cell (its merged range) with a value, on its page in the library PDF's coordinates."""

    box: Box
    horizontal: str  # "left", "center" or "right"
    vertical: str  # "top", "center" or "bottom"
    address: str

    def __post_init__(self) -> None:
        if self.horizontal not in HORIZONTAL_ALIGNMENTS or self.vertical not in VERTICAL_ALIGNMENTS:
            raise ValueError(f"unknown alignment {self.horizontal!r}/{self.vertical!r}")


@dataclass(frozen=True)
class Match:
    """One piece and the glyphs Excel printed for it.

    ``rule`` is how the position is judged: "ordinary" (dx and dy within
    the tolerance), "merged_range" (an FR021 piece in a cell merged over
    rows: dx within the tolerance, Excel's text centre within half a cell of
    the range's middle, ``range_dy``, and the library piece on the range's
    rows within half a cell, ``library_in_range``) or "no_cell" (an FR021
    piece, or a piece neither a header, footer nor figure text, Excel printed
    in no cell: always outside). Every rule also needs a size ratio (Excel's
    glyph size over the piece's) within SIZE_RATIO_LIMITS (an unknown size
    fails) and no glyph whose ink Excel's clip took (``clipped_glyphs``).
    """

    piece: Piece
    box: Box  # Excel's glyphs
    dx: float  # Excel minus library, pt
    dy: float
    tolerance: tuple[float, float]
    anchor: tuple[str, str]  # the cell's alignment, the decoration's, or ("nearest", "nearest")
    cell: str | None
    rule: str = "ordinary"
    range_dy: float | None = None  # Excel's text centre minus the merged range's middle, pt
    library_in_range: bool | None = None
    size_ratio: float | None = None  # None: a size is unknown
    clipped_glyphs: tuple[str, ...] = ()  # the characters of the glyphs Excel's clip cut

    def __post_init__(self) -> None:
        if self.rule not in RULES:
            raise ValueError(f"rule must be one of {', '.join(RULES)}")
        if (self.rule == "merged_range") != (self.range_dy is not None) or (
                self.rule == "merged_range") != (self.library_in_range is not None):
            raise ValueError("range_dy and library_in_range are given exactly for the merged_range rule")
        if self.rule == "merged_range" and self.piece.category != FR021:
            raise ValueError("only fr021_vertical_merge pieces take the merged_range rule")
        if len(self.tolerance) != 2 or not all(_finite(value, "tolerance") > 0 for value in self.tolerance):
            raise ValueError("the tolerance is a pair of positive numbers")
        if (len(self.anchor) != 2 or self.anchor[0] not in (*HORIZONTAL_ALIGNMENTS, NEAREST)
                or self.anchor[1] not in (*VERTICAL_ALIGNMENTS, NEAREST)):
            raise ValueError(f"unknown anchor {self.anchor!r}")
        if self.size_ratio is not None and _finite(self.size_ratio, "size_ratio") <= 0:
            raise ValueError("size_ratio must be positive")

    @property
    def failures(self) -> tuple[str, ...]:
        """What fails: "no_cell", "dx", "dy", "range_dy", "library_off_range", "size", "size_unknown", "clipped"."""
        found = []
        if self.rule == "no_cell":
            found.append("no_cell")
        if abs(self.dx) > self.tolerance[0] + EPSILON_PT:
            found.append("dx")
        if self.rule == "ordinary" and abs(self.dy) > self.tolerance[1] + EPSILON_PT:
            found.append("dy")
        if self.rule == "merged_range":
            if abs(self.range_dy) > self.tolerance[1] + EPSILON_PT:
                found.append("range_dy")
            if not self.library_in_range:
                found.append("library_off_range")
        if self.size_ratio is None:
            found.append("size_unknown")
        elif not SIZE_RATIO_LIMITS[0] <= self.size_ratio <= SIZE_RATIO_LIMITS[1]:
            found.append("size")
        if self.clipped_glyphs:
            found.append("clipped")
        return tuple(found)

    @property
    def within(self) -> bool:
        return not self.failures

    @property
    def distance(self) -> float:
        """The library-to-Excel offset in tolerances, which orders the candidates of one text.

        Always the physical offset, also for the merged_range rule: the range
        offset is small in any merged range, so it would let a piece take
        another equation's glyphs.
        """
        return math.hypot(self.dx / self.tolerance[0], self.dy / self.tolerance[1])


@dataclass(frozen=True)
class PageMatch:
    matches: tuple[Match, ...]  # in the library's piece order
    missing: tuple[Piece, ...]
    extra: tuple[str, ...] = ()  # Excel's runs of glyphs no piece took, in row order

    @property
    def outside(self) -> tuple[Match, ...]:
        return tuple(item for item in self.matches if not item.within)


_HORIZONTAL = {"left": 0.0, "center": 0.5, "right": 1.0}
_VERTICAL = {"top": 0.0, "center": 0.5, "bottom": 1.0}


def anchor_point(box: Box, horizontal: str, vertical: str) -> tuple[float, float]:
    """The point of ``box`` a cell with this alignment keeps in place (paper axes, also for turned text)."""
    try:
        fx, fy = _HORIZONTAL[horizontal], _VERTICAL[vertical]
    except KeyError:
        raise ValueError(f"unknown alignment {horizontal!r}/{vertical!r}") from None
    return box.x0 + fx * (box.x1 - box.x0), box.y0 + fy * (box.y1 - box.y0)


def tolerance(piece: Piece, pitch_x: float, pitch_y: float) -> tuple[float, float]:
    """(x, y) tolerance: one cell along the reading direction, half a cell across it.

    The plan's "one cell across, half a cell down" read along the reading
    direction: a turned table is laid out on the page grid with its axes
    swapped, so its local column is the page's pitch_y (implementation
    choice, not confirmed by the user; the paper-axes count is reported too).
    """
    return (pitch_x / 2, pitch_y) if piece.vertical else (pitch_x, pitch_y / 2)


def find_cell(printed: Box, cells: Sequence[CellArea], lookup: tuple[float, float] | None = None) -> CellArea | None:
    """The cell holding the centre of the text Excel printed, else the one holding ``lookup``.

    Excel's text first: the map puts a piece in the row of its baseline,
    while the glyph box of a text centred on a two-row band has its centre
    in the row above, and Excel prints a cell's text inside the cell (within
    about 1 pt of the PDF's grid). When Excel's text left every cell, the
    library piece's own cell (``lookup``, on its baseline) is taken, so a
    shift out of the cell is still measured against the cell's alignment.
    """
    for point in (printed.centre, lookup):
        if point is None:
            continue
        for area in cells:
            if area.box.holds(point):
                return area
    return None


def _offset(piece: Piece, excel: Box, area: CellArea | None) -> tuple[float, float, tuple[str, str]]:
    """Excel minus library at the anchor: the decoration's alignment, else the cell's, else the nearest."""
    library = piece.box
    if piece.align is not None or area is not None:
        horizontal, vertical = (piece.align, "center") if piece.align is not None else (area.horizontal,
                                                                                         area.vertical)
        lx, ly = anchor_point(library, horizontal, vertical)
        ex, ey = anchor_point(excel, horizontal, vertical)
        return ex - lx, ey - ly, (horizontal, vertical)
    dx = min(((excel.x0 + f * (excel.x1 - excel.x0)) - (library.x0 + f * (library.x1 - library.x0))
              for f in (0.0, 0.5, 1.0)), key=abs)
    dy = min(((excel.y0 + f * (excel.y1 - excel.y0)) - (library.y0 + f * (library.y1 - library.y0))
              for f in (0.0, 0.5, 1.0)), key=abs)
    return dx, dy, (NEAREST, NEAREST)


def _rule(piece: Piece, printed: Box, area: CellArea | None,
          pitch_y: float) -> tuple[str, float | None, bool | None]:
    """How ``piece`` is judged (see Match), with ``range_dy`` and ``library_in_range`` for merged_range.

    FR021 pieces in a range merged over rows take the range rule; Excel's
    text is in the range (the cell is looked up by it alone, see match_page)
    and the library piece must lie on the range's rows within half a cell,
    so a match to another equation's range fails. An FR021 piece in a
    one-row cell has no vertical merge to excuse an offset, so it takes the
    ordinary rule. A piece in no cell is outside (no_cell) unless it is a
    header's or footer's (its decoration's alignment) or a figure's (nearest
    anchors) text, both known from the tags.
    """
    if area is None:
        if piece.category != FR021 and (piece.align is not None or piece.figure):
            return "ordinary", None, None
        return "no_cell", None, None
    if piece.category != FR021:
        return "ordinary", None, None
    if area.box.y1 - area.box.y0 <= 1.5 * pitch_y:
        return "ordinary", None, None
    y = (piece.lookup or piece.box.centre)[1]
    inside = area.box.y0 - pitch_y / 2 - EPSILON_PT <= y <= area.box.y1 + pitch_y / 2 + EPSILON_PT
    return "merged_range", printed.centre[1] - (area.box.y0 + area.box.y1) / 2, inside


def _occurrences(text: str, row: str) -> list[int]:
    found, start = [], row.find(text)
    while start >= 0:
        found.append(start)
        start = row.find(text, start + 1)
    return found


def _size_ratio(piece: Piece, glyphs: Sequence[Glyph]) -> float | None:
    """Excel's mean glyph size over the piece's; None when either side does not know a size (which fails)."""
    if piece.size is None or any(glyph.size is None for glyph in glyphs):
        return None
    return sum(glyph.size for glyph in glyphs) / len(glyphs) / piece.size


def match_page(pieces: Sequence[Piece], rows: Sequence[Sequence[Glyph]], cells: Sequence[CellArea],
               pitch_x: float, pitch_y: float) -> PageMatch:
    """Match the library's pieces to Excel's glyph rows of the same page by text.

    Longer texts are matched first (a superscript "3" must not take the "3"
    of "12.375"); the pieces of one text take Excel's occurrences nearest
    first (distance in tolerances); a glyph serves one piece; a piece takes
    only glyphs read in its direction. Glyphs no piece took are ``extra``.
    The cell is the one of Excel's text, else (not for FR021 pieces, whose
    rule is the merged range Excel printed them in) the library piece's.
    """
    strings = ["".join(glyph.char for glyph in row) for row in rows]
    by_text: dict[str, list[int]] = {}
    for index, piece in enumerate(pieces):
        by_text.setdefault(piece.text, []).append(index)
    used: set[tuple[int, int]] = set()
    found: dict[int, Match] = {}
    for text in sorted(by_text, key=lambda item: (-len(item), item)):
        places = []
        for row, string in enumerate(strings):
            for start in _occurrences(text, string):
                glyphs = rows[row][start:start + len(text)]
                box = Box.around(glyph.box for glyph in glyphs)
                places.append((row, start, glyphs, box, glyphs[0].vertical))
        candidates = []
        for index in by_text[text]:
            piece = pieces[index]
            limit = tolerance(piece, pitch_x, pitch_y)
            for row, start, glyphs, box, vertical in places:
                if vertical != piece.vertical:
                    continue
                # Figure text is drawn as shapes, never in a cell; FR021 pieces take only Excel's cell.
                area = None if piece.figure else find_cell(box, cells, None if piece.category == FR021
                                                           else piece.lookup)
                dx, dy, anchor = _offset(piece, box, area)
                rule, range_dy, inside = _rule(piece, box, area, pitch_y)
                item = Match(piece, box, dx, dy, limit, anchor, None if area is None else area.address, rule,
                             range_dy, inside, _size_ratio(piece, glyphs),
                             tuple(glyph.char for glyph in glyphs if glyph.clipped))
                candidates.append((item.distance, index, row, start, item))
        candidates.sort(key=lambda candidate: candidate[:4])
        for _, index, row, start, item in candidates:
            taken = {(row, k) for k in range(start, start + len(text))}
            if index in found or taken & used:
                continue
            used |= taken
            found[index] = item
    extra = []
    for row, string in enumerate(strings):
        run = ""
        for k, char in enumerate(string):
            if (row, k) in used:
                if run:
                    extra.append(run)
                run = ""
            else:
                run += char
        if run:
            extra.append(run)
    return PageMatch(tuple(found[index] for index in sorted(found)),
                     tuple(piece for index, piece in enumerate(pieces) if index not in found), tuple(extra))


def _largest(matches: Sequence[Match], axis: str) -> tuple[float | None, str | None]:
    """The largest offset and its piece; no_cell pieces have no anchor to measure at and are left out."""
    matches = [match for match in matches if match.rule != "no_cell"]
    if not matches:
        return None, None
    item = max(matches, key=lambda match: abs(getattr(match, axis)))
    return round(abs(getattr(item, axis)), 3), item.piece.text


def _paper_outside(item: Match) -> bool:
    """A vertical piece out of the paper-axes tolerance (one cell across, half a cell down)."""
    if not item.piece.vertical or item.rule != "ordinary":
        return False
    pitch_x, pitch_y = item.tolerance[0] * 2, item.tolerance[1]  # tolerance() of a vertical piece
    return abs(item.dx) > pitch_x + EPSILON_PT or abs(item.dy) > pitch_y / 2 + EPSILON_PT


def page_summary(number: int, sheet: str, result: PageMatch) -> dict:
    """The JSON record of one page (offsets in pt, Excel minus library).

    The counts and largest offsets at the top are the standard pieces';
    the FR021 pieces have theirs under FR021 (``max_range_dy_pt`` is the
    largest offset of Excel's text centre from its merged range's middle).
    ``missing`` and ``outside`` list both, each with its category.
    ``nearest`` counts the figure pieces compared at their nearest anchors;
    ``outside_if_paper_axes`` the vertical pieces that the paper-axes reading
    of the tolerance would put outside; ``extra`` lists Excel's glyph runs no
    piece took. Each ``outside`` item lists its ``failures`` (offsets,
    ``no_cell``, ``size``, ``clipped`` with its ``clipped_glyphs``). The
    page's ``clipped_out`` glyphs are added by the run.
    """
    standard = [item for item in result.matches if item.piece.category == STANDARD]
    merged = [item for item in result.matches if item.piece.category == FR021]
    max_dx, dx_text = _largest(standard, "dx")
    max_dy, dy_text = _largest(standard, "dy")
    ranged = [item for item in merged if item.range_dy is not None]
    nearest = [item.piece.text for item in result.matches if item.piece.figure and item.anchor == (NEAREST, NEAREST)]
    return {
        "page": number, "sheet": sheet,
        "pieces": len(standard) + sum(piece.category == STANDARD for piece in result.missing),
        "matched": len(standard),
        "missing_count": sum(piece.category == STANDARD for piece in result.missing),
        "outside_count": sum(not item.within for item in standard),
        "max_dx_pt": max_dx, "max_dx_text": dx_text, "max_dy_pt": max_dy, "max_dy_text": dy_text,
        FR021: {"pieces": len(merged) + sum(piece.category == FR021 for piece in result.missing),
                "matched": len(merged), "missing": sum(piece.category == FR021 for piece in result.missing),
                "outside": sum(not item.within for item in merged),
                "max_dx_pt": _largest(merged, "dx")[0], "max_dy_pt": _largest(merged, "dy")[0],
                "max_range_dy_pt": _largest(ranged, "range_dy")[0]},
        "nearest": len(nearest), "nearest_texts": nearest[:RECORD_LIMIT],
        "outside_if_paper_axes": sum(_paper_outside(item) for item in result.matches),
        "missing": [{"text": piece.text, "category": piece.category, "box": piece.box.as_list(),
                     "vertical": piece.vertical} for piece in result.missing],
        "outside": [{"text": item.piece.text, "category": item.piece.category, "rule": item.rule,
                     "failures": list(item.failures), "clipped_glyphs": list(item.clipped_glyphs),
                     # no_cell has no anchor to measure at (the nearest anchors would understate it)
                     "dx_pt": None if item.rule == "no_cell" else round(item.dx, 3),
                     "dy_pt": None if item.rule == "no_cell" else round(item.dy, 3),
                     "range_dy_pt": None if item.range_dy is None else round(item.range_dy, 3),
                     "size_ratio": None if item.size_ratio is None else round(item.size_ratio, 3),
                     "tolerance_pt": [round(value, 3) for value in item.tolerance], "anchor": list(item.anchor),
                     "cell": item.cell, "library_box": item.piece.box.as_list(), "excel_box": item.box.as_list(),
                     "vertical": item.piece.vertical}
                    for item in result.outside],
        "extra_count": len(result.extra), "extra": list(result.extra[:RECORD_LIMIT]),
    }


# --- reading the PDFs ----------------------------------------------------------------------------------

def _text_lines(page: pymupdf.Page):
    for block in page.get_text("rawdict", flags=TEXT_FLAGS)["blocks"]:
        if block.get("type") == 0:
            yield from block["lines"]


class CommandInfo(NamedTuple):
    """What the tags of one Text command say, and where its baseline starts (PDF coordinates)."""

    x: float
    y: float
    category: str  # STANDARD or FR021
    align: str | None  # its header's or footer's horizontal alignment
    figure: bool  # drawn by a Figure
    turns: int  # quarter turns counterclockwise, 0 to 3
    glyphs: int  # its non-whitespace characters (the self-check of pdf_pieces)


# The reading direction (PyMuPDF's line "dir") of a Text command turned counterclockwise by 0-3 quarters.
_DIRECTIONS = {0: (1, 0), 1: (0, -1), 2: (-1, 0), 3: (0, 1)}


def text_categories(scene: ScenePage) -> list[CommandInfo]:
    """The Text commands of a ScenePage as CommandInfo, from their tags alone.

    A command is FR021 when it is upright, its innermost equation line (or
    math block, as the map groups them) holds a fraction, it lies in an
    equation part of FR021_PARTS of that line, and not in the fraction: the
    map merges such upright pieces over the line's rows (FR-021). Placements
    number the structures, so the line is keyed by the placement as the map
    keys it. A header's or footer's command carries its decoration's
    ``align``; a figure's command is marked ``figure``.
    """
    owners: list = [("page",)] * len(scene.commands)
    for number, placement in enumerate(scene.metadata.get("placements", ())):
        start, stop = placement["commands"]
        owners[start:stop] = [("placement", number)] * (stop - start)
    texts, with_fraction = [], set()
    for index, command in enumerate(scene.commands):
        if not isinstance(command, Text):
            continue
        path = command.tag.path if command.tag is not None else ()
        decoration = next((scope for scope in path if scope.kind is Kind.DECORATION), None)
        align = decoration.data.get("align") if decoration is not None else None
        align = align if align in HORIZONTAL_ALIGNMENTS else None
        figure = any(scope.kind is Kind.FIGURE for scope in path)
        position = max((k for k, scope in enumerate(path) if scope.kind in (Kind.EQUATION_LINE, Kind.MATH_BLOCK)),
                       default=None)
        line, in_fraction, part = None, False, None
        if position is not None:
            line = (owners[index], path[:position], replace(path[position], part=None))
            inner = path[position + 1:]
            in_fraction = any(scope.kind is Kind.FRACTION for scope in inner)
            if in_fraction:
                with_fraction.add(line)
            part = next((scope.part for scope in inner if scope.kind is Kind.EQUATION_PART), None)
        texts.append((command, line, in_fraction, part, align, figure))
    return [CommandInfo(command.x, command.y,
                        FR021 if (command.quarter_turns % 4 == 0 and line is not None and line in with_fraction
                                  and not in_fraction and part in FR021_PARTS) else STANDARD,
                        align, figure, command.quarter_turns % 4, sum(not char.isspace() for char in command.text))
            for command, line, in_fraction, part, align, figure in texts]


def _command_of(char: dict, direction: tuple[float, float], commands: Sequence[CommandInfo]) -> int | None:
    """The index of the command a glyph came from (or None).

    The command reads in the glyph's direction, its baseline is the glyph's
    (within ORIGIN_TOLERANCE_PT), and of those it starts last at or before
    the glyph along the reading direction.
    """
    cos, sin = direction
    x, y = char["origin"]
    along, across = cos * x + sin * y, -sin * x + cos * y
    best, start = None, -math.inf
    for index, command in enumerate(commands):
        dx, dy = _DIRECTIONS[command.turns]
        if abs(dx - cos) > 1e-3 or abs(dy - sin) > 1e-3:
            continue
        command_along, command_across = cos * command.x + sin * command.y, -sin * command.x + cos * command.y
        if abs(command_across - across) <= ORIGIN_TOLERANCE_PT and start < command_along <= along + ORIGIN_TOLERANCE_PT:
            best, start = index, command_along
    return best


def piece_category(owners: Sequence[CommandInfo | None]) -> str:
    """FR021 when every glyph came from an FR021 command, else STANDARD (the strict side)."""
    return FR021 if owners and all(owner is not None and owner.category == FR021 for owner in owners) else STANDARD


def pdf_pieces(page: pymupdf.Page, scene: ScenePage | None = None) -> list[Piece]:
    """The library page's text pieces: each PyMuPDF span, whitespace removed, in the page's order.

    With the page's ``scene``, every glyph is given to the Text command that
    drew it; a command that got another number of glyphs than its text has,
    or a glyph no command drew, means the scene is not this page's
    (ValueError). A piece is FR021 by ``piece_category``, takes its
    decoration's alignment when all its glyphs came from one, and is figure
    text when all came from a figure's commands.
    """
    commands = text_categories(scene) if scene is not None else None
    counts = [0] * len(commands or ())
    pieces = []
    for line in _text_lines(page):
        cos, sin = line["dir"]
        vertical = abs(sin) > abs(cos)
        up = (sin, -cos)  # the glyphs' side of the baseline
        for span in line["spans"]:
            chars = [char for char in span["chars"] if not char["c"].isspace()]
            if not chars:
                continue
            owners: list[CommandInfo | None] = []
            if commands is not None:
                for char in chars:
                    index = _command_of(char, (cos, sin), commands)
                    if index is None:
                        raise ValueError(f"glyphs {char['c']!r} at {char['origin']} were not drawn by a Text command "
                                         "of the page's scene")
                    counts[index] += 1
                    owners.append(commands[index])
            aligns = {owner.align if owner is not None else None for owner in owners or [None]}
            box = Box.around(Box.of(char["bbox"]) for char in chars)
            ox, oy = chars[0]["origin"]
            centre_x, centre_y = box.centre
            lookup = ((ox + up[0] * LOOKUP_NUDGE_PT, centre_y) if vertical
                      else (centre_x, oy + up[1] * LOOKUP_NUDGE_PT))
            pieces.append(Piece("".join(char["c"] for char in chars), box, vertical, piece_category(owners),
                                span["size"], lookup, aligns.pop() if len(aligns) == 1 else None,
                                bool(owners) and all(owner.figure for owner in owners)))
    wrong = [(command, count) for command, count in zip(commands or (), counts) if count != command.glyphs]
    if wrong:
        command, count = wrong[0]
        raise ValueError(f"the page's glyphs do not match its scene: the Text command at ({command.x}, {command.y}) "
                         f"has {command.glyphs} glyphs, the PDF {count} ({len(wrong)} command(s) differ)")
    return pieces


def _judged(clipped: Callable[[Box], bool], box: Box, char: str, page: pymupdf.Page) -> bool:
    try:
        return clipped(box)
    except ValueError as exc:  # name the glyph: the error ends the run (exit code 6)
        raise ValueError(f"{exc} (glyph {char!r} on page {page.number + 1})") from exc


def pdf_glyph_rows(page: pymupdf.Page, clip_page: "ClipPage | None" = None) -> list[tuple[Glyph, ...]]:
    """Excel's glyphs (whitespace removed) as rows: one direction and baseline, in reading order.

    Excel's PDF splits spans and lines differently from the library's (cells
    next to each other may share a span), so the rows are rebuilt from the
    glyphs' origins. Superscripts and subscripts, off the baseline, are rows
    of their own on both sides (the library's are spans of their own).
    ``clip_page`` (ClipLoss.page of this page) marks the glyphs whose box
    holds ink Excel's clip took.
    """

    def clipped(box: Box) -> bool:
        return clip_page is not None and clip_page.clipped(box)

    groups: dict[tuple[float, float], list[tuple[float, float, Glyph]]] = {}
    for line in _text_lines(page):
        cos, sin = line["dir"]
        vertical = abs(sin) > abs(cos)
        for span in line["spans"]:
            for char in span["chars"]:
                if char["c"].isspace():
                    continue
                x, y = char["origin"]
                box = Box.of(char["bbox"])
                groups.setdefault((round(cos, 3), round(sin, 3)), []).append(
                    (-sin * x + cos * y, cos * x + sin * y,
                     Glyph(char["c"], box, span["size"], vertical, _judged(clipped, box, char["c"], page))))
    rows = []
    for key in sorted(groups):
        current: list[tuple[float, float, Glyph]] = []
        for item in sorted(groups[key], key=lambda item: (item[0], item[1])):
            if current and item[0] - current[-1][0] > BASELINE_TOLERANCE_PT:
                rows.append(current)
                current = []
            current.append(item)
        if current:
            rows.append(current)
    return [tuple(glyph for _, _, glyph in sorted(row, key=lambda item: item[1])) for row in rows]


_WHITESPACE = b" \t\r\n\f\x00"
_DELIMITERS = b"()<>[]{}/%"


def _literal_end(content: bytes, start: int) -> int:
    """The index after the literal string starting at ``start`` (balanced parentheses, backslash escapes)."""
    depth, index = 0, start
    while index < len(content):
        byte = content[index]
        if byte == 0x5C:  # backslash: the next byte is escaped
            index += 2
            continue
        depth += {0x28: 1, 0x29: -1}.get(byte, 0)
        index += 1
        if depth == 0:
            return index
    raise ValueError("a content stream has an unclosed string")


def without_clipping(content: bytes) -> bytes:
    """A content stream with its clip operators (W, W*) removed and everything else kept byte for byte.

    The stream is tokenized (strings, names, comments, arrays, dictionaries),
    so "W" inside a string is never touched. An inline image (BI) holds raw
    bytes the tokenizer cannot skip safely: ValueError.
    """
    out, index, end = bytearray(), 0, len(content)
    while index < end:
        byte = content[index]
        if byte in _WHITESPACE:
            stop = index + 1
        elif byte == 0x25:  # % comment to the end of the line
            stop = index + 1
            while stop < end and content[stop] not in b"\r\n":
                stop += 1
        elif byte == 0x28:  # ( literal string )
            stop = _literal_end(content, index)
        elif content.startswith(b"<<", index) or content.startswith(b">>", index):
            stop = index + 2
        elif byte == 0x3C:  # < hex string >
            stop = content.find(b">", index)
            if stop < 0:
                raise ValueError("a content stream has an unclosed hex string")
            stop += 1
        elif byte in b"[]{})>":
            stop = index + 1
        else:  # a name, a number or an operator
            stop = index + 1
            while stop < end and content[stop] not in _WHITESPACE and content[stop] not in _DELIMITERS:
                stop += 1
            token = content[index:stop]
            if token == b"BI":
                raise ValueError("a content stream holds an inline image; its clipping cannot be removed safely")
            if token in (b"W", b"W*"):
                out += b" " * len(token)
                index = stop
                continue
        out += content[index:stop]
        index = stop
    return bytes(out)


def _page_glyphs(page: pymupdf.Page) -> list[tuple[str, Box]]:
    return [(char["c"], Box.of(char["bbox"])) for line in _text_lines(page) for span in line["spans"]
            for char in span["chars"] if not char["c"].isspace()]


class ClipPage:
    """One page of Excel's PDF as printed and without clipping: which glyph boxes lost ink, which glyphs vanished.

    ``dropped`` lists the glyphs (character and box) found only without the
    clip: Excel's clip took them so far that they are not extracted at all.
    """

    def __init__(self, printed: pymupdf.Page, free: pymupdf.Page) -> None:
        self._printed, self._free = printed.get_displaylist(), free.get_displaylist()
        self._page = Box(printed.rect.x0, printed.rect.y0, printed.rect.x1, printed.rect.y1)
        kept = Counter((char, tuple(box.as_list())) for char, box in _page_glyphs(printed))
        dropped = []
        for char, box in _page_glyphs(free):
            key = (char, tuple(box.as_list()))
            if kept[key]:
                kept[key] -= 1
            else:
                dropped.append((char, box))
        self.dropped: tuple[tuple[str, Box], ...] = tuple(dropped)

    def clipped(self, box: Box) -> bool:
        """Did the clip take ink inside ``box`` (not counting the boxes of dropped glyphs, whose ink it is)?"""
        rect = pymupdf.Rect(box.x0, box.y0, box.x1, box.y1)
        if rect.is_empty:  # nothing to render: judging it unclipped would be a guess
            raise ValueError(f"the glyph box {box} has no area")
        page = self._page
        if box.x0 < page.x0 - EPSILON_PT or box.y0 < page.y0 - EPSILON_PT or box.x1 > page.x1 + EPSILON_PT \
                or box.y1 > page.y1 + EPSILON_PT:
            return True  # the paper cuts it (the renders stop at the page, so this is not measured)
        matrix, images, origin = pymupdf.Matrix(CLIP_ZOOM, CLIP_ZOOM), [], None
        for display in (self._printed, self._free):
            pixmap = display.get_pixmap(matrix=matrix, clip=rect, colorspace=pymupdf.csGRAY, alpha=False)
            images.append(Image.frombytes("L", (pixmap.width, pixmap.height), pixmap.samples))
            origin = (pixmap.x, pixmap.y)
        if images[0].size != images[1].size:
            raise ValueError(f"the renders of the glyph box {box} differ in size")
        difference = ImageChops.difference(*images).point(lambda value: 255 if value > CLIP_LEVEL else 0)
        painter = ImageDraw.Draw(difference)
        for _, other in self.dropped:
            if other.x0 < box.x1 and other.x1 > box.x0 and other.y0 < box.y1 and other.y1 > box.y0:
                painter.rectangle((math.floor(other.x0 * CLIP_ZOOM) - origin[0], math.floor(other.y0 * CLIP_ZOOM) - origin[1],
                                   math.ceil(other.x1 * CLIP_ZOOM) - origin[0], math.ceil(other.y1 * CLIP_ZOOM) - origin[1]),
                                  fill=0)
        return difference.getbbox() is not None


class ClipLoss:
    """Excel's PDF with and without its clipping (``page`` gives a ClipPage).

    Every content stream of the pages and every form XObject loses its clip
    operators in the second copy, and every form's /BBox (which clips the
    form too) is widened, so a glyph box renders differently only where a
    clip took ink. Use as a context manager.
    """

    def __init__(self, path: Path) -> None:
        self.printed = pymupdf.open(path)
        try:
            self.free = pymupdf.open(path)
            forms = {xref for xref in range(1, self.free.xref_length())
                     if self.free.xref_get_key(xref, "Subtype") == ("name", "/Form")}
            streams = {xref for page in self.free for xref in page.get_contents()} | forms
            for xref in sorted(streams):
                self.free.update_stream(xref, without_clipping(self.free.xref_stream(xref)))
            for xref in sorted(forms):
                self.free.xref_set_key(xref, "BBox", "[-100000 -100000 100000 100000]")
        except BaseException:
            self.printed.close()
            if getattr(self, "free", None) is not None:
                self.free.close()
            raise

    def __enter__(self) -> "ClipLoss":
        return self

    def __exit__(self, *exc) -> None:
        self.printed.close()
        self.free.close()

    def page(self, index: int) -> ClipPage:
        return ClipPage(self.printed[index], self.free[index])


# --- the workbook --------------------------------------------------------------------------------------

def _horizontal(alignment: str | None) -> str:
    if alignment in HORIZONTAL_ALIGNMENTS:
        return alignment
    # The map writes every value cell's alignment; anything else (general included) is not judged on a guess.
    raise ValueError(f"the workbook has a horizontal alignment the comparison does not know: {alignment!r}")


def _vertical(alignment: str | None) -> str:
    if alignment in VERTICAL_ALIGNMENTS:
        return alignment
    raise ValueError(f"the workbook has a vertical alignment the comparison does not know: {alignment!r}")


def workbook_cells(result: XlsxResult, pages: Sequence[ScenePage]) -> dict[int, list[CellArea]]:
    """The written workbook's cells with values, by 0-based page of ``pages``, on that page's grid.

    Each sheet's pages start at row 0 and after each of its row breaks; the
    columns and rows are the page grid's (one cell per grid cell, FR-021).
    """
    import openpyxl

    book = openpyxl.load_workbook(result.path)
    if [sheet.title for sheet in book.worksheets] != [info.name for info in result.sheets]:
        raise ValueError("the workbook's sheets are not the ones render_xlsx reported")
    cells: dict[int, list[CellArea]] = {}
    for info, sheet in zip(result.sheets, book.worksheets):
        starts = [0] + sorted(item.id for item in sheet.row_breaks.brk)
        if len(starts) != info.last_page - info.first_page + 1:
            raise ValueError(f"sheet {info.name}: its row breaks make {len(starts)} pages, not "
                             f"{info.last_page - info.first_page + 1}")
        merges = {(item.min_row, item.min_col): item for item in sheet.merged_cells.ranges}
        for row in sheet.iter_rows():
            for item in row:
                if item.value is None:
                    continue
                top, left = item.row - 1, item.column - 1
                merged = merges.get((item.row, item.column))
                bottom, right = (merged.max_row - 1, merged.max_col - 1) if merged else (top, left)
                part = bisect.bisect_right(starts, top) - 1
                index = info.first_page - 1 + part
                grid, first = pages[index].profile.grid, starts[part]
                box = Box(grid.x(left), grid.y(top - first), grid.x(right + 1), grid.y(bottom + 1 - first))
                cells.setdefault(index, []).append(CellArea(
                    box, _horizontal(item.alignment.horizontal), _vertical(item.alignment.vertical), item.coordinate))
    return cells


def workbook_number_cells(path: Path) -> dict[str, int]:
    """Number cells per sheet as the file holds them (Excel's count is checked against these)."""
    import openpyxl

    book = openpyxl.load_workbook(path)
    return {sheet.title: sum(isinstance(item.value, (int, float)) and not isinstance(item.value, bool)
                             for row in sheet.iter_rows() for item in row) for sheet in book.worksheets}


_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_RELATIONSHIPS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_OFFICE_RELATIONSHIP = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_ANCHORS = {"{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}" + name
            for name in ("twoCellAnchor", "oneCellAnchor", "absoluteAnchor")}


_DRAWING_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"
_DOCUMENT_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"


def _part(base: str, target: str) -> str:
    """The package path of a relationship target relative to the part ``base`` ("/..." is from the root)."""
    parts = [] if target.startswith("/") else base.split("/")[:-1]
    for piece in target.split("/"):
        if piece == "..":
            if not parts:
                raise ValueError(f"the relationship target {target!r} of {base!r} leaves the package")
            parts.pop()
        elif piece not in ("", "."):
            parts.append(piece)
    return "/".join(parts)


def _relations(package: zipfile.ZipFile, part: str) -> list[tuple[str, str, str]]:
    """(Id, Type, target part) of the internal relationships of ``part`` ("" is the package)."""
    folder, _, name = part.rpartition("/")
    rels = f"{folder}/_rels/{name}.rels" if folder else f"_rels/{name}.rels"
    if rels not in package.namelist():
        return []
    root = ElementTree.fromstring(package.read(rels))
    return [(item.get("Id"), item.get("Type"), _part(part, item.get("Target")))
            for item in root.iter(f"{_RELATIONSHIPS}Relationship") if item.get("TargetMode") != "External"]


def workbook_shapes(path: Path) -> dict[str, int]:
    """Top-level shapes per sheet as the file's drawings hold them (one per anchor; a group is one)."""
    with zipfile.ZipFile(path) as package:
        workbook = next((target for _, kind, target in _relations(package, "") if kind == _DOCUMENT_TYPE), None)
        if workbook is None:
            raise ValueError(f"{path} has no workbook part (no officeDocument relationship)")
        targets = {identifier: target for identifier, _, target in _relations(package, workbook)}
        shapes = {}
        for sheet in ElementTree.fromstring(package.read(workbook)).iter(f"{_MAIN}sheet"):
            part = targets[sheet.get(f"{_OFFICE_RELATIONSHIP}id")]
            drawings = [target for _, kind, target in _relations(package, part) if kind == _DRAWING_TYPE]
            shapes[sheet.get("name")] = sum(
                sum(child.tag in _ANCHORS for child in ElementTree.fromstring(package.read(drawing)))
                for drawing in drawings)
    return shapes


# --- Excel ----------------------------------------------------------------------------------------------

def excel_missing_reason() -> str | None:
    """Why Excel cannot be driven here, or None."""
    if sys.platform != "win32":
        return "Excel COM checks need Windows"
    import winreg

    try:
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"Excel.Application\CLSID"))
    except OSError:
        return "Microsoft Excel is not installed (no Excel.Application COM class)"
    return None


class ExcelRunner(Protocol):
    def __call__(self, workbook: Path, sheet_pdfs: Sequence[Path], timeout: float) -> dict: ...


_EXCEL_SCRIPT = r"""
$ErrorActionPreference = 'Stop'
$params = Get-Content -Raw -Encoding UTF8 -LiteralPath $args[0] | ConvertFrom-Json
Add-Type -Namespace CompareXlsxPdf -Name Native -MemberDefinition '[DllImport("user32.dll")] public static extern uint GetWindowThreadProcessId(System.IntPtr window, out uint process);'
$logs = @(Get-ChildItem -LiteralPath $env:TEMP -Filter 'error*.xml' -File -ErrorAction SilentlyContinue | ForEach-Object { $_.FullName })
$clock = [System.Diagnostics.Stopwatch]::StartNew()
$xl = New-Object -ComObject Excel.Application
$book = $null
try {
  [uint32]$excelProcess = 0
  [void][CompareXlsxPdf.Native]::GetWindowThreadProcessId([System.IntPtr]$xl.Hwnd, [ref]$excelProcess)
  Set-Content -Encoding ASCII -LiteralPath $args[2] -Value $excelProcess
  $xl.Visible = $false
  $xl.DisplayAlerts = $false
  $xl.ScreenUpdating = $false
  $result = [ordered]@{ version = [string]$xl.Version; build = [string]$xl.Build; printer = [string]$xl.ActivePrinter;
    open_error = $null; repair_logs = @(); sheets = @(); temp = [string]$env:TEMP;
    seconds = [ordered]@{ start = $clock.Elapsed.TotalSeconds } }
  try { $book = $xl.Workbooks.Open($params.path, 0, $true) } catch { $result.open_error = $_.Exception.Message }
  $result.seconds.open = $clock.Elapsed.TotalSeconds
  if ($null -ne $book) {
    $pdfs = @($params.pdfs)
    if ($book.Worksheets.Count -ne $pdfs.Count) { throw "the workbook has $($book.Worksheets.Count) sheets, not $($pdfs.Count)" }
    for ($k = 1; $k -le $book.Worksheets.Count; $k++) {
      $ws = $book.Worksheets.Item($k)
      $begin = $clock.Elapsed.TotalSeconds
      $ws.ExportAsFixedFormat(0, $pdfs[$k - 1])
      $exported = $clock.Elapsed.TotalSeconds
      $numbers = 0; $hashes = @(); $found = $null; $lookup = $null
      # SpecialCells(xlCellTypeConstants, xlNumbers) fails when the sheet has no number; the count is checked
      try { $found = $ws.UsedRange.SpecialCells(2, 1) } catch { $lookup = $_.Exception.Message }
      if ($null -ne $found) {
        foreach ($c in $found.Cells) {
          # SpecialCells also returns the other cells of a merged range; count its top-left cell only
          if ($c.MergeCells -and $c.MergeArea.Cells.Item(1).Address() -ne $c.Address()) { continue }
          $numbers++
          $text = [string]$c.Text
          # ### in a cell too narrow for the number; nothing at all in one too narrow for a single #
          if ($text.Contains('#') -or $text.Trim().Length -eq 0) {
            $hashes += [ordered]@{ address = $c.Address($false, $false); text = $text; value = $c.Value2 } }
        }
      }
      $result.sheets += [ordered]@{ name = $ws.Name; number_cells = $numbers; shapes = $ws.Shapes.Count;
        hashes = $hashes; lookup_error = $lookup; export_seconds = $exported - $begin;
        numbers_seconds = $clock.Elapsed.TotalSeconds - $exported }
    }
  }
  $result.repair_logs = @(Get-ChildItem -LiteralPath $env:TEMP -Filter 'error*.xml' -File -ErrorAction SilentlyContinue |
    Where-Object { $logs -notcontains $_.FullName } | ForEach-Object { $_.FullName })
  $result.seconds.total = $clock.Elapsed.TotalSeconds
  ConvertTo-Json -InputObject $result -Depth 6 -Compress | Set-Content -Encoding UTF8 -LiteralPath $args[1]
} finally {
  # Only the workbook this script opened; a failing Close must not skip Quit.
  if ($null -ne $book) { try { $book.Close($false) } catch { } }
  $xl.Quit()
  [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($xl)
  [GC]::Collect(); [GC]::WaitForPendingFinalizers()
}
"""

# The fields of Excel's report and of each sheet in it, with their types (None allowed where listed).
_REPORT_FIELDS = {"version": (str,), "build": (str,), "printer": (str,), "open_error": (str, type(None)),
                  "repair_logs": (list,), "sheets": (list,), "temp": (str,)}
_SHEET_FIELDS = {"name": (str,), "number_cells": (int,), "shapes": (int,), "hashes": (list,),
                 "lookup_error": (str, type(None))}
_HASH_FIELDS = {"address": (str,), "text": (str,), "value": (int, float)}


def _fields(item, fields_: dict, where: str) -> None:
    if not isinstance(item, dict):
        raise ExcelError(f"{where} of Excel's report is not an object: {item!r}")
    for name, types in fields_.items():
        if name not in item:
            raise ExcelError(f"{where} of Excel's report has no {name!r}")
        value = item[name]
        if isinstance(value, bool) or not isinstance(value, types) or (isinstance(value, int) and value < 0
                                                                        and name in ("number_cells", "shapes")):
            raise ExcelError(f"{where} of Excel's report has an unexpected {name!r}: {value!r}")


def validate_excel_report(report) -> dict:
    """Excel's report as the comparison reads it, or ExcelError naming the field that is missing or odd."""
    _fields(report, _REPORT_FIELDS, "the report")
    for index, sheet in enumerate(report["sheets"], 1):
        _fields(sheet, _SHEET_FIELDS, f"sheet {index}")
        for number, cell in enumerate(sheet["hashes"], 1):
            _fields(cell, _HASH_FIELDS, f"number cell {number} of sheet {index}")
    if report["open_error"] is None and not report["sheets"]:
        raise ExcelError("Excel opened the workbook but reported no sheet")
    return report


def _excel_running(process: int) -> bool:
    """Is the Excel with this PID still running? (The image name guards against a reused PID.)"""
    completed = subprocess.run(["tasklist", "/FI", f"PID eq {process}", "/FI", "IMAGENAME eq EXCEL.EXE",
                                "/FO", "CSV", "/NH"], capture_output=True)
    if completed.returncode != 0:
        raise ExcelError(f"tasklist failed with exit code {completed.returncode}")
    return f'"{process}"' in completed.stdout.decode("ascii", "replace")


def _stop_excel(process_file: Path, failed: bool) -> str:
    """Wait for the Excel this script started to quit; stop it after 20 s. Returns what went wrong."""
    if not process_file.exists():
        return "the Excel process ID was not written; check that no hidden Excel is left" if failed else ""
    text = process_file.read_text(encoding="ascii").strip()
    process = int(text) if text.isdigit() else 0
    if process <= 0:
        return "the Excel process ID could not be read; check that no hidden Excel is left"
    deadline = time.monotonic() + 20
    while _excel_running(process) and time.monotonic() < deadline:
        time.sleep(0.5)
    if not _excel_running(process):
        return ""
    killed = subprocess.run(["taskkill", "/PID", str(process), "/T", "/F"], capture_output=True)
    return f"Excel (PID {process}) did not quit and " + (
        "was stopped" if killed.returncode == 0 else f"could not be stopped (exit {killed.returncode})")


def run_excel(workbook: Path, sheet_pdfs: Sequence[Path], timeout: float) -> dict:
    """Open ``workbook`` in a hidden Excel, export sheet k to ``sheet_pdfs[k]``, count shapes and read numbers."""
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder:
        directory = Path(folder)
        script, inputs, outputs, process_file = (directory / name for name in (
            "compare.ps1", "params.json", "result.json", "excel.pid"))
        script.write_text(_EXCEL_SCRIPT, encoding="utf-8-sig")  # PowerShell 5.1 needs the BOM for non-ASCII
        inputs.write_text(json.dumps({"path": str(Path(workbook).resolve()),
                                      "pdfs": [str(Path(pdf).resolve()) for pdf in sheet_pdfs]},
                                     ensure_ascii=False), encoding="utf-8")
        error = ""
        try:
            completed = subprocess.run(
                ["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(script),
                 str(inputs), str(outputs), str(process_file)], capture_output=True, timeout=timeout)
            if completed.returncode != 0:
                error = (completed.stderr + completed.stdout).decode("oem", "replace")  # the console code page
        except subprocess.TimeoutExpired:
            error = f"the Excel export did not finish in {timeout} s"
        except OSError as exc:  # no PowerShell
            raise ExcelError(f"PowerShell could not be started: {exc}") from exc
        try:
            stopped = _stop_excel(process_file, bool(error))
        except OSError as exc:  # no tasklist or taskkill
            stopped = f"the Excel process could not be checked: {exc}"
        error = "\n".join(part for part in (error.strip(), stopped) if part)
        if error:
            raise ExcelError(error)
        if not outputs.exists():
            raise ExcelError("the Excel script ended without writing its result")
        try:
            report = json.loads(outputs.read_text(encoding="utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ExcelError(f"Excel's report is not JSON: {exc}") from exc
        return validate_excel_report(report)


# --- targets --------------------------------------------------------------------------------------------

@dataclass(frozen=True)
class Target:
    """A report to compare: ``build()`` returns its ScenePages and the FontRegistry (None: the default)."""

    name: str
    sheet_name: str
    build: Callable[[], tuple[Sequence[ScenePage], FontRegistry | None]]

    def __post_init__(self) -> None:
        if not self.name or _SAFE_NAME.search(self.name) or self.name in (".", ".."):
            raise ValueError(f"a target's name is a folder name of letters, digits, '_', '.' and '-': {self.name!r}")


def _load(key: str, path: Path):
    if key not in sys.modules:
        spec = importlib.util.spec_from_file_location(key, path)
        if spec is None or spec.loader is None:
            raise ImportError(f"{path} cannot be loaded as a module")
        module = importlib.util.module_from_spec(spec)
        sys.modules[key] = module
        try:
            spec.loader.exec_module(module)
        except BaseException:
            del sys.modules[key]
            raise
    return sys.modules[key]


def _safe(name: str) -> str:
    return _SAFE_NAME.sub("_", name)


def representative_target(variant: str = "standard", font: str = DEFAULT_FONT) -> Target:
    """The representative example's grid version (term_columns="grid"), as write_report builds it."""

    def build():
        example = _load("representative_report", REPRESENTATIVE)
        if example.SHEET_NAME != REPRESENTATIVE_SHEET:
            raise ValueError(f"the example's sheet is {example.SHEET_NAME!r}, not {REPRESENTATIVE_SHEET!r}")
        image = REPRESENTATIVE.parent / "assets" / "bridge-overview.png"
        pages = example.build_report(image, font=font, variant=variant, term_columns="grid").layout()
        fallen = example.grid_fallbacks(pages)
        if fallen:
            raise example.GridTermFallbackError(f"grid term columns do not fit in {', '.join(fallen)} with the font "
                                                f"{font!r}; they fell back to packed terms", fallen)
        return pages, None

    name = f"representative-{variant}" if font == DEFAULT_FONT else f"representative-{variant}-{font}"
    return Target(_safe(name), REPRESENTATIVE_SHEET, build)


def cantilever_target(source_pages: Sequence[int] | None = None) -> Target:
    """The cantilever p1-37 fixture, laid out through its public entry (builder.layout_report)."""

    def build():
        builder = _load("cantilever_fixture_builder", CANTILEVER_BUILDER)
        fonts = builder.fixture_fonts()
        selected = builder.SOURCE_PAGES if source_pages is None else tuple(source_pages)
        return builder.layout_report(fonts, source_pages=selected), fonts

    name = "cantilever" if source_pages is None else "cantilever-p" + "-".join(f"{page:02}" for page in source_pages)
    return Target(name, CANTILEVER_SHEET, build)


def default_output(name: str) -> Path:
    return REPOSITORY_ROOT / "build" / "reports" / "xlsx-compare" / name


# --- the comparison -------------------------------------------------------------------------------------

CHECKS = ("opens_without_repair", "readback", "strict", "no_hashes", "page_count", "page_size", "positions",
          "no_extra_text")


def verdict(checks: dict[str, bool | None]) -> int:
    """EXIT_PASSED when at least one check ran and every check passed; a check not run (None) fails."""
    return EXIT_PASSED if checks and all(checks.values()) else EXIT_FAILED


def _clear(out: Path) -> None:
    """Remove this script's own files of an earlier run (nothing else in ``out``)."""
    for path in [out / "result.json", out / "library.pdf", out / "workbook.xlsx",
                 *(out / "excel").glob("sheet-*.pdf"), *(out / "side-by-side").glob("page-*.png")]:
        path.unlink(missing_ok=True)


def _finish(out: Path, report: dict, code: int) -> tuple[int, dict]:
    report["exit_code"] = code
    report["passed"] = code == EXIT_PASSED
    (out / "result.json").write_text(json.dumps(report, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return code, report


def _records(items) -> list[dict]:
    def plain(value):
        return list(value) if isinstance(value, tuple) else value

    return [{field.name: plain(getattr(item, field.name)) for field in fields(item)} for item in items[:RECORD_LIMIT]]


def _error(exc: BaseException) -> dict:
    error = {"type": type(exc).__name__, "message": str(exc), "traceback": traceback.format_exception(exc)[-6:]}
    if isinstance(exc, XlsxStrictError):
        error.update(shifts=len(exc.shifts), fallbacks=len(exc.fallbacks),
                     shift_records=_records(exc.shifts), fallback_records=_records(exc.fallbacks))
    return error


def _page_size(page: pymupdf.Page) -> list[float]:
    return [round(page.rect.width, 3), round(page.rect.height, 3)]


def side_by_side(library: pymupdf.Document, excel_pages: Sequence[pymupdf.Page | None], folder: Path,
                 zoom: float) -> list[str]:
    """One PNG per page: the library's page on the left, Excel's on the right (blank when Excel has none)."""
    folder.mkdir(parents=True, exist_ok=True)
    matrix, written = pymupdf.Matrix(zoom, zoom), []
    for number, page in enumerate(library, 1):
        images = []
        for source in (page, excel_pages[number - 1]):
            if source is None:
                images.append(None)
                continue
            pixmap = source.get_pixmap(matrix=matrix, alpha=False)
            images.append(Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples))
        left = images[0]
        right = images[1] or Image.new("RGB", left.size, (235, 235, 235))
        gap = max(4, round(8 * zoom))
        canvas = Image.new("RGB", (left.width + gap + right.width, max(left.height, right.height)), (128, 128, 128))
        canvas.paste(left, (0, 0))
        canvas.paste(right, (left.width + gap, 0))
        path = folder / f"page-{number:02}.png"
        canvas.save(path)
        written.append(str(path))
    return written


def _largest_of(values) -> float | None:
    return max((value for value in values if value is not None), default=None)


def _summary(pages: list[dict]) -> dict:
    def count(category, key):
        return sum(item["category"] == category for page in pages for item in page[key])

    return {
        "pieces": sum(page["pieces"] for page in pages), "matched": sum(page["matched"] for page in pages),
        "missing": count(STANDARD, "missing"), "outside": count(STANDARD, "outside"),
        "max_dx_pt": _largest_of(page["max_dx_pt"] for page in pages),
        "max_dy_pt": _largest_of(page["max_dy_pt"] for page in pages),
        "nearest": sum(page["nearest"] for page in pages),
        "outside_if_paper_axes": sum(page["outside_if_paper_axes"] for page in pages),
        "extra": sum(page["extra_count"] for page in pages),
        "clipped_out": sum(page["clipped_out_count"] for page in pages),
        FR021: {
            "pieces": sum(page[FR021]["pieces"] for page in pages),
            "matched": sum(page[FR021]["matched"] for page in pages),
            "missing": count(FR021, "missing"), "outside": sum(page[FR021]["outside"] for page in pages),
            "max_dx_pt": _largest_of(page[FR021]["max_dx_pt"] for page in pages),
            "max_dy_pt": _largest_of(page[FR021]["max_dy_pt"] for page in pages),
            "max_range_dy_pt": _largest_of(page[FR021]["max_range_dy_pt"] for page in pages),
        },
    }


def _compare(report: dict, checks: dict, out: Path, pages: Sequence[ScenePage], result: XlsxResult,
             printed: dict, sheet_pdfs: Sequence[Path], zoom: float) -> None:
    """The checks on Excel's output (readback, numbers, page count and size, positions) into ``report``."""
    cells = workbook_cells(result, pages)
    numbers, shapes = workbook_number_cells(result.path), workbook_shapes(result.path)
    sheets = printed["sheets"]
    report["shapes"] = {sheet["name"]: {"workbook": shapes.get(sheet["name"]), "excel": sheet["shapes"]}
                        for sheet in sheets}
    checks["readback"] = [sheet["name"] for sheet in sheets] == list(shapes) and all(
        item["workbook"] == item["excel"] for item in report["shapes"].values())
    report["hash_cells"] = [{"sheet": sheet["name"], **cell} for sheet in sheets for cell in sheet["hashes"]]
    report["number_cells"] = {sheet["name"]: sheet["number_cells"] for sheet in sheets}
    lookup_errors = {sheet["name"]: sheet["lookup_error"] for sheet in sheets if sheet["lookup_error"]}
    if lookup_errors:
        report["number_lookup_errors"] = lookup_errors
    checks["no_hashes"] = not report["hash_cells"] and report["number_cells"] == numbers
    with pymupdf.open(result.path.with_name("library.pdf")) as library, ExitStack() as stack:
        losses = [stack.enter_context(ClipLoss(path)) if path.exists() else None for path in sheet_pdfs]
        counts = [0 if loss is None else len(loss.printed) for loss in losses]
        report["page_count"] = {
            "library": len(library), "excel": sum(counts),
            "sheets": [{"name": info.name, "expected": info.last_page - info.first_page + 1, "excel": count}
                       for info, count in zip(result.sheets, counts)]}
        checks["page_count"] = len(library) == len(pages) and all(
            item["expected"] == item["excel"] for item in report["page_count"]["sheets"])
        excel_pages: list[tuple[ClipLoss, int] | None] = [None] * len(library)
        for info, loss in zip(result.sheets, losses):
            for offset in range(info.last_page - info.first_page + 1):
                if loss is not None and offset < len(loss.printed):
                    excel_pages[info.first_page - 1 + offset] = (loss, offset)
        report["pages"] = []
        sizes_match, sizes_compared = True, 0
        for index, page in enumerate(library):
            info = next(item for item in result.sheets if item.first_page <= index + 1 <= item.last_page)
            grid = pages[index].profile.grid
            shown = excel_pages[index]  # (ClipLoss, page of the sheet's PDF) or None
            printed_page = None if shown is None else shown[0].printed[shown[1]]
            clip_page = None if shown is None else shown[0].page(shown[1])
            rows = [] if shown is None else pdf_glyph_rows(printed_page, clip_page)
            compared = match_page(pdf_pieces(page, pages[index]), rows, cells.get(index, ()), grid.pitch_x,
                                  grid.pitch_y)
            summary = page_summary(index + 1, info.name, compared)
            dropped = () if clip_page is None else clip_page.dropped
            summary.update(excel_page=None if shown is None else shown[1] + 1,
                           clipped_out_count=len(dropped),
                           clipped_out=[{"char": char, "box": box.as_list()} for char, box in dropped[:RECORD_LIMIT]],
                           pitch_pt=[grid.pitch_x, grid.pitch_y], library_size_pt=_page_size(page),
                           excel_size_pt=None if printed_page is None else _page_size(printed_page),
                           text_commands=sum(isinstance(command, Text) and bool(command.text.strip())
                                             for command in pages[index].commands))
            if printed_page is not None:
                sizes_compared += 1
                sizes_match = sizes_match and all(
                    abs(a - b) <= PAGE_SIZE_TOLERANCE_PT
                    for a, b in zip(summary["library_size_pt"], summary["excel_size_pt"]))
            report["pages"].append(summary)
        checks["page_size"] = sizes_match and sizes_compared > 0
        report["summary"] = _summary(report["pages"])
        # Every page that draws text yields pieces (else the PDF's text could not be read), all fit, and
        # Excel's clip took no glyph away entirely.
        checks["positions"] = (report["summary"]["pieces"] + report["summary"][FR021]["pieces"] > 0 and all(
            not page["missing"] and not page["outside"] and not page["clipped_out_count"]
            and (page["pieces"] + page[FR021]["pieces"] > 0 or page["text_commands"] == 0)
            for page in report["pages"]))
        checks["no_extra_text"] = report["summary"]["extra"] == 0
        report["side_by_side"] = side_by_side(
            library, [None if item is None else item[0].printed[item[1]] for item in excel_pages],
            out / "side-by-side", zoom)


def run(target: Target, out: Path, *, excel: ExcelRunner = run_excel, zoom: float = 1.5, timeout: float = 900.0,
        strict: bool = True) -> tuple[int, dict]:
    """Write, export and compare ``target`` in ``out``; returns the exit code and the report (also result.json)."""
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    _clear(out)
    library_pdf, workbook = out / "library.pdf", out / "workbook.xlsx"
    report: dict = {"target": target.name, "strict": strict, "passed": False, "exit_code": None,
                    "files": {"library_pdf": str(library_pdf), "workbook": str(workbook)}}
    try:
        built, fonts = target.build()
        pages = tuple(built)
        result = render_xlsx([(target.sheet_name, pages)], workbook, fonts, strict=strict)
        render_pdf(pages, library_pdf, fonts)
    except OSError as exc:  # the output (a file held open, a full disk), not the library's report
        report["internal_error"], report["checks_complete"] = _error(exc), False
        print(f"{target.name}: the output could not be written: {type(exc).__name__}: {exc}", file=sys.stderr)
        return _finish(out, report, EXIT_INTERNAL_ERROR)
    except Exception as exc:  # reported as it is: the library could not write this report
        report["library_error"], report["checks_complete"] = _error(exc), False
        print(f"{target.name}: the library could not write the report: {type(exc).__name__}: {exc}", file=sys.stderr)
        return _finish(out, report, EXIT_LIBRARY_ERROR)
    checks: dict[str, bool | None] = {}
    try:
        report["xlsx"] = {
            "sheets": [{"name": info.name, "first_page": info.first_page, "last_page": info.last_page,
                        "scale": info.scale} for info in result.sheets],
            "shifts": len(result.shifts), "fallbacks": len(result.fallbacks),
            "numbers_as_text": [{"page": item.page, "detail": item.detail} for item in result.numbers_as_text],
            "texts_clipped": [{"page": item.page, "detail": item.detail} for item in result.texts_clipped],
            "dropped": len(result.dropped), "untagged": result.untagged_count,
        }
        sheet_pdfs = [out / "excel" / f"sheet-{index:02}.pdf" for index in range(1, len(result.sheets) + 1)]
        sheet_pdfs[0].parent.mkdir(parents=True, exist_ok=True)
        report["files"]["excel_pdfs"] = [str(path) for path in sheet_pdfs]
        try:
            printed = validate_excel_report(excel(workbook, sheet_pdfs, timeout))
        except ExcelError as exc:
            report["excel_error"], report["checks_complete"] = str(exc), False
            print(f"{target.name}: Excel did not finish: {exc}", file=sys.stderr)
            return _finish(out, report, EXIT_EXCEL_ERROR)
        report["excel"] = {name: printed.get(name) for name in ("version", "build", "printer", "open_error",
                                                                 "repair_logs", "temp", "seconds")}
        checks.update(dict.fromkeys(CHECKS))  # None until run: a check left None never passes
        checks["opens_without_repair"] = printed["open_error"] is None and not printed["repair_logs"]
        checks["strict"] = not result.shifts and not result.fallbacks
        if printed["open_error"] is not None:  # nothing else could be checked: null, and not complete
            checks.update(readback=None, no_hashes=None, page_count=None, page_size=None, positions=None,
                          no_extra_text=None)
            report["checks"], report["checks_complete"] = checks, False
            return _finish(out, report, EXIT_FAILED)
        _compare(report, checks, out, pages, result, printed, sheet_pdfs, zoom)
    except Exception as exc:  # the comparison itself failed: never reported as a check result
        report["internal_error"] = _error(exc)
        report["checks"], report["checks_complete"] = checks, False
        print(f"{target.name}: the comparison failed: {type(exc).__name__}: {exc}", file=sys.stderr)
        return _finish(out, report, EXIT_INTERNAL_ERROR)
    report["checks"], report["checks_complete"] = checks, True
    return _finish(out, report, verdict(checks))


# --- command line ---------------------------------------------------------------------------------------

def _pages(text: str) -> tuple[int, ...]:
    try:
        pages = tuple(int(item) for item in text.split(","))
    except ValueError:
        raise argparse.ArgumentTypeError("pages must be comma-separated page numbers") from None
    if any(page not in CANTILEVER_PAGES for page in pages) or any(a >= b for a, b in zip(pages, pages[1:])):
        raise argparse.ArgumentTypeError(
            f"pages must be increasing original pages from {CANTILEVER_PAGES[0]} to {CANTILEVER_PAGES[-1]}")
    return pages


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--example", choices=("representative", "cantilever"), required=True)
    parser.add_argument("--variant", choices=("standard", "long"), help="representative only (default standard)")
    parser.add_argument("--font", help=f"representative only (default {DEFAULT_FONT})")
    parser.add_argument("--pages", type=_pages, help="cantilever only: increasing original pages (default all 37)")
    parser.add_argument("--out", type=Path, help="output folder (default build/reports/xlsx-compare/<target>/)")
    parser.add_argument("--strict", action=argparse.BooleanOptionalAction, default=True,
                        help="write the workbook with strict=True (default); --no-strict still fails on shifts")
    parser.add_argument("--zoom", type=float, default=1.5, help="side-by-side image scale, 0.1 to 8 (default 1.5)")
    parser.add_argument("--timeout", type=float, default=900.0, help="seconds for Excel (default 900)")
    options = parser.parse_args(argv)
    if options.example == "cantilever" and (options.variant is not None or options.font is not None):
        parser.error("--variant and --font apply to --example representative")
    if options.example == "representative" and options.pages is not None:
        parser.error("--pages applies to --example cantilever")
    if options.font is not None and not options.font.strip():
        parser.error("--font must name a font")
    if not 0.1 <= options.zoom <= 8 or not (math.isfinite(options.timeout) and options.timeout > 0):
        parser.error("--zoom must be from 0.1 to 8 and --timeout a positive number of seconds")
    if options.example == "representative":
        target = representative_target(options.variant or "standard", options.font or DEFAULT_FONT)
    else:
        target = cantilever_target(options.pages)
    out = options.out or default_output(target.name)
    missing = excel_missing_reason()
    try:
        if missing is not None:
            out.mkdir(parents=True, exist_ok=True)
            _clear(out)
            print(f"{target.name}: not compared: {missing}", file=sys.stderr)
            _finish(out, {"target": target.name, "strict": options.strict, "reason": missing,
                          "checks_complete": False}, EXIT_NO_EXCEL)
            return EXIT_NO_EXCEL
        code, report = run(target, out, zoom=options.zoom, timeout=options.timeout, strict=options.strict)
    except OSError as exc:  # the output folder: no result.json can be written there
        print(f"{target.name}: the output in {out} could not be written: {exc}", file=sys.stderr)
        return EXIT_INTERNAL_ERROR
    print(f"{target.name}: {'passed' if code == EXIT_PASSED else 'FAILED'} (exit {code}); {out / 'result.json'}")
    if "checks" in report:
        print("checks: " + ", ".join(f"{name}={value}" for name, value in report["checks"].items()))
    if "excel" in report:
        print(f"Excel {report['excel']['version']} build {report['excel']['build']}, "
              f"printer {report['excel']['printer']}")
    if "summary" in report:
        print("standard: pieces {pieces}, matched {matched}, missing {missing}, outside {outside}, "
              "max dx {max_dx_pt} pt, max dy {max_dy_pt} pt; nearest anchors {nearest}, "
              "outside on paper axes {outside_if_paper_axes}, Excel texts no piece took {extra}, "
              "glyphs Excel's clip took away {clipped_out}"
              .format(**report["summary"]))
        print(f"{FR021}: " + "pieces {pieces}, matched {matched}, missing {missing}, outside {outside}, "
              "max dx {max_dx_pt} pt, max dy {max_dy_pt} pt, max dy from the merged range's middle "
              "{max_range_dy_pt} pt".format(**report["summary"][FR021]))
    return code


if __name__ == "__main__":
    raise SystemExit(main())
